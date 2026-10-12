"""Read-only, instance-aware assembly traversal and atomic XLSX export."""

from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
import os
from pathlib import Path, PureWindowsPath
import re
from tempfile import NamedTemporaryFile

from dxf_dwg import exporter as sw
from solidworks.connection import connect_to_solidworks

pythoncom, win32com = sw.pythoncom, sw.win32com


class Cancelled(Exception):
    pass


def _check_cancel(stop_event):
    if stop_event is not None and stop_event.is_set():
        raise Cancelled()


def name_key(value):
    """Turkish alphabet, numeric segments: Parça 2 precedes Parça 10."""
    alphabet = 'abcçdefgğhıijklmnoöprsştuüvyz'
    text = str(value).replace('I', 'ı').replace('İ', 'i').casefold()
    return tuple((0, int(chunk)) if chunk.isdecimal() else
                 (1, tuple(alphabet.index(c) if c in alphabet else ord(c)+100 for c in chunk))
                 for chunk in re.split(r'(\d+)', text) if chunk)


@dataclass(frozen=True)
class Part:
    name: str
    path: str
    configuration: str
    quantity: int


@dataclass(frozen=True)
class Node:
    number: str
    parent: str
    level: int
    name: str
    instance: str
    kind: str
    path: str
    configuration: str


@dataclass(frozen=True)
class Result:
    assembly: str
    configuration: str
    parts: tuple
    tree: tuple
    suppressed: int
    excluded: int
    include_excluded: bool

    @property
    def total_parts(self):
        return sum(part.quantity for part in self.parts)


def _assembly(app):
    model = sw._value(app, 'ActiveDoc')
    if model is None or sw._value(model, 'GetType') != sw.ASSEMBLY:
        raise RuntimeError("BOM için SolidWorks'te ana montajı etkinleştirin.")
    api = model
    if not hasattr(api, 'ResolveAllLightWeightComponents'):
        for interface in ('IAssemblyDoc', 'AssemblyDoc'):
            try:
                api = win32com.client.CastTo(model, interface)
                if hasattr(api, 'ResolveAllLightWeightComponents'):
                    break
            except Exception:
                continue
        else:
            raise RuntimeError('SolidWorks montaj API’sine erişilemiyor.')
    return model, api


def _excluded(component):
    if hasattr(component, 'GetExcludeFromBOM2'):
        values = component.GetExcludeFromBOM2(1, None)  # swThisConfiguration
        if not isinstance(values, (tuple, list)) or len(values) != 1:
            raise RuntimeError('Bileşenin aktif konfigürasyondaki BOM dışı durumu okunamadı.')
        return bool(values[0])
    # Documented legacy property for older dispatch interfaces.
    return bool(sw._value(component, 'ExcludeFromBOM'))


def collect(app, stop_event=None, *, include_excluded=False, progress=None):
    """Count each physical occurrence; never deduplicate subassembly traversal.

    Hidden components are included; suppressed branches are omitted. Excluded
    BOM components and their descendants are omitted unless explicitly enabled.
    No feature edits, configuration switches, file saves or document closes.
    """
    _check_cancel(stop_event)
    assembly, api = _assembly(app)
    resolved = api.ResolveAllLightWeightComponents(True)
    if resolved not in (0, 2):
        raise RuntimeError(f'Montaj bileşenleri çözümlenemedi. API hata kodu: {resolved}')
    config = assembly.ConfigurationManager.ActiveConfiguration
    if hasattr(config, 'IsSpeedPak') and sw._value(config, 'IsSpeedPak'):
        raise RuntimeError('SpeedPak yerine tam montaj konfigürasyonunu etkinleştirin.')
    root = config.GetRootComponent3(True)
    if root is None:
        raise RuntimeError('Montaj bileşen ağacı okunamadı.')
    configuration = str(sw._value(config, 'Name'))
    assembly_name = str(sw._document_name(assembly))
    counts, originals, tree = Counter(), {}, []
    suppressed = excluded = visited = 0
    root_key = sw._part_key(assembly_name, configuration)
    # Sibling records are read once and sorted before entering the traversal.
    def children(component):
        nonlocal suppressed, excluded
        records = []
        for child in sw._value(component, 'GetChildren') or ():
            _check_cancel(stop_event)
            state = sw._value(child, 'GetSuppression')
            if state == 0:
                suppressed += 1
                continue
            if state not in (2, 3):
                raise RuntimeError('Çözümlenmemiş bileşen var; eksik BOM oluşturulmadı.')
            if not include_excluded and _excluded(child):
                excluded += 1
                continue
            path = str(sw._value(child, 'GetPathName') or '')
            ref = str(sw._value(child, 'ReferencedConfiguration') or '')
            if not path or not ref:
                raise RuntimeError('Bir bileşenin dosya yolu veya konfigürasyonu okunamadı.')
            filename = PureWindowsPath(path)
            kind = {'.sldprt': 'Parça', '.sldasm': 'Montaj'}.get(filename.suffix.lower())
            if kind is None:
                raise RuntimeError(f'Tanınmayan bileşen dosyası: {path}')
            instance = str(sw._value(child, 'Name2'))
            records.append((child, path, ref, filename.stem, instance, kind))
        return sorted(records, key=lambda r: (name_key(r[3]), name_key(r[2]), r[1], name_key(r[4])))

    stack = [(record, str(index), '', 1, frozenset((root_key,)))
             for index, record in reversed(list(enumerate(children(root), 1)))]
    while stack:
        _check_cancel(stop_event)
        (component, path, ref, name, instance, kind), number, parent, level, ancestors = stack.pop()
        key = sw._part_key(path, ref)
        tree.append(Node(number, parent, level, name, instance, kind, path, ref))
        if kind == 'Parça':
            counts[key] += 1
            originals.setdefault(key, (name, path, ref))
        else:
            if key in ancestors:
                raise RuntimeError(f'Montaj ağacında döngü bulundu: {path}')
            if sw._value(component, 'GetModelDoc2') is None:
                raise RuntimeError(f'Alt montaj yüklenmemiş: {path}')
            if hasattr(component, 'IsSpeedPak') and sw._value(component, 'IsSpeedPak'):
                raise RuntimeError(f'SpeedPak alt montajın tam ağacını açın: {path}')
            stack.extend((record, f'{number}.{index}', number, level+1, ancestors | {key})
                         for index, record in reversed(list(enumerate(children(component), 1))))
        visited += 1
        if progress and visited % 100 == 0:
            progress(f'Montaj taranıyor · {visited} bileşen okundu')
    _check_cancel(stop_event)
    parts = tuple(sorted((Part(*originals[key], count) for key, count in counts.items()),
                         key=lambda p: (name_key(p.name), name_key(p.configuration), p.path)))
    return Result(assembly_name, configuration, parts, tuple(tree), suppressed, excluded, include_excluded)


def export_excel(result, destination, stop_event=None):
    """Write literal strings and numeric quantities; replace only after success."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    _check_cancel(stop_event)
    destination = Path(destination)
    if destination.suffix.lower() != '.xlsx':
        raise ValueError('Excel dosyası .xlsx uzantılı olmalı.')
    if max(len(result.parts), len(result.tree)) > 1048572:
        raise ValueError('Excel satır sınırı aşıldı; dosya oluşturulmadı.')
    book = Workbook()
    book.remove(book.active)
    created = datetime.now(timezone.utc).astimezone(timezone(timedelta(hours=3)))
    def sheet(title, headers, rows, widths):
        ws = book.create_sheet(title)
        ws.append([title])
        ws.append([result.assembly, result.configuration, created.isoformat(timespec='seconds')])
        ws.append([f'Bastırılmış dal: {result.suppressed}', f'BOM dışı dal: {result.excluded}',
                   'BOM dışında bırakılanlar dahil' if result.include_excluded else 'BOM dışında bırakılanlar hariç'])
        ws.append(headers)
        for values in rows:
            _check_cancel(stop_event)
            ws.append(values)
        # Names and paths are user-controlled strings, never Excel formulas.
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top')
        for cell in ws[4]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='177C65')
        ws['A1'].font = Font(size=16, bold=True, color='177C65')
        ws.freeze_panes = 'A5'
        ws.auto_filter.ref = f'A4:{get_column_letter(len(headers))}{max(4, ws.max_row)}'
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        return ws
    temporary = None
    try:
        sheet('Parça listesi', ['Sıra', 'Parça adı', 'Adet', 'Konfigürasyon', 'Dosya adı', 'Dosya yolu'],
              ([i, p.name, p.quantity, p.configuration, PureWindowsPath(p.path).name, p.path]
               for i, p in enumerate(result.parts, 1)), (8, 34, 10, 28, 40, 80))
        tree_sheet = sheet('Ürün ağacı', ['Pozisyon', 'Üst pozisyon', 'Seviye', 'Parça / montaj adı',
                            'Bileşen örneği', 'Tür', 'Adet', 'Konfigürasyon', 'Dosya yolu'],
              ([n.number, n.parent, n.level, n.name, n.instance, n.kind, 1, n.configuration, n.path]
               for n in result.tree), (16, 16, 10, 34, 40, 12, 10, 28, 80))
        tree_sheet.sheet_properties.outlinePr.summaryBelow = False
        for row, node in enumerate(result.tree, 5):
            _check_cancel(stop_event)
            tree_sheet.row_dimensions[row].outlineLevel = min(node.level-1, 7)
            tree_sheet.cell(row, 4).alignment = Alignment(indent=min(node.level-1, 250))
            if node.kind == 'Montaj':
                for cell in tree_sheet[row]:
                    cell.font = Font(bold=True, color='177C65')
                    cell.fill = PatternFill('solid', fgColor='E7F4EE')
        _check_cancel(stop_event)
        with NamedTemporaryFile(dir=destination.parent, prefix='.gokworks-bom-', suffix='.xlsx', delete=False) as stream:
            temporary = Path(stream.name)
        book.save(temporary)
        _check_cancel(stop_event)
        os.replace(temporary, destination)
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def scan_worker(events, stop_event, *, include_excluded=False):
    initialized = False
    try:
        if pythoncom is None or win32com is None:
            raise RuntimeError('Montaj taraması Windows, pywin32 ve açık SolidWorks oturumu gerektirir.')
        pythoncom.CoInitialize()
        initialized = True
        result = collect(connect_to_solidworks(), stop_event, include_excluded=include_excluded,
                         progress=lambda text: events.put(('bom_phase', text)))
        outcome = dict(state='completed', result=result)
    except Cancelled:
        outcome = dict(state='cancelled')
    except Exception as error:
        outcome = dict(state='error', error=str(error))
    finally:
        if initialized:
            try:
                pythoncom.CoUninitialize()
            except Exception as error:
                outcome = dict(state='error', error=f'SolidWorks bağlantısı temizlenemedi: {error}')
        events.put(('bom_done', outcome))
