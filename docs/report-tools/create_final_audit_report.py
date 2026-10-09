"""Son kontrolün test kanıtlarını ve gerçek ekranlarını Word'e aktar."""

from datetime import datetime
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor

from create_word_report import paragraph, table, field

ROOT = Path(__file__).resolve().parents[2]
FOLDER = ROOT/'docs/reports/final-audit'


def main():
    data = json.loads((FOLDER/'gui-verification-results.json').read_text())
    assert data['successful'] and data['outcomes'] == {'passed': data['tests_run']}
    for filename, digest in data['sha256'].items():
        assert hashlib.sha256((ROOT/filename).read_bytes()).hexdigest() == digest, filename
    screenshots = json.loads((FOLDER/'screenshots.json').read_text())
    for item in screenshots:
        assert hashlib.sha256((ROOT/item['file']).read_bytes()).hexdigest() == item['sha256']
    refs = json.loads((FOLDER/'reference-comparison.json').read_text())
    coverage = json.loads((FOLDER/'coverage.json').read_text())['totals']
    gui = [r for r in data['tests'] if r['id'].startswith(('test_export_window.', 'test_gui_audit.', 'test_final_gui.'))
           or '.KGuiTests.' in r['id'] or '.ManualGuiTests.' in r['id']]
    assert len(gui) == 49
    matrix = [s for s in data['subcases'] if s['test'].endswith('resolution_scale_matrix')]
    menus = [s for s in data['subcases'] if s['test'].endswith('real_menu_command')]
    when = datetime.fromisoformat(data['finished_utc']).astimezone(ZoneInfo('Europe/Istanbul'))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in [('Normal', 10), ('Title', 26), ('Heading 1', 19), ('Heading 2', 13)]:
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string('176B57' if name.startswith('Heading') else '243041')
    section.header.paragraphs[0].text = 'GOKWORKS · Son program ve GUI kontrolü'
    footer = section.footer.paragraphs[0]
    footer.add_run('GOKWORKS | '); field(footer, 'PAGE'); footer.add_run(' / '); field(footer, 'NUMPAGES')
    doc.add_heading('Son program ve GUI\nkontrol raporu', 0)
    paragraph(doc, when.strftime('%d.%m.%Y %H:%M:%S — Türkiye saati (UTC+3)'))
    paragraph(doc, f"{data['tests_run']} test geçti; 0 başarısız, 0 atlanan. {len(gui)} gerçek GUI testi.", True)
    paragraph(doc, f"{len(matrix)} boyut/tema/ekran/sekme kombinasyonu ve {len(menus)} gerçek menü seçimi geçti. Yeşil vurgu renkleri korundu.")
    paragraph(doc, 'Testler gerçek Tk pencereleri, gerçek ezdxf okuyucusu ve SolidWorks COM davranışını temsil eden nesnelerle Linux ortamında çalıştırıldı. Gerçek Windows/SolidWorks montajı ve üretim açınım geometrisi burada doğrulanmadı.')
    table(doc, ('Alan', 'Sonuç'), [
        ('PDF', '34 test; belge kimliği, aktivasyon, uzun geçici ad, klasör hedefi ve COM kapanış hatası kontrol edildi.'),
        ('DXF/DWG', '69 test; aktarım, dosya güvenliği, açınım seçenekleri, kalınlık ve alt montaj adedi.'),
        ('K/R', '39 hesaplama/COM testi; tablo, BA, manuel eşleştirme, geri okuma, konfigürasyon ve geri alma.'),
        ('GUI', '49 test; PDF, DXF/DWG, otomatik ve manuel K/R, filtre, CSV, ayrıntılar, tema, iptal ve kapanış.'),
        ('Diğer', '8 test; ayar dosyaları, SolidWorks bağlantısı ve giriş noktası.'),
    ], (3.4, 13.6))
    paragraph(doc, f"Satır kapsamı %{coverage['percent_statements_covered']:.2f}; satır/dal birleşik kapsamı %{coverage['percent_covered']:.2f}. Uygulama {data['lines']['application']} fiziksel Python satırı, testler {data['lines']['tests']} satır.")
    doc.add_heading('GitHub kaynaklarıyla karşılaştırma', 1)
    for text in [
        'CodeStack ve xCAD: ExportToDWG2, PDF sayfa seçimi/SaveAs, parça arayüzü, aktif belge ve montaj konfigürasyonları karşılaştırıldı. Windows gerektiren dış makrolar çalıştırılmadı.',
        'ExportDxf ve pywin32 b312: kalınlık/metre dönüşümü, COM varsayılan üyesi, geç bağlama ve BYREF davranışı incelendi.',
        f"FreeCAD SheetMetal: {refs['formula_cases_passed']} BA kombinasyonunda dış formül gerçekten yürütüldü ve sonuçlar 1e-12 toleransla eşleşti. Genel ANSI tablosundaki 9 hücre eşleşti. Bu, üretim kalibrasyonu değildir.",
        'SolidDNA 2023 interop tanımları: SheetMetal, özel büküm payı ve K tipi=2 sözleşmesi önceki metadata kanıtlarıyla karşılaştırıldı.',
        'CustomTkinter 6.0: dış complex_example.py gerçekten açıldı; iki tema ve %100/%125/%150 ölçek seçimi callback hatası vermeden çalıştı.',
        'Kaynak commitleri, dosya eşleştirmesi ve tüm kapsam sınırları docs/reports/final-audit/README.md içinde; sayısal kaynak kanıtı reference-comparison.json dosyasında.',
    ]: paragraph(doc, text)
    doc.add_heading('Çözünürlük ve DPI kontrolü', 1)
    paragraph(doc, '%100: 1000×690, 1240×860, 1366×768, 1600×900, 1920×1080. %125: 1400×900. %150: 1920×1080. İki tema, dört görünüm ve Dosyalar/Günlük sekmeleri sınandı. Asgari çalışma alanı 1000×690 mantıksal pikseldir.')
    paragraph(doc, 'İlk pencere ekran ölçülerine göre seçiliyor. 1366×768/%100 ve 1920×1080/%125-%150 açılışları, temsil edilen ekran ölçüleri ve gerçek Tk ölçeklemesiyle sınandı. Gerçek Windows DPI ve çoklu monitör çalıştırması değildir.')
    for title, stem in [('PDF', 'pdf'), ('DXF / DWG', 'dxf-dwg'), ('K FAKTÖR - RADIUS · Malzeme / Takım', 'k-radius'), ('K FAKTÖR - RADIUS · Manuel K / R', 'manual-k-radius')]:
        doc.add_page_break(); doc.add_heading(title, 1)
        for suffix, caption in [('light', 'Açık tema'), ('dark', 'Koyu tema')]:
            doc.add_picture(str(FOLDER/'screenshots'/f'{stem}-{suffix}.png'), width=Inches(6.35))
            paragraph(doc, caption+' · Gerçek 1240×860 Tk görüntüsü; tablo boş, üretim CAD verisi gösterilmez.')
    doc.add_page_break(); doc.add_heading('Akşam SolidWorks kontrolü', 1)
    paragraph(doc, 'Aktif montajda 2 mm ve farklı kalınlıkta iki sac kullanın. Manuel K/R sekmesine 2 / 0,274 / 2 girin. Önce Montajı tara, sonra Montaja K/R uygula komutunu çalıştırın. 2 mm sacın SheetMetal K/R değerlerini ve farklı kalınlıktaki sacın korunmasını kontrol edin. Kayıt seçeneği kapalıysa değişiklik bellektedir.')
    paragraph(doc, 'DXF/DWG açınımının ölçülerini, kontur/büküm çizgilerini ve dosya adındaki toplam adedi; PDF sayfa ve çizim içeriğini gerçek dosyalarda kontrol edin. DWG doğrulaması başlık/boyut kontrolüdür; tam DWG ayrıştırması yapılmaz.')
    paragraph(doc, 'Kanıtlar: gui-verification-results.json, test-run.log, coverage.json, reference-comparison.json, screenshots.json. Bu raporun paket ve içerik bütünlüğü python-docx/ZIP ile kontrol edildi; Microsoft Word içinde yerel sayfa görünümü çalıştırılmadı.')
    path = FOLDER/'GOKWORKS_Son_Kontrol_2026-10-09.docx'
    doc.save(path)
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        images = [name for name in archive.namelist() if name.startswith('word/media/')]
        assert len(images) == 8
    readback = '\n'.join(p.text for p in Document(path).paragraphs)
    assert all(term in readback for term in ('199 test', '49 gerçek GUI', '108 BA', '112 boyut', '144 gerçek menü'))
    (FOLDER/'word-report-validation.json').write_text(json.dumps(dict(
        file=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(), bytes=path.stat().st_size,
        tests=data['tests_run'], gui_tests=len(gui), embedded_images=len(images), zip_crc_ok=True,
        python_docx_readback_ok=True, native_word_render_verified=False), indent=2)+'\n')
    print(path)


if __name__ == '__main__':
    main()
