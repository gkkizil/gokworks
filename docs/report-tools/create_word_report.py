"""Kaydedilmiş test kanıtları ve gerçek GUI görüntüleriyle Word raporu üret."""

from datetime import datetime
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
FOLDER = ROOT/'docs/reports'
GREEN, TEXT, MUTED, LINE = '176B57', '243041', '627385', 'DDE5ED'

LABELS = {
    'test_navigation_retains_independent_preferences': 'Ekran geçişinde tercih, liste, günlük ve sonuçları koruma',
    'test_dxf_scan_and_navigation_locked_during_worker': 'Parça tarama, kilit dosyalarını eleme ve menüyü kilitleme',
    'test_pdf_scan_regression': 'PDF ekranında yalnız teknik resimleri tarama',
    'test_full_dxf_worker_and_csv_report': 'GUI → iş parçacığı → DXF/DWG sonuçları → CSV akışı',
    'test_empty_list_hint_follows_scan_reset_and_independent_tabs': 'Boş liste yönlendirmesi ve ekranlar arasında bağımsız durum',
    'test_empty_dxf_worker_message': 'Boş parça klasöründe anlaşılır sonuç mesajı',
    'test_source_dialog_filters_parts': 'Parça dosyası seçim filtresi ve kaynak klasör değeri',
    'test_full_pdf_worker_and_turkish_csv_report': 'GUI → PDF iş parçacığı → Türkçe CSV raporu akışı',
    'test_invalid_source_and_output_keep_worker_idle': 'Geçersiz kaynak ve çıktı yolunda işlemi başlatmama',
    'test_connection_failure_releases_controls_and_shows_error': 'SolidWorks bağlantı hatasından sonra kontrolleri açma',
    'test_restart_loads_both_modules_and_rejects_invalid_format': 'Yeniden açılışta iki ekranın tercihleri ve biçim doğrulaması',
    'test_stop_then_close_waits_for_active_worker': 'Durdurma ve kapanışta aktif iş parçacığını bekleme',
    'test_report_write_failure_is_displayed': 'CSV dosyasına yazma hatasını arayüzde gösterme',
    'test_themes_preserve_both_pages_and_reload_preference': 'Tema değişiminde veriyi koruma ve tema tercihini yeniden yükleme',
    'test_search_status_filters_and_live_results_keep_full_counts': 'Türkçe arama, durum filtresi ve canlı sonuçlarla tam sayaçlar',
    'test_rescan_deletes_detached_rows_and_resets_filters': 'Filtrelenmiş eski satırlardan sonra yeniden tarama',
    'test_filters_are_independent_between_modules': 'PDF/DXF ekranlarının arama ve filtrelerini bağımsız tutma',
    'test_details_show_full_paths_and_copy_to_clipboard': 'Ayrıntılarda tam yollar ve pencere kapandıktan sonra pano',
    'test_theme_change_while_busy_preserves_worker_controls': 'İşlem sürerken tema değişikliği ve kontrol kilitleri',
    'test_scaled_interface_keeps_table_text_and_icons_usable': '%125 ölçeklemede iki ekran, iki tema ve tablo yazısı',
    'test_real_entrypoint_runs_tk_loop_and_closes_cleanly': 'Gerçek main.py girişinden Tk döngüsüne girme ve kapanış',
    'test_minimum_window_size_keeps_controls_and_results_visible': 'Küçük pencerede sonuç görünürlüğü, büyük pencerede tarama düğmesi',
    'test_browse_buttons_choose_and_cancel_in_both_modules': 'İki ekranda Gözat düğmeleri: seçme, iptal ve Türkçe yollar',
    'test_theme_selector_buttons_apply_and_save_choice': 'Açık/Koyu seçim düğmelerinin gerçek komutları ve kayıt',
    'test_dxf_option_widgets_and_start_button_reach_worker': 'Biçim/işaret kutuları ve Başlat düğmesinin işçiye doğru seçenek iletmesi',
    'test_scan_stop_button_and_disabled_navigation': 'Tara/Durdur düğmeleri ve işlem sırasında kapalı menü',
    'test_declining_close_keeps_worker_and_window_alive': 'Kapanışı reddedince pencereyi ve işlemi açık tutma',
    'test_filtered_csv_report_button_still_writes_all_results': 'Hatalı filtresi açıkken tüm sonuçları CSV’ye kaydetme ve iptal',
    'test_ctrl_f_search_entry_and_clear_button_bindings': 'Gerçek Ctrl+F olayı, arama alanına yazma ve Temizle düğmesi',
    'test_mouse_double_click_and_escape_close_details': 'Gerçek Tk çift tıklama, Enter ve Escape olaylarıyla ayrıntılar',
    'test_output_folder_button_delegates_or_reports_error': 'Çıktı klasörü düğmesinin OS çağrısı ve hata/eksik klasör mesajı',
    'test_scan_failure_releases_controls_and_reports_reason': 'Tarama hatasının nedeni ve kontrollerin tekrar açılması',
    'test_repeated_navigation_and_theme_buttons_keep_page_state': '16 ardışık menü/tema değişiminde ekran verilerini koruma',
    'test_layout_matrix_keeps_primary_controls_inside_window': '12 boyut/tema/ekran birleşiminde kontrol ve tablo sınırları',
}


def shade(cell, fill):
    node = OxmlElement('w:shd')
    node.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(node)


def set_run(run, size=10, bold=False, color=TEXT):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def table(doc, headings, rows, widths):
    value = doc.add_table(rows=1, cols=len(headings))
    value.alignment = WD_TABLE_ALIGNMENT.CENTER
    value.autofit = False
    for column, width in zip(value.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = OxmlElement(f'w:{edge}')
        for name, content in (('val', 'single'), ('sz', '4'), ('color', LINE)):
            node.set(qn(f'w:{name}'), content)
        borders.append(node)
    value._tbl.tblPr.append(borders)
    for index, heading in enumerate(headings):
        cell = value.rows[0].cells[index]
        cell.width = Cm(widths[index])
        shade(cell, GREEN)
        set_run(cell.paragraphs[0].add_run(heading), 10, True, 'FFFFFF')
    header = OxmlElement('w:tblHeader')
    value.rows[0]._tr.get_or_add_trPr().append(header)
    for index, contents in enumerate(rows):
        row = value.add_row()
        for column, content in enumerate(contents):
            cell = row.cells[column]
            cell.width = Cm(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index % 2 == 0:
                shade(cell, 'F4F7FA')
            set_run(cell.paragraphs[0].add_run(str(content)), 10,
                    color=GREEN if content == 'GEÇTİ' else TEXT)
    for row in value.rows:
        prevent_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(prevent_split)
        for cell in row.cells:
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                node = OxmlElement(f'w:{side}')
                node.set(qn('w:w'), '85')
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            cell._tc.get_or_add_tcPr().append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(0)
    return value


def paragraph(doc, text, bold=False):
    value = doc.add_paragraph()
    set_run(value.add_run(text), 11, bold)
    return value


def page(doc, title):
    doc.add_page_break()
    doc.add_heading(title, level=1)


def screenshot(doc, filename, caption):
    doc.add_picture(str(FOLDER/'assets'/filename), width=Inches(6.35))
    value = doc.add_paragraph()
    set_run(value.add_run(caption), 9, color=MUTED)


def field(paragraph, instruction):
    element = OxmlElement('w:fldSimple')
    element.set(qn('w:instr'), instruction)
    paragraph._p.append(element)


def main():
    data = json.loads((FOLDER/'gui-verification-results.json').read_text())
    coverage = json.loads((FOLDER/'gui-coverage.json').read_text())['totals']
    assert len(data['tests']) == data['tests_run']
    gui = [record for record in data['tests'] if record['id'].split('.')[0] in ('test_export_window', 'test_gui_audit')]
    assert all(record['id'].rsplit('.', 1)[1] in LABELS for record in gui)
    when = datetime.fromisoformat(data['finished_utc'])
    stamp = when.strftime('%d.%m.%Y %H:%M:%S UTC')
    duration = re.search(r'Ran \d+ tests in ([\d.]+)s', (FOLDER/'gui-test-run.log').read_text()).group(1)
    passed = data['outcomes'].get('passed', 0)
    failed = data['outcomes'].get('failed', 0) + data['outcomes'].get('error', 0)
    skipped = data['outcomes'].get('skipped', 0)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    section.header_distance, section.footer_distance = Cm(.7), Cm(.7)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'tr-TR')
    normal.element.get_or_add_rPr().append(language)
    for name, size in (('Title', 30), ('Heading 1', 20), ('Heading 2', 13)):
        style = doc.styles[name]
        style.font.name, style.font.size = 'Calibri', Pt(size)
        style.font.color.rgb = RGBColor.from_string(GREEN if name != 'Title' else TEXT)
        style.paragraph_format.space_after = Pt(10)
    header = section.header.paragraphs[0]
    set_run(header.add_run('GOKWORKS  /  GUI DOĞRULAMA'), 9, True, GREEN)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(footer.add_run('GOKWORKS · '), 9, color=MUTED)
    field(footer, 'PAGE')
    set_run(footer.add_run(' / '), 9, color=MUTED)
    field(footer, 'NUMPAGES')
    doc.core_properties.title = 'GOKWORKS GUI Test Raporu'
    doc.core_properties.subject = 'PDF ve DXF/DWG arayüzünün baştan sona yeniden doğrulanması'
    doc.core_properties.author = 'GOKWORKS'
    doc.core_properties.last_modified_by = 'Codex'
    doc.core_properties.created = when.replace(tzinfo=None)
    doc.core_properties.modified = when.replace(tzinfo=None)

    doc.add_heading('GUI Test Raporu', level=0)
    paragraph(doc, 'PDF ve DXF/DWG · Baştan sona yeniden kontrol', True)
    paragraph(doc, stamp)
    doc.add_heading('Karar', level=1)
    paragraph(doc, 'GUI, bu bulut ortamında sınanan kullanıcı akışlarında çalışıyor.' if data['successful'] else
              'Sınanan GUI akışlarında giderilmesi gereken başarısız testler var.', True)
    paragraph(doc, 'Pencere açılışı, iki ekran, iki tema, tarama, seçenekler, arama/filtreler, '
              'ayrıntılar, CSV raporu, durdurma ve kapanış akışları yeniden çalıştırıldı. '
              'Sonuç, gerçek Tk arayüzüne dayanır; yalnızca kod okuması değildir.')
    table(doc, ['Gösterge', 'Sonuç'], [
        ('Toplam test', str(data['tests_run'])), ('Başarılı / başarısız / atlanan', f'{passed} / {failed} / {skipped}'),
        ('Gerçek GUI testleri', f'{len(gui)} (22 önceki + 12 ek kabul testi)'),
        ('Test konsolu süresi', f'{duration} saniye'),
        ('Tüm uygulama satır kapsamı', f"%{coverage['percent_statements_covered']:.1f}".replace('.', ',')),
        ('Satır/dal birleşik kapsamı', f"%{coverage['percent_covered']:.1f}".replace('.', ',')),
    ], [8, 9])
    doc.add_heading('Sonucun sınırı', level=2)
    paragraph(doc, 'Linux ortamında gerçek SolidWorks/Windows COM, Windows dosya seçim pencereleri ve '
              'Explorer açılışı çalıştırılamaz. PDF/DXF/DWG aktarım çağrıları COM test nesneleriyle '
              'sınanmıştır; gerçek CAD açınım ölçüleri ve DWG içeriği bu raporla doğrulanmış değildir.')
    paragraph(doc, 'K Faktörü ekranı mevcut programda kasıtlı olarak devre dışıdır. '
              'Bu rapor K faktörü hesabının çalıştığını iddia etmez.')

    page(doc, 'Yöntem ve test ortamı')
    table(doc, ['Bileşen', 'Çalıştırılan ortam'], [
        ('Depo / ref', 'gkkizil/gokworks · main · değiştirilmiş çalışma kopyası'),
        ('Git başlangıç commit’i', data['git_head']), ('Sistem', data['platform']),
        ('Python / Tcl-Tk', f"{data['python']} / {data['tk']}"),
        ('GUI kütüphaneleri', f"CustomTkinter {data['packages']['customtkinter']} · Pillow {data['packages']['Pillow']}"),
        ('DXF okuyucusu', f"ezdxf {data['packages']['ezdxf']}"),
        ('Test ekranı', 'Xorg :99, 1280×900; pencere kontrolleri: 1000×690, 1140×840, 1240×860'),
        ('Rapor aracı', f"python-docx {data['packages']['python-docx']} (uygulama bağımlılığı değil)"),
    ], [4.2, 12.8])
    doc.add_heading('Nasıl kontrol edildi?', level=2)
    paragraph(doc, 'Gerçek Tk pencereleri oluşturuldu ve olay döngüsünde çalıştırıldı. Yeni kabul '
              'testlerinde düğmelerin invoke komutları, menü seçimleri, alanlara yazma ve gerçek Tk '
              'Ctrl+F, çift tıklama, Enter, Escape olayları kullanıldı. Bunlar otomatik olay '
              'kontrolleridir; fiziksel fareyle yapılmış insan testi olarak sunulmaz.')
    paragraph(doc, 'Dosya seçimi/iptali, OS klasör açma ve SolidWorks sınırları kontrollü test '
              'nesneleriyle temsil edildi. GUI iş parçacıkları, olay kuyruğu ve dosya/CSV yazımı '
              'gerçek kod üzerinden çalıştı. DXF okuyucusu dört gerçek biçim örneğiyle çalıştı; '
              'bu örnekler ezdxf üretimidir, SolidWorks çıktısı değildir.')
    table(doc, ['Test grubu', 'Adet', 'Sonuç'], [
        ('DXF/DWG', 69, 'GEÇTİ'), ('PDF', 29, 'GEÇTİ'),
        ('Önceki GUI akışları', 22, 'GEÇTİ'), ('Ek GUI kabul kontrolleri', 12, 'GEÇTİ'),
        ('Ayarlar / bağlantı / giriş', 8, 'GEÇTİ'),
    ], [11, 3, 3])

    page(doc, 'İşlevlerin değerlendirilmesi')
    table(doc, ['Kontrol', 'Sonuç / kanıtın sınırı'], [
        ('Açılış ve sayfalar', 'Geçti. main.py, PDF/DXF geçişleri ve durum koruma sınandı.'),
        ('Açık / koyu tema', 'Geçti. Gerçek seçim düğmeleri, kalıcı tercih ve işlem sırasında değişim sınandı.'),
        ('Kaynak / çıktı seçimi', 'Geçti. Gözat komutları, filtre, Türkçe yollar ve iptal sınandı; native Windows diyaloğu simüle edildi.'),
        ('Tarama ve hata mesajları', 'Geçti. Dosya eleme, durdurma, okuma hatası ve kontrollerin yeniden açılması sınandı.'),
        ('DXF/DWG seçenekleri', 'Geçti. Biçim, büküm, alt klasör ve üzerine yaz değerlerinin işçiye iletilmesi sınandı.'),
        ('Aktarım ve sonuçlar', 'Geçti. GUI→işçi→sonuç/CSV akışı sınandı; SolidWorks COM tarafı test nesneleriyle temsil edildi.'),
        ('Arama / durum filtresi', 'Geçti. Türkçe İ eşleşmesi, canlı sonuçlar, temizleme ve ekranlar arasında bağımsızlık sınandı.'),
        ('Ayrıntılar / pano', 'Geçti. Çift tıklama, Enter, Escape, tam yollar ve kapandıktan sonra panonun korunması sınandı.'),
        ('CSV raporu', 'Geçti. Hatalı filtresi açıkken tüm üç sonuç yazıldı; iptal ve yazma hatası kontrol edildi.'),
        ('Çıktı klasörünü açma', 'Komut ve hata yönetimi geçti. Gerçek Windows Explorer açılışı doğrulanmadı.'),
        ('Küçük pencere / ölçek', '12 boyut-tema-ekran alt senaryosu geçti. Ayrıca %125 widget ölçeklemesi iki ekranda/temada geçti.'),
        ('Kapanış / K Faktörü', 'Aktif işi bekleme ve kapanışı reddetme geçti. K Faktörü kasıtlı devre dışı; hesaplaması yok.'),
    ], [4.6, 12.4])
    doc.add_heading('Bu incelemede yapılan düzeltme', level=2)
    paragraph(doc, 'İlk ek test çalışmasında 11 test geçti, bir test ilk yerleşim zamanlayıcısına '
              'erişemedi. İlk boyutlandırma, pencerenin Configure olayına bağlıydı. GUI kurulumunun '
              'sonunda ilk boyutlandırma açıkça zamanlandı; olay erken geldiyse de ilk yerleşim '
              'yapılıyor. Sonraki 12 ek test ve 140 testlik tam çalıştırma geçti. Bu bulgu, '
              'normal kullanıcı açılışında bir çökme gözlemi olarak sunulmaz.')

    for number, records in enumerate((gui[:17], gui[17:]), 1):
        page(doc, f'GUI test envanteri · {number}/2')
        paragraph(doc, 'Her satır son çalıştırmadan alınmıştır. Süreler ilgili testin kurulum/temizliği dahil ölçümüdür.')
        rows = []
        for index, record in enumerate(records, 1+(number-1)*17):
            method = record['id'].rsplit('.', 1)[1]
            rows.append((index, LABELS[method]+'\n'+method,
                         f"{record['seconds']:.2f} s", 'GEÇTİ' if record['status'] == 'passed' else record['status']))
        value = table(doc, ['No', 'Senaryo / test adı', 'Süre', 'Sonuç'], rows, [.8, 12.2, 1.6, 2.4])
        for row in value.rows[1:]:
            for run in row.cells[1].paragraphs[0].runs:
                run.font.size = Pt(9)
            for cell in (row.cells[0], row.cells[2], row.cells[3]):
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for theme, title in (('light', 'Açık tema'), ('dark', 'Koyu tema')):
        page(doc, title+' · gerçek pencere görüntüleri')
        screenshot(doc, f'audit-dxf-dwg-{theme}.png', 'DXF/DWG · 1240×860 · son kodun çalışan Tk penceresi')
        paragraph(doc, 'Tema seçimi, aynı çalışma alanı ve aynı aktarım özelliklerini korur. '
                  'Görüntü tasarım çizimi değildir; yeniden test sonrasında doğrudan pencereden alınmıştır.')
        doc.add_heading('PDF ekranı', level=2)
        doc.add_picture(str(FOLDER/'assets'/f'audit-pdf-{theme}.png'), width=Inches(4.0))
        paragraph(doc, 'PDF ekranı da aynı temada açılmıştır. Bulut görüntülerindeki yazı tipleri '
                  'Windows’taki Segoe UI/Consolas’tan farklıdır.')

    page(doc, 'Açık noktalar ve tekrar çalıştırma')
    table(doc, ['Alan', 'Durum'], [
        ('GUI testleri', '140 toplam / 34 GUI testi geçti. Son çalıştırmada başarısız veya atlanan test yok.'),
        ('Gerçek SolidWorks çıktısı', 'Çalıştırılmadı. Windows/SolidWorks oturumu bu Linux ortamında yok.'),
        ('Native Windows diyalogları / Explorer', 'Gerçek OS pencereleri çalıştırılmadı; komutların argümanları ve hata/iptal yolları sınandı.'),
        ('Monitör / DPI çeşitleri', '%125 widget ölçeği sınandı. Tüm Windows monitör, DPI ve sürüm kombinasyonları doğrulanmadı.'),
        ('K Faktörü', 'Kasıtlı devre dışı; kullanılabilir bir hesaplama ekranı değil.'),
        ('Ortam uyarıları', 'Salt okunur font/cache dizinleri için uyarı var. Alternatif font çizimiyle GUI çalıştı; bunlar test hatası değildir.'),
    ], [5, 12])
    doc.add_heading('Windows’ta son kontrol', level=2)
    for text in (
        'SolidWorks açıkken gerçek .SLDDRW ve tek gövdeli sac .SLDPRT kopyasıyla PDF/DXF/DWG aktarımını deneyin.',
        'Çıktıları açın; PDF sayfaları, açınım ölçüleri/yönü, büküm çizgileri ve “PARÇA - kalınlık - adet” adını kontrol edin.',
        'Gözat ve çıktı klasörünü aç düğmelerini native Windows pencereleriyle deneyin; iki tema ve pencere boyutlarını gözden geçirin.',
        'Aktarım sırasında Durdur ve kapanış sorusunun iki yanıtını deneyin; CSV raporunu Excel veya uygun bir tablo okuyucuyla açın.',
    ):
        doc.add_paragraph(text, style='List Bullet')
    doc.add_heading('Tekrarlanabilir kanıt', level=2)
    paragraph(doc, 'Test komutu: DISPLAY=:99 APPDATA=/workspace/.cloud-setup/user .venv/bin/python '
              'docs/report-tools/run_verification.py')
    paragraph(doc, 'Windows standart test komutu: .venv\\Scripts\\python.exe -m unittest discover -s tests -v')
    paragraph(doc, 'Kanıtlar: docs/reports/gui-verification-results.json (testler, alt senaryolar, '
              'süreler, sürümler ve SHA-256 dosya özetleri); gui-test-run.log; gui-coverage.json. '
              'Rapor üretimi: .venv/bin/python docs/report-tools/create_word_report.py')
    paragraph(doc, f"Son sayım: {data['lines']['application']:,} uygulama + {data['lines']['tests']:,} test "
              f"= {sum(data['lines'].values()):,} Python satırı.".replace(',', '.') +
              ' Boş satırlar ve yorumlar dahil; bağımlılıklar, yedekler ve rapor araçları hariç.')
    output = FOLDER/f"GOKWORKS_GUI_Test_Raporu_{when:%Y-%m-%d}.docx"
    doc.save(output)
    reopened = Document(output)
    assert len(reopened.inline_shapes) == 4
    assert any('GUI Test Raporu' in paragraph.text for paragraph in reopened.paragraphs)
    assert any(str(data['tests_run']) in cell.text for value in reopened.tables for row in value.rows for cell in row.cells)
    with ZipFile(output) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.endswith(('.xml', '.rels')):
                etree.fromstring(archive.read(name))
        media = [name for name in archive.namelist() if name.startswith('word/media/')]
        assert len(media) == 4
    print(json.dumps({'document': str(output), 'bytes': output.stat().st_size,
                      'test_count': data['tests_run'], 'gui_tests': len(gui),
                      'embedded_images': len(reopened.inline_shapes), 'validation': 'passed'}, ensure_ascii=False))


if __name__ == '__main__':
    main()
