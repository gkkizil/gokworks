"""Yeni K/R test kanıtları ve üç ekranın gerçek görüntüleriyle Word raporu."""

from datetime import datetime
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor

from create_word_report import paragraph, table, field

ROOT = Path(__file__).resolve().parents[2]
FOLDER = ROOT / 'docs/reports/k-radius'


def main():
    data = json.loads((FOLDER/'gui-verification-results.json').read_text())
    assert data['successful'] and data['outcomes'] == {'passed': data['tests_run']}
    for filename, digest in data['sha256'].items():
        assert hashlib.sha256((ROOT/filename).read_bytes()).hexdigest() == digest, filename
    coverage = json.loads((FOLDER/'coverage.json').read_text())['totals']
    screenshots = json.loads((FOLDER/'screenshots.json').read_text())
    for screenshot in screenshots:
        assert hashlib.sha256((ROOT/screenshot['file']).read_bytes()).hexdigest() == screenshot['sha256']
    when = datetime.fromisoformat(data['finished_utc']).astimezone(ZoneInfo('Europe/Istanbul'))
    gui = [r for r in data['tests'] if r['id'].split('.')[0] in ('test_export_window','test_gui_audit')
           or '.KGuiTests.' in r['id']]
    kr = [r for r in data['tests'] if r['id'].startswith('modules.k_factor.tests.test_k_factor.')]
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in (('Normal',10),('Title',26),('Heading 1',19),('Heading 2',13)):
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string('176B57' if name.startswith('Heading') else '243041')
    section.header.paragraphs[0].text = 'GOKWORKS · K/R algoritması ve güncel GUI denetimi'
    footer = section.footer.paragraphs[0]
    footer.add_run('GOKWORKS | ');field(footer,'PAGE');footer.add_run(' / ');field(footer,'NUMPAGES')
    doc.add_heading('K FAKTÖR - RADIUS\nGüncel GUI ve test raporu',0)
    paragraph(doc,when.strftime('%d.%m.%Y %H:%M:%S — Türkiye saati (UTC+3)'))
    paragraph(doc,f"{data['tests_run']} test geçti, 0 başarısız, 0 atlanan. {len(gui)} gerçek Tk GUI testi; K/R için {len(kr)} yeni test (GUI dahil).",True)
    paragraph(doc,'Hesaplama ve GUI Linux sanal ekranında gerçekten çalıştırıldı. SolidWorks COM test nesneleriyle temsil edildi; gerçek Windows/SolidWorks montajı veya üretim numunesi bu ortamda doğrulanmadı.')
    table(doc,('Özellik','Davranış'),[
        ('Malzeme','Genel çelik, paslanmaz, yumuşak alüminyum; ANSI K tablosu'),
        ('Takım','Bıçak uç R ve açı, alt kalıp V ve açı; standart ölçüler + özel giriş'),
        ('Otomatik V','Her parçanın gerçek kalınlığına göre katalogdan yaklaşık 8t'),
        ('BA ölçümü','Gerçek iç R + bağımsız BA veya A/B/düz boydan K kalibrasyonu'),
        ('Montaj','Aktif konfigürasyon, alt montajlar, dosya+konfigürasyonda tek işlem'),
        ('Yazma','K tipi=2, R metreye çevrilir; rebuild ve geri okuma'),
        ('Hata','Konfigürasyona yayılan/hatalı yazım geri alınır; geri alma hatası görünür'),
        ('Kaydetme','Varsayılan bellek değişikliği; açık seçimle Save3'),
    ],(3.6,13.4))
    paragraph(doc,f"Satır/dal birleşik kapsamı: %{coverage['percent_covered']:.1f}. Uygulama: {data['lines']['application']} fiziksel Python satırı; testler: {data['lines']['tests']} satır.")
    for title, stem in (('K FAKTÖR - RADIUS','gui-k-radius'),('PDF ekranı','gui-current-pdf'),('DXF / DWG ekranı','gui-current-dxf-dwg')):
        doc.add_page_break();doc.add_heading(title,1)
        for suffix, label in (('light','Açık tema'),('dark','Koyu tema')):
            doc.add_picture(str(ROOT/'docs/screenshots'/f'{stem}-{suffix}.png'),width=Inches(6.35))
            paragraph(doc,label+' · Gerçek 1240 × 860 Tk penceresi; üretim CAD verisi gösterilmez.')
    doc.add_page_break();doc.add_heading('Elle bend allowance ve K kontrolü',1)
    for text in (
        'BA = (θ × π / 180) × (R + K × t)',
        'K = (BA / (θ × π / 180) − R) / t',
        'BD = 2 × (R + t) × tan(θ / 2) − BA',
        'Açı, düz sacdan dönülen büküm açısıdır. İç açı 60° ise θ=120° kullanılır.',
        'Elle örnek: t=2 mm, R=2 mm, K=0.4, θ=90° → BA=4.39822971502571 mm; BD=3.60177028497429 mm.',
        'A=50 mm, B=40 mm, ölçülmüş düz boy=86.3982297150257 mm → BA=4.3982297150257 mm → K≈0.4.',
        'Bu sayısal örnek fiziksel ölçüm değildir. Üretim teyidi için BA, R ve t gerçek bağımsız numuneden alınmalıdır. Tablodan BA hesaplayıp ters formülle aynı K’ye dönmek üretim doğrulaması sayılmaz.',
        'FreeCAD tools/calc-unfold.py gerçekten çalıştırıldı: t=2, R=1.64, K=0.38, θ=90° → her iki programda BA=3.7699111843077517 mm.',
        'R tahmin katsayıları yaklaşık başlangıçlardır; üretici kalibrasyonuyla teyit edilmedi. Havada büküm modeli tonaj, çarpışma, minimum flanş veya geri esneme hesabı yapmaz.',
    ): paragraph(doc,text)
    doc.add_picture(str(ROOT/'docs/screenshots/gui-k-radius-ba-light.png'),width=Inches(6.35))
    doc.add_page_break();doc.add_heading('K/R test envanteri ve kanıtlar',1)
    for record in kr:
        paragraph(doc,'GEÇTİ · '+record['id'])
    paragraph(doc,'Kanıtlar: docs/reports/k-radius/gui-verification-results.json, test-run.log, coverage.json, api-contract.json, reference-comparison.json ve screenshots.json. Kullanım, kaynak commitleri ve kapsam sınırları: docs/k-factor-radius.md.')
    path=FOLDER/'GOKWORKS_K_RADIUS_GUI_Raporu_2026-10-09.docx'
    doc.save(path)
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        images=[n for n in archive.namelist() if n.startswith('word/media/')]
        assert len(images)==7
    text='\n'.join(p.text for p in Document(path).paragraphs)
    assert all(r['id'] in text for r in kr)
    (FOLDER/'word-report-validation.json').write_text(json.dumps(dict(
        file=path.name,bytes=path.stat().st_size,sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        embedded_images=len(images),tests=data['tests_run'],gui_tests=len(gui),kr_tests=len(kr),
        zip_crc_ok=True,python_docx_readback_ok=True,native_word_render_verified=False),indent=2)+'\n')
    print(path)


if __name__ == '__main__': main()
