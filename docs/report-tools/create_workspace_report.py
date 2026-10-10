"""Build the K/R, ISO and GUI report from verified hashes and test evidence."""

from collections import Counter
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
FOLDER = ROOT/'docs/reports/kr-workspace-fix'


def main():
    data = json.loads((FOLDER/'gui-verification-results.json').read_text())
    assert data['successful'] and data['outcomes'] == {'passed': data['tests_run']}
    for filename, digest in data['sha256'].items():
        assert hashlib.sha256((ROOT/filename).read_bytes()).hexdigest() == digest, filename
    screenshots = json.loads((FOLDER/'screenshots.json').read_text())
    for item in screenshots:
        assert hashlib.sha256((ROOT/item['file']).read_bytes()).hexdigest() == item['sha256']
    source = json.loads((FOLDER/'iso-source-comparison.json').read_text())
    assert source['unexpected_differences'] == 0
    assert source['table_sha256'] == data['sha256']['iso_fits/tables.json']
    coverage = json.loads((FOLDER/'coverage.json').read_text())['totals']
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(1.8)
    for name, size in [('Normal', 10), ('Title', 24), ('Heading 1', 16)]:
        font = doc.styles[name].font
        font.name, font.size = 'Calibri', Pt(size)
        font.color.rgb = RGBColor.from_string('177C65' if name == 'Heading 1' else '243041')
    section.header.paragraphs[0].text = 'GOKWORKS · K/R, ISO ve GUI doğrulaması'
    doc.add_heading('K/R, ISO geçmeleri ve\nGUI kontrol raporu', 0)
    doc.add_paragraph('10.10.2026 · Son çalışma: '+data['finished_utc'])
    doc.add_paragraph(f"{data['tests_run']} test geçti; başarısız/atlanan test yok. 93 gerçek Tk GUI testi.")
    doc.add_paragraph('Gerçek Linux/Tk arayüzü ve DXF okuyucusu çalıştırıldı. SolidWorks COM davranışı test nesneleriyle temsil edildi. Gerçek Windows/SolidWorks 2025 montajında K/R ataması ve CAD çıktıları bu ortamda çalıştırılmadı.')
    doc.add_heading('Düzeltmeler', 1)
    for text in [
        'K/R: çağrılabilir dönen COM nesnesi varsayılan üye olarak çağrılmaz. Yalnız aktif büküm payı alanı okunur/yazılır; rebuild, geri okuma ve geri alma korunur. Hata aşama ve üye adı içerir.',
        'Manuel örnek: hedef sac 2 mm, K=0,274, R=2 mm. SolidWorks radius birimi metre: 0,002. Diğer kalınlıklar atlanır. Kaydetme ayrı seçilir.',
        'PDF/DXF/K: solda bağımsız kaydırılan form, sağda tam yükseklikte sonuç alanı. K/R tarama/hesap/kayıt seçimi sabit kalır. ISO/GD&T kendi çalışma alanlarını korur.',
        'ISO: 23 delik/mil esaslı örnek, Tümü/Delik esaslı/Mil esaslı filtresi. Tümü görünümünde iki esas dönüşümlü gösterilir. 198 çift menülerden seçilebilir.',
        'Manuel ISO: H7 ağırlıklı ayrı hazır örnek önerisi kaldırıldı; tam eşleşme, yakın referans ve eşit uzaklıktaki sınıflar ayrılır. Giriş limitleri değişmez. Kısa filtreye geçişte boş kaydırma alanı sorunu giderildi.',
    ]:
        doc.add_paragraph(text, style='List Bullet')
    doc.add_heading('ISO bağımsız kontrol örnekleri', 1)
    table = doc.add_table(rows=1, cols=3); table.style = 'Light Shading Accent 1'
    for cell, title in zip(table.rows[0].cells, ['Ø50 mm', 'Delik − mil aralığı (mm)', 'Sonuç']):
        cell.text = title
    for row in [('JS7/h6', '−0,0125 … +0,0285', 'Ara'), ('K7/h6', '−0,018 … +0,023', 'Ara'),
                ('N7/h6', '−0,033 … +0,008', 'Ara'), ('P7/h6', '−0,042 … −0,001', 'Sıkı'),
                ('Manuel delik 0/+0,02; mil −0,02/0', '0 … +0,04', 'Boşluklu; tam ISO değil'),
                ('Manuel delik −0,05/0; mil ±0,1', '−0,15 … +0,10', 'Ara; eşit delik referansları N7/P7')]:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    doc.add_paragraph(f"Sabit GitHub kaynaklarıyla {source['comparisons_total']} sayısal kontrol: 0 beklenmeyen fark; bilinen F8 ≤3 mm kaynak satırından iki fark. Ø50 için 11 delik sınıfı bağımsız sayılarla GUI/raporda sınandı. Resmî ISO sertifikası iddia edilmez.")
    doc.add_paragraph('H7/p6 Ø3’te ara, Ø50’de sıkı; P7/h6 Ø3’te −0,016…0 mm sıkılık/sınırda temastır. Gruplama seçilen çaptaki sayısal limitlerden yapılır.')
    doc.add_heading('Test kanıtı ve kapsam sınırı', 1)
    modules = Counter(item['id'].split('.')[0] for item in data['tests'])
    for name, count in sorted(modules.items()):
        doc.add_paragraph(f'{name}: {count} başarılı test')
    doc.add_paragraph(f"Uygulama {data['lines']['application']} Python satırı; testler {data['lines']['tests']} satırı. Satır/dal birleşik kapsamı %{coverage['percent_covered']:.1f}; bu oran gerçek CAD doğruluğu ölçüsü değildir.")
    doc.add_paragraph('1000×690 ve 1366×768 (%100), 1400×900 (%125), 1920×1080 (%150), iki tema, sayfa/alt sekme geçişleri ve kaydırma sınandı. Gerçek Windows monitör/DPI davranışı ayrıca kontrol edilmelidir.')
    for title, name in [('K/R manuel · koyu', 'k-factor-manual-dark'), ('K/R manuel · açık', 'k-factor-manual-light'),
                        ('ISO mil esaslı · koyu', 'iso-fits-shaft-basis-dark'), ('ISO mil esaslı · açık', 'iso-fits-shaft-basis-light'),
                        ('ISO manuel karşılaştırma', 'iso-fits-manual-comparison-dark'), ('DXF/DWG', 'dxf-dwg-dark'),
                        ('PDF', 'pdf-light'), ('GD&T', 'geometric-tolerances-dark')]:
        doc.add_page_break(); doc.add_heading(title, 1)
        doc.add_picture(str(FOLDER/'screenshots'/f'{name}.png'), width=Cm(17))
        doc.add_paragraph('Gerçek 1240×860 Tk penceresi; SolidWorks verisi veya gerçek CAD başarı sonucu gösterilmez.')
    doc.add_page_break(); doc.add_heading('Yerel SolidWorks 2025 kabul kontrolü', 1)
    doc.add_paragraph('GitHub dalı codex/dxf-dwg-gui-report. Git ile açılan projede doğru dalda Pull; ZIP ile açılan projede güncel ZIP ayrı klasöre açılmalı. Aynı PyCharm yorumlayıcısına requirements.txt kurulur; main.py çalıştırılır.')
    doc.add_paragraph('Aktif montajda 2 mm / K=0,274 / R=2 mm ile tara → uygula. İlk kontrol dosya kaydı kapalıyken yapılır. SheetMetal unsurunun doğru konfigürasyonunda K, R, kalınlık ve büküm payı türü kontrol edilir. Başarısızsa yeni aşama/API üyesini içeren tam günlük/CSV gerekir.')
    doc.add_paragraph('Kanıtlar: docs/reports/kr-workspace-fix/README.md, test-run.log, gui-verification-results.json, coverage.json, dispatch-reproduction.json, iso-matching-reproduction.json, iso-source-comparison.json, screenshots.json. Yerel adımlar: docs/local-solidworks-check.md.')
    path = FOLDER/'GOKWORKS_KR_ISO_GUI_Kontrol_2026-10-10.docx'
    doc.save(path)
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        count = sum(name.startswith('word/media/') for name in archive.namelist())
    readback = Document(path)
    assert str(data['tests_run']) in '\n'.join(p.text for p in readback.paragraphs)
    assert count == 8
    validation = dict(file=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(), bytes=path.stat().st_size,
                      tests=data['tests_run'], embedded_images=count, zip_crc_ok=True,
                      python_docx_readback_ok=True, native_word_render_verified=False)
    (FOLDER/'word-report-validation.json').write_text(json.dumps(validation, indent=2)+'\n')
    print(path)


if __name__ == '__main__':
    main()
