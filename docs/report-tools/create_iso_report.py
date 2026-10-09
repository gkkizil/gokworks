"""ISO eklemesinin doğrulanmış kanıtlarından resimli Word raporu üret."""

from datetime import datetime
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from create_word_report import paragraph, table, field

ROOT = Path(__file__).resolve().parents[2]
FOLDER = ROOT/'docs/reports/iso-fits'


def main():
    results = json.loads((FOLDER/'gui-verification-results.json').read_text())
    sources = json.loads((FOLDER/'source-comparison.json').read_text())
    screenshots = json.loads((FOLDER/'screenshots.json').read_text())
    coverage = json.loads((FOLDER/'coverage.json').read_text())['totals']
    assert results['successful'] and results['outcomes'] == {'passed': results['tests_run']}
    for file, digest in results['sha256'].items():
        assert hashlib.sha256((ROOT/file).read_bytes()).hexdigest() == digest, file
    for item in screenshots:
        assert hashlib.sha256((ROOT/item['file']).read_bytes()).hexdigest() == item['sha256']
    assert sources['unexpected_differences'] == 0
    assert sources['table_sha256'] == results['sha256']['iso_fits/tables.json']
    gui = [r for r in results['tests'] if r['id'].startswith(('test_export_window.', 'test_gui_audit.', 'test_final_gui.'))
           or any(f'.{name}.' in r['id'] for name in ('KGuiTests', 'ManualGuiTests', 'IsoGuiTests'))]
    when = datetime.fromisoformat(results['finished_utc']).astimezone(ZoneInfo('Europe/Istanbul'))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in [('Normal', 10), ('Title', 25), ('Heading 1', 19)]:
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string('176B57' if name.startswith('Heading') else '243041')
    section.header.paragraphs[0].text = 'GOKWORKS · ISO geçme toleransları'
    footer = section.footer.paragraphs[0]
    footer.add_run('GOKWORKS | '); field(footer, 'PAGE'); footer.add_run(' / '); field(footer, 'NUMPAGES')
    doc.add_heading('ISO geçme toleransları\nve program kontrol raporu', 0)
    paragraph(doc, when.strftime('%d.%m.%Y %H:%M:%S — Türkiye saati (UTC+3)'))
    paragraph(doc, f"{results['tests_run']} test geçti; 0 başarısız, 0 atlanan. {len(gui)} gerçek Tk GUI testi.", True)
    paragraph(doc, 'ISO GEÇMELERİ: nominal çap, 11 delik / 18 mil sınıfı, boşluk/sıkılık limitleri, şematik kesit, gerçek sapma bölgeleri, üç örnek ölçü, mm/µm, arama, CSV ve pano. Açık/koyu tema ile seçimlerin kalıcılığı sınandı.')
    paragraph(doc, 'Çap kapsamı 1–400 mm. ISO 286-1:2010 ve ISO 286-2:2010 temel alınan standart adlarıdır. Resmî ISO/üretici sitelerine bulut ağ politikası nedeniyle erişilemedi. Tam standart metni indirilmedi veya incelenmiş sayılmadı; güncel baskı/düzeltmeler ISO kataloğundan ayrıca kontrol edilmelidir. Açık kaynak karşılaştırması ISO sertifikası değildir.')
    table(doc, ('Ø50 mm', 'Cmin / Cmax (mm)', 'Tür'), [
        ('H8/f7', '+0,025 / +0,089', 'Boşluklu'), ('H7/h6', '0 / +0,041', 'Boşluklu; sınırda temas'),
        ('H7/k6', '−0,018 / +0,023', 'Ara'), ('H7/p6', '−0,042 / −0,001', 'Sıkı; 0,001–0,042 mm')], (3, 7, 7))
    paragraph(doc, 'Elle kontrol: Dmin=D+EI; Dmax=D+ES; dmin=D+ei; dmax=D+es. Cmin=EI−es; Cmax=ES−ei. H8/f7 için EI=0, ES=39, es=−25, ei=−50 µm: Cmin=25, Cmax=89 µm. Tasarım örneğindeki 0,074–0,161 mm sayıları kullanılmadı.')
    doc.add_heading('Kaynakların karşılaştırılması', 1)
    paragraph(doc, f"{sources['comparisons_total']} sayısal kontrol; {len(sources['known_source_errors'])} bilinen hatalı kaynak farkı, 0 beklenmeyen fark.")
    for name, count in sources['comparisons'].items():
        paragraph(doc, f'{name}: {count} kontrol.')
    paragraph(doc, 'isofits isotol Python fonksiyonu D>3 mm için gerçekten yürütüldü. Passungsrechner JSON tablosu karşılaştırıldı; JavaScript GUI çalıştırılmadı. Rust projesinin sayısal tabloları incelendi; Rust uygulaması çalıştırılmadı. ISO286 kaynağının IT genişlikleri ve temel sapmaları bağımsız kontrol edildi.')
    paragraph(doc, 'F8 ≤3 mm: Passungsrechner EI=+10, ES=+20 µm veriyor; IT8 genişliği 14 µm ve F temel sapması +6 µm olmalı. Uygulamada EI=+6, ES=+20 µm kullanılır. 1,5 ve 3 mm karşılaştırmalarındaki iki kaynak farkı kanıtta korunur. Kaynak commitleri ve dosya hashleri source-comparison.json içindedir.')
    doc.add_heading('Test kapsamı ve sınırlar', 1)
    paragraph(doc, 'ISO için 8 hesap testi ve 9 GUI testi eklendi. Bütün 198 delik/mil kombinasyonunun beş çaptaki örnek ölçüleri limitler içinde kontrol edildi. 3/50 mm sınırları, virgüllü giriş, yanlış sınıf, kapsam dışı çap, yarım mikrometre, sıfır boşluk, geçersiz girişte eski sonuçların temizlenmesi, CSV iptal/kayıt hatası ve tercihler sınandı.')
    paragraph(doc, 'ISO GUI matrisi iki tema × beş boyut/ölçek: 1000×690, 1240×860, 1366×768 (%100), 1400×900 (%125), 1920×1080 (%150). Gerçek pencere boyutu ve kaydırılarak sonuç düğmesine erişim doğrulandı. Önceki PDF/DXF/K-R testleri de tekrar geçti. Bu Linux çalışması gerçek Windows DPI, SolidWorks COM üretimi veya yerel Word sayfa görünümü değildir.')
    paragraph(doc, f"Satır kapsamı %{coverage['percent_statements_covered']:.2f}; satır/dal birleşik kapsamı %{coverage['percent_covered']:.2f}. Uygulama {results['lines']['application']} fiziksel Python satırı; testler {results['lines']['tests']} satır. Veri JSON'u ve rapor araçları bu Python sayımında değildir.")
    paragraph(doc, 'Kesit şematiktir; sapmalar büyütülür. Bölge grafiği doğrusal sapma ölçeğidir. Orta ölçü üretim dağılımı varsayımı değildir. Montaj kuvveti, yağlama, malzeme, sıcaklık ve geometrik toleranslar ayrıca değerlendirilir. CSV ölçü/tolerans raporudur; CAD modeli değiştirilmez.')
    for item in screenshots:
        doc.add_page_break()
        doc.add_heading(f"{item['hole']}/{item['shaft']} · {item['theme']} · {item['view']}", 1)
        doc.add_picture(str(ROOT/item['file']), width=Inches(6.35))
        paragraph(doc, 'Gerçek 1240×860 Tk penceresi. Ø50 mm sayısal hesap; SolidWorks bağlantısı gerekmez.')
    doc.add_heading('Kanıtlar', 1)
    paragraph(doc, 'docs/iso-fits.md; gui-verification-results.json; test-run.log; coverage.json; source-comparison.json; screenshots.json. Önceki CAD referans incelemesi docs/reports/final-audit/README.md içinde tarihsel kayıt olarak korunur.')
    path = FOLDER/'GOKWORKS_ISO_Gecme_Raporu_2026-10-09.docx'
    doc.save(path)
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        embedded = [n for n in archive.namelist() if n.startswith('word/media/')]
        assert len(embedded) == len(screenshots)
    readback = '\n'.join(p.text for p in Document(path).paragraphs)
    assert all(term in readback for term in (f"{results['tests_run']} test", f'{len(gui)} gerçek', 'ISO 286-2:2010', '3217'))
    (FOLDER/'word-report-validation.json').write_text(json.dumps({
        'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'bytes': path.stat().st_size,
        'tests': results['tests_run'], 'gui_tests': len(gui), 'embedded_images': len(embedded),
        'zip_crc_ok': True, 'python_docx_readback_ok': True, 'native_word_render_verified': False}, indent=2)+'\n')
    print(path)


if __name__ == '__main__':
    main()
