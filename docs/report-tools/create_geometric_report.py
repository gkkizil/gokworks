"""Geometrik tolerans eklemesini test ve ekran kanıtlarından Word'e aktar."""

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from create_word_report import paragraph, table, field

ROOT = Path(__file__).resolve().parents[2]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, default=ROOT/'docs/reports/geometric-tolerances')
    parser.add_argument('--dimension-audit', action='store_true')
    args = parser.parse_args()
    folder = args.output_dir.resolve()
    results = json.loads((folder/'gui-verification-results.json').read_text())
    sources = json.loads((folder/'source-comparison.json').read_text())
    screens = json.loads((folder/'screenshots.json').read_text())
    coverage = json.loads((folder/'coverage.json').read_text())['totals']
    assert results['successful'] and results['outcomes'] == {'passed': results['tests_run']}
    for file, digest in results['sha256'].items():
        assert hashlib.sha256((ROOT/file).read_bytes()).hexdigest() == digest, file
    for item in screens:
        assert hashlib.sha256((ROOT/item['file']).read_bytes()).hexdigest() == item['sha256']
    assert sources['unexpected_differences'] == 0 and sources['comparisons_total'] == 14
    primary = sources['primary_standard']
    assert sources['active_characteristics'] == 12 and len(primary['section_anchor_checks']) >= 8
    if args.dimension_audit:
        drawing = json.loads((folder/'drawing-before-after.json').read_text())
        assert drawing['current_diagram_sha256'] == results['sha256']['modules/geometric_tolerances/diagram.py']
        assert drawing['native_tk_measurement'] and drawing['after']['long_value_fits_with_5px_padding']
    gui = [r for r in results['tests'] if r['id'].startswith(('test_export_window.', 'test_gui_audit.', 'test_final_gui.'))
           or any(f'.{name}.' in r['id'] for name in ('KGuiTests', 'ManualGuiTests', 'IsoGuiTests', 'ManualIsoGuiTests', 'GeometricGuiTests'))]
    when = datetime.fromisoformat(results['finished_utc']).astimezone(ZoneInfo('Europe/Istanbul'))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in [('Normal', 10), ('Title', 25), ('Heading 1', 19)]:
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string('176B57' if name.startswith('Heading') else '243041')
    section.header.paragraphs[0].text = 'GOKWORKS · Geometrik toleranslar'
    footer = section.footer.paragraphs[0]
    footer.add_run('GOKWORKS | '); field(footer, 'PAGE'); footer.add_run(' / '); field(footer, 'NUMPAGES')
    doc.add_heading('GD&T ölçü ve şema denetimi' if args.dimension_audit else
                    'Geometrik toleranslar\nve program kontrol raporu', 0)
    paragraph(doc, when.strftime('%d.%m.%Y %H:%M:%S — Türkiye saati (UTC+3)'))
    paragraph(doc, f"{results['tests_run']} test geçti; 0 başarısız, 0 atlanan. {len(gui)} gerçek Tk GUI testi.", True)
    if args.dimension_audit:
        doc.add_heading('Bulunan ve düzeltilen sorunlar', 1)
        table(doc, ('Önce', 'Sonuç ve kontrol'), [
            ('45° yazan açısallık sınırı yaklaşık 29,39° çizilmişti.',
             'İki sınır gerçek 45°; t oku her iki sınıra ulaşır ve sınırlara diktir. Tk koordinatlarından açı ve skaler çarpım kontrol edildi.'),
            ('Konum çap oku 140 birimlik bölgede 120 birim uzunluktaydı.',
             'Çap oku çemberin iki sınırına ulaşır; uzunluğu çizilen çapa eşittir.'),
            ('Profil sınırları eğrinin düşey kopyalarıydı; normal yönde uniform değildi.',
             'Teorik dairesel profil ve eş merkezli normal ofset sınırları çizildi. Her örnek noktada genişlik sabit, teorik profil iki sınırın ortasında.'),
            ('Çerçeve hücresi sabit genişlikteydi; uzun sayılar sığmayabiliyordu.',
             'Gerçek Canvas yazı sınırları ölçülüyor. En uzun 999.999999 ve Ø, A/B/C iki tema ve %100/%125/%150 ölçeklerde sığıyor.'),
            ('ASME teknik çerçevesinde yerelleştirilmiş ondalık virgül kullanılıyordu.',
             '4.3.1/4.3.3: teknik çerçevede ondalık nokta, 1 mm altında başta sıfır; Türkçe giriş/açıklamalar virgülü destekler.'),
            ('Ölçüm alanında her özellik için toplam hatanın anlamı ayrı gösterilmiyordu.',
             '12 güncel özellik için ölçüm tanımı ekrana ve CSV/panoya eklendi; konum çapı, merkezlenmiş profil 2 x maksimum mutlak normal sapma, radyal şekil ve TIR ayrıldı.')], (6, 11))
        doc.add_heading('Bağımsız elle hesaplanan örnekler', 1)
        table(doc, ('Giriş', 'Hesap ve sonuç'), [
            ('Konum dx=0,015; dy=0,020; t=0,05 mm', '2 x sqrt(0,015² + 0,020²) = 0,05 mm; sınırda.'),
            ('Konum dx=0,018; dy=0,024; t=0,05 mm', 'Radyal 0,03; çap cinsinden hata 0,06 mm; sınır dışında.'),
            ('Merkezlenmiş profil sapmaları +0,01 / +0,04; t=0,05 mm',
             'Her nokta ±0,025 içinde olmalı. 2 x max(|sapma|)=0,08 mm; sınır dışında. Max-min 0,03 kullanmak hatalıdır.'),
            ('Merkezlenmiş profil sapmaları -0,025 / +0,025; t=0,05 mm', 'Eşdeğer toplam hata 0,05; sınırda.'),
            ('Değerlendirilmiş sınır yarıçapları 24,975 / 25,025 mm', 'Radyal şekil bölgesi 0,05 mm; çap farkıyla veya ikiyle çarparak karıştırılmaz.'),
            ('Değerlendirilmiş toplam hata 0,050001; t=0,05 mm', 'Kalan pay -0,000001 mm; sınır dışında. Decimal karşılaştırması yuvarlamaz.')], (7, 10))
        paragraph(doc, 'Bu örnekler testlerde bağımsız hesaplanıp arayüze değerlendirilmiş hata olarak verildi. Arayüz ham dx/dy veya nokta bulutu hesaplayıcısı değildir. Datumlu profil hizalaması datum sistemini, datumsuz profil izin verilen ilişkilendirmeyi korumalıdır. 2B konum örneği bütün AME eksenini/eksen eğimini doğrulamaz.')
    paragraph(doc, 'GEOMETRİK TOL. beşinci çalışma alanıdır: 12 güncel ASME özelliği ve iki kaldırılmış eğitim kartı, Türkçe/İngilizce ad ve arama, vektör sembol, örnek tolerans çerçevesi, bölge şeması, datum kullanımı, ölçüm yöntemi, üretim notu, açık tolerans girişi, isteğe bağlı toplam hata karşılaştırması, CSV/pano ve standart okuma rehberi.')
    doc.add_heading('Sayısal anlam ve belirtim kapsamı', 1)
    paragraph(doc, 'ASME Y14.5-2018: Rule #1 (5.8.1), Rule #2 RFS/RMB (5.9.1), datum önceliği (Bölüm 7), basic ölçüler, yüzey şekli (8.4), yön (9.3), RFS konum / unrelated AME ekseni (10.3.2), uniform profil (11.3) ve radyal salgı bölgeleri (12.4) yüklenen belgeden incelendi. Genel f/m/c/v sınıfından geometrik değer türetilmez.')
    table(doc, ('Örnek', 'Doğru okuma'), [
        ('Silindiriklik t=0,3 mm', 'İki eş eksenli silindir arasındaki toplam radyal mesafe 0,3 mm; ±0,3 değildir. Datum yoktur.'),
        ('Konum Ø0.3 | A | B | C', 'Silindirik bölge çapı 0,3 mm; radyal pay 0,15 mm. Datum önceliği A, B, C. Teorik konum ayrıca tanımlanır.'),
        ('Hata 0,050000; t=0,05 mm', 'Aritmetik sınırda: hata ≤ t. Tam resim/model veya ölçüm belirsizliğiyle kalite kabulü değildir.'),
        ('Hata 0,050001; t=0,05 mm', 'Sınır dışında; kalan pay -0,000001 mm. Yuvarlayarak uygun sayılmaz.')], (5, 12))
    paragraph(doc, 'Gösterim boyutu sadece örnek bağlamıdır; tolerans tasarımcı tarafından girilir. Girişler mm, virgül/nokta ve en çok 6 ondalık basamak kullanır. Boyut 0 < D ≤ 10000, tolerans 0 < t ≤ 1000, hata 0 ≤ e ≤ 1000 mm; bunlar arayüz sınırlarıdır, ASME uygulama aralığı değildir. Sıfır MMC konum belirtimi, bonus ve özel bölge değiştiricileri bu örneklere dahil değildir.')
    paragraph(doc, 'Dört şekil toleransında datum hücresi üretilmez. Profilde datum belirtime bağlıdır. Bu RFS delik konumu örneği datum ister; özel datumsuz koaksiyellik (10.6.2.3) üretilmez. Yön ve salgı örnekleri referans ister. Eş merkezlilik/simetri 2018’de kaldırılmıştır; sayısal giriş/çerçeve/bölge/çıktı engellenir (Önsöz ve Ek D). A / A|B / A|B|C temsili datum seçenekleridir; bütün olası resim belirtimleri için çerçeve editörü değildir. Gerçek çizimde özellik türü, teorik ölçüler ve datum unsurları ayrıca tanımlanır.')
    paragraph(doc, 'İsteğe bağlı hata, ilgili geometrik özellik için ölçüm sisteminin değerlendirdiği toplam hatadır. Tek çap, tek nokta veya ± boyut sapmasıyla geometrik uygunluk bulunmaz. Özellik değişince önceki hata temizlenir; uygulama yeniden açılınca ölçüm geri yüklenmez. Geçersiz giriş eski sonucu, şemayı, CSV/pano işlemini devre dışı bırakır.')
    doc.add_heading('Kaynak karşılaştırması', 1)
    paragraph(doc, f"FreeCAD-GDT: {sources['commit']}. FreeCAD eklenti kataloğundaki kaynak üzerinden 14 özellik adı eşleştirildi; açık tolerans değeri, ayrı datum sistemi ve 14 sembol SVG kaynağı kayıt altına alındı. Silindiriklik ve toplam salgı işaretleri yapısal olarak incelendi. Kaynak hashleri source-comparison.json içindedir. Kod kopyalanmadı. FreeCAD çalıştırılmadı; bu karşılaştırma metrologik/normatif doğrulama değildir.")
    paragraph(doc, f"Birincil belge: {primary['filename']}, {primary['pdf_pages']} PDF sayfası; SHA-256 {primary['sha256']}. İlgili bölümler metinden incelendi; {len(primary['section_anchor_checks'])} bölüm çıpası PDF sayfa indeksleriyle kayıtlıdır. Belgenin tamamı veya figürleri GitHub’a yüklenmedi.")
    paragraph(doc, 'Dosya adında R2024 var; kapak ve ana metin 2018 baskısıdır. İçerikte R2024 yeniden onay bilgisi teyit edilmedi. Farklı PDF metadata başlığı baskı kanıtı kabul edilmedi. İnceleme ilgili bölümler içindir; tüm belge/yeniden onay kaydı veya tam standart uygunluğu sertifikalandırılmaz.')
    paragraph(doc, 'ASME 8.4 normal, ilgili boyutla ilişkili yüzey formu: isteğe bağlı toplam boyut toleransı girildiğinde t < boyut toleransı aranır. t=0,05 ve boyut toleransı 0,05 reddedilir; 0,050001 ilişkiyi sağlar. Bu kontrol MMC kusursuz şekil zarfını ölçmez. Özel free state/independency/medyan uygulamaları üretilmez. Tür değişince boyut toleransı temizlenir.')
    paragraph(doc, 'Rule #2: işaretsiz boyut özelliği geometrik toleransında RFS, datum boyut özelliği referansında RMB varsayılır; M/L, bonus, datum mobility ve virtual condition hesabı yoktur. Yüzey formu/profili/yönü veya salgıya uygunsuz MMC/LMC eklenmez. Kaldırılan semboller eski GitHub kaynağında bulunduğu için kaynak doğrudan normatif kabul edilmedi; ASME belgesi esas alındı.')
    doc.add_heading('Test ve arayüz kanıtları', 1)
    math_count = sum('.GeometricRulesTests.' in r['id'] for r in results['tests'])
    gui_count = sum('.GeometricGuiTests.' in r['id'] for r in results['tests'])
    paragraph(doc, f'GD&T için {math_count} hesap/içerik ve {gui_count} gerçek GUI testi: kaldırılmış sembollerin engellenmesi, normal yüzeyde t < boyut toleransı, RFS/RMB ve AME ayrımı, toplam genişlik/çap ayrımı, datum, 1 nm eşik, geçersiz değerler, 12 güncel türün iki temada çizimi, datum menüleri, arama, CSV/pano/iptal/dosya hatası, sayfa önbelleği, tema, tercihler, rehber, ölçüm temizliği ve sayısal şema kontrolleri.')
    paragraph(doc, 'Yeni sayfa matrisi: açık/koyu × 1000×690, 1240×860, 1366×768 (%100), 1400×900 (%125), 1920×1080 (%150). Gerçek pencere boyutları doğrulandı. İlgili boyut toleransı dahil bütün girişler ve kopyalama düğmesi kaydırılarak erişilir; listedeki son özellik de erişilebilir. Dairesellik şemasında çemberin oranı korunur. Önceki PDF/DXF/DWG/K-R/ISO testleri de tekrar yürütüldü.')
    paragraph(doc, f"Uygulama {results['lines']['application']} fiziksel Python satırı; testler {results['lines']['tests']} satır. Satır kapsamı %{coverage['percent_statements_covered']:.2f}; satır/dal birleşik %{coverage['percent_covered']:.2f}. JSON verisi, rapor araçları ve bağımlılıklar bu satır sayımına dahil değildir.")
    paragraph(doc, 'Sınırlar: Tk Linux sanal ekranda çalıştırıldı. Gerçek Windows DPI, SolidWorks COM üretimi ve Microsoft Word sayfa görünümü doğrulanmadı. Şemalar öğreticidir ve ölçekli değildir; CMM simülasyonu yapılmaz, geometrik tolerans CAD modeline yazılmaz. Hata ≤ t karşılaştırması tam ASME resim/model uygunluğu veya belirsizlikli kalite kabulü değildir.')
    for item in screens:
        doc.add_page_break()
        doc.add_heading(f"{item['name']} · {item['theme']}", 1)
        doc.add_picture(str(ROOT/item['file']), width=Inches(6.35))
        paragraph(doc, f"Gerçek Tk ekranı {item['size'][0]}×{item['size'][1]}; {item['view']}. t={item['tolerance_mm']} mm. Görüntü ve kaynak hashleri kaydedildi.")
    doc.add_heading('Kanıt dosyaları', 1)
    paragraph(doc, 'docs/geometric-tolerances.md; gui-verification-results.json; test-run.log; coverage.json; source-comparison.json; screenshots.json; ölçü denetiminde drawing-before-after.json. Rapor üretimi bütün uygulama/test ve görüntü hashlerini yeniden doğrular. Önceki raporlar kendi sürümlerinin tarihsel kayıtlarıdır.')
    path = folder/('GOKWORKS_GDT_Olcu_Denetimi_2026-10-09.docx' if args.dimension_audit else
                   'GOKWORKS_Geometrik_Tolerans_Raporu_2026-10-09.docx')
    doc.save(path)
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        embedded = [name for name in archive.namelist() if name.startswith('word/media/')]
        assert len(embedded) == len(screens)
    readback = '\n'.join(p.text for p in Document(path).paragraphs)
    assert all(term in readback for term in (f"{results['tests_run']} test", f'{len(gui)} gerçek', 'ASME Y14.5-2018', 'Rule #1'))
    (folder/'word-report-validation.json').write_text(json.dumps({
        'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'bytes': path.stat().st_size,
        'tests': results['tests_run'], 'gui_tests': len(gui), 'embedded_images': len(embedded),
        'zip_crc_ok': True, 'python_docx_readback_ok': True, 'native_word_render_verified': False}, indent=2)+'\n')
    print(path)


if __name__ == '__main__':
    main()
