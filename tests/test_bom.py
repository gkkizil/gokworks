"""BOM occurrence counts, hierarchy, Excel content and failure atomicity."""

from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
import queue
import tempfile
import threading
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from bom import exporter as bom
from support.solidworks import FakeAssembly, FakeApp, FakeComponent, install_com


class Component(FakeComponent):
    def __init__(self, path, *args, name=None, excluded=False, mass_kg=1.0, **kwargs):
        super().__init__(path, *args, **kwargs)
        self.Name2 = name or Path(path).stem+'-1'
        self.ExcludeFromBOM = excluded
        self.mass_kg = mass_kg


class MassProperties:
    """A live assembly-context calculator; never reads the part's active doc."""
    def __init__(self):
        self.UseSystemUnits = False
        self.IncludeHiddenBodiesOrComponents = False
        self.SelectedItems = ()
        self.calls = []
        self.on_recalculate = None

    def Recalculate(self):
        assert self.UseSystemUnits is True
        assert self.IncludeHiddenBodiesOrComponents is True
        assert len(self.SelectedItems) == 1
        component = self.SelectedItems[0]
        self.calls.append(component)
        if self.on_recalculate:
            self.on_recalculate(component)
        return not getattr(component, 'mass_failed', False)

    @property
    def Mass(self):
        component = self.SelectedItems[0]
        masses = component.mass_kg
        return masses[component.ReferencedConfiguration] if isinstance(masses, dict) else masses


def attach_mass_properties(assembly):
    properties = MassProperties()
    assembly.Extension = SimpleNamespace(CreateMassProperty2=lambda: properties)
    return properties


class BomTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        plate = Component(r'C:\İmalat\Parça 2.SLDPRT', name='Parça 2-1')
        bolt = Component(r'C:\İmalat\Parça 10.SLDPRT', name='Parça 10-1')
        branch = Component(r'C:\İmalat\ÖRNEK-M001.SLDASM', children=[bolt, plate])
        hidden = Component(r'C:\İmalat\Parça 2.SLDPRT', name='Parça 2-2')
        hidden.IsHidden = True
        self.assembly = FakeAssembly(r'C:\İmalat\Ana.SLDASM', [branch, branch, hidden])
        self.app = FakeApp(self.assembly)
        self.mass_properties = attach_mass_properties(self.assembly)

    def test_mass_uses_referenced_occurrence_in_kg_and_preserves_document_state(self):
        thin = Component('/plate.sldprt', 'İnce', mass_kg={'İnce': .25, 'Kalın': 1.5})
        thick = Component('/plate.sldprt', 'Kalın', mass_kg={'İnce': .25, 'Kalın': 1.5})
        thin.model = SimpleNamespace(ConfigurationManager=SimpleNamespace(
            ActiveConfiguration=SimpleNamespace(Name='Kalın')))
        self.assembly.root.children = [thin, thick, thin]
        self.mass_properties.SelectedItems = (Component('/unrelated.sldprt', mass_kg=99),)
        selection = self.assembly.selection = ['Existing user selection']
        result = bom.collect(self.app)
        by_config = {part.configuration: part for part in result.parts}
        self.assertEqual((by_config['İnce'].quantity, by_config['İnce'].mass_kg,
                          by_config['İnce'].total_mass_kg), (2, .25, .5))
        self.assertEqual(by_config['Kalın'].mass_kg, 1.5)
        self.assertEqual(result.total_mass_kg, 2.0)
        self.assertEqual(result.missing_mass_parts, 0)
        self.assertEqual(len(self.mass_properties.calls), 3)
        self.assertIs(self.assembly.selection, selection)
        self.assertEqual(thin.model.ConfigurationManager.ActiveConfiguration.Name, 'Kalın')
        self.assertEqual(self.assembly.ConfigurationManager.ActiveConfiguration.Name, 'Assembly-Default')
        self.assertEqual((self.app.open_count, self.app.closed), (0, []))

    def test_instance_weights_are_summed_instead_of_reusing_file_weight(self):
        self.assembly.root.children = [Component('/same.sldprt', mass_kg=1.0),
                                       Component('/same.sldprt', mass_kg=1.5)]
        result = bom.collect(self.app)
        part, = result.parts
        self.assertEqual(part.quantity, 2)
        self.assertIsNone(part.mass_kg)
        self.assertEqual(part.total_mass_kg, 2.5)
        self.assertIn('farklı', part.mass_note)
        self.assertEqual([node.mass_kg for node in result.tree], [1.0, 1.5])
        self.assertEqual(result.total_mass_kg, 2.5)

    def test_com_component_selection_uses_dispatch_array(self):
        wrapped = []
        def variant(kind, value):
            wrapped.append(kind)
            return SimpleNamespace(value=value)
        class Calculator(MassProperties):
            @property
            def SelectedItems(self):
                return self.items
            @SelectedItems.setter
            def SelectedItems(self, value):
                self.items = value.value if hasattr(value, 'value') else value
        self.assembly.Extension.CreateMassProperty2 = Calculator
        part = Component('/native.sldprt', mass_kg=.025)
        part._oleobj_ = object()
        self.assembly.root.children = [part]
        with patch.multiple(bom, pythoncom=SimpleNamespace(VT_ARRAY=8192, VT_DISPATCH=9),
                            win32com=SimpleNamespace(client=SimpleNamespace(VARIANT=variant))):
            result = bom.collect(self.app)
        self.assertEqual(wrapped, [8192 | 9])
        self.assertEqual(result.parts[0].mass_kg, .025)

    def test_nested_mass_respects_bom_exclusions_and_does_not_double_count(self):
        inner = Component('/inner.sldasm', children=[Component('/a.sldprt', mass_kg=.5),
            Component('/excluded.sldprt', mass_kg=7, excluded=True)])
        outer = Component('/outer.sldasm', children=[inner, Component('/b.sldprt', mass_kg=.25)])
        self.assembly.root.children = [outer, outer, Component('/suppressed.sldprt', state=0, mass_kg=10)]
        result = bom.collect(self.app)
        self.assertEqual(result.total_mass_kg, 1.5)
        self.assertEqual([node.mass_kg for node in result.tree if node.name == 'outer'], [.75, .75])
        self.assertEqual([node.mass_kg for node in result.tree if node.name == 'inner'], [.5, .5])
        included = bom.collect(self.app, include_excluded=True)
        self.assertEqual(included.total_mass_kg, 15.5)

    def test_invalid_or_unavailable_mass_is_explicit_and_never_a_partial_total(self):
        for value in (None, float('nan'), float('inf'), -.1, True, 'bad'):
            with self.subTest(value=value):
                self.assembly.root.children = [Component('/sub.sldasm', children=[
                    Component('/unknown.sldprt', mass_kg=value), Component('/known.sldprt', mass_kg=2)])]
                result = bom.collect(self.app)
                self.assertEqual(result.total_parts, 2)
                self.assertIsNone(result.total_mass_kg)
                self.assertEqual(result.missing_mass_parts, 1)
                self.assertIsNone(result.tree[0].mass_kg)
                self.assertIn('okunamadı', next(p.mass_note for p in result.parts if p.name == 'unknown'))
        self.assembly.root.children = [Component('/a.sldprt')]
        self.assembly.Extension = SimpleNamespace(CreateMassProperty2=lambda: None)
        result = bom.collect(self.app)
        self.assertIsNone(result.parts[0].total_mass_kg)
        self.assertIn('API', result.parts[0].mass_note)

    def test_recalculation_failure_or_wrong_scope_does_not_export_assembly_mass_as_part_mass(self):
        part = Component('/failed.sldprt', mass_kg=10)
        part.mass_failed = True
        self.assembly.root.children = [part]
        self.assertIsNone(bom.collect(self.app).parts[0].mass_kg)
        class IgnoredSelection(MassProperties):
            @property
            def SelectedItems(self):
                return ()
            @SelectedItems.setter
            def SelectedItems(self, value):
                pass
        self.assembly.Extension.CreateMassProperty2 = IgnoredSelection
        result = bom.collect(self.app)
        self.assertIsNone(result.total_mass_kg)
        self.assertIn('tek bileşen', result.parts[0].mass_note)

    def test_zero_mass_is_valid_and_cancellation_during_mass_read_returns_no_result(self):
        self.assembly.root.children = [Component('/zero.sldprt', mass_kg=0)]
        result = bom.collect(self.app)
        self.assertEqual(result.parts[0].mass_kg, 0)
        self.assertEqual(result.total_mass_kg, 0)
        stop = threading.Event()
        self.mass_properties.on_recalculate = lambda _: stop.set()
        with self.assertRaises(bom.Cancelled):
            bom.collect(self.app, stop_event=stop)

    def test_excel_weights_are_numeric_and_unknowns_and_variable_instances_are_explicit(self):
        self.assembly.root.children = [Component('/sub.sldasm', children=[
            Component('/a.sldprt', mass_kg=.123456789), Component('/a.sldprt', mass_kg=.123456789),
            Component('/b.sldprt', mass_kg=1), Component('/b.sldprt', mass_kg=2),
            Component('/c.sldprt', mass_kg=None)])]
        destination = self.folder/'weighted.xlsx'
        bom.export_excel(bom.collect(self.app), destination)
        book = load_workbook(destination)
        self.addCleanup(book.close)
        parts, tree = book.worksheets
        self.assertEqual(parts['G4'].value, 'Birim ağırlık (kg)')
        self.assertEqual(parts['H4'].value, 'Toplam ağırlık (kg)')
        self.assertAlmostEqual(parts['G5'].value, .123456789)
        self.assertAlmostEqual(parts['H5'].value, .246913578)
        self.assertEqual(parts['G5'].data_type, 'n')
        self.assertEqual(parts['G5'].number_format, '0.000000')
        self.assertEqual(parts['G6'].value, 'Değişken')
        self.assertEqual(parts['H6'].value, 3)
        self.assertEqual(parts['G7'].value, 'Okunamadı')
        self.assertEqual(parts['H7'].value, 'Okunamadı')
        self.assertEqual(parts['E1'].value, 'Eksik ağırlık bilgisi')
        self.assertEqual(tree['J5'].value, 'Okunamadı')
        self.assertEqual(tree['J6'].value, .123456789)
        self.assertEqual(tree['J6'].data_type, 'n')
        self.assertIn('okunamadı', parts['I7'].value)

    def test_excel_bom_total_matches_leaf_parts_with_repeated_subassemblies(self):
        result = bom.collect(self.app)
        destination = self.folder/'total.xlsx'
        bom.export_excel(result, destination)
        book = load_workbook(destination, data_only=True)
        self.addCleanup(book.close)
        for sheet in book:
            self.assertEqual(sheet['E1'].value, 5)
            self.assertEqual(sheet['E1'].data_type, 'n')

    def test_repeated_subassemblies_keep_occurrences_hierarchy_and_natural_name_order(self):
        result = bom.collect(self.app)
        self.assertEqual([p.name for p in result.parts], ['Parça 2', 'Parça 10'])
        self.assertEqual([p.quantity for p in result.parts], [3, 2])
        self.assertEqual(result.total_parts, 5)
        self.assertEqual(len(result.tree), 7)
        assemblies = [n for n in result.tree if n.kind == 'Montaj']
        self.assertEqual(len(assemblies), 2)
        for assembly in assemblies:
            children = [n for n in result.tree if n.parent == assembly.number]
            self.assertEqual([n.name for n in children], ['Parça 2', 'Parça 10'])
            self.assertTrue(all(n.level == 2 for n in children))
        self.assertEqual(result.parts[0].path, r'C:\İmalat\Parça 2.SLDPRT')
        self.assertEqual(self.app.closed, [])
        self.assertEqual(self.app.open_count, 0)
        self.assertEqual(self.assembly.ConfigurationManager.ActiveConfiguration.Name, 'Assembly-Default')

    def test_same_names_different_files_and_configurations_do_not_merge(self):
        self.assembly.root.children = [Component('/a/same.sldprt', 'Default'),
            Component('/a/same.sldprt', 'İnce'), Component('/b/same.sldprt', 'Default'),
            Component('/A/SAME.SLDPRT', 'Default')]
        result = bom.collect(self.app)
        self.assertEqual(len(result.parts), 3)
        self.assertEqual(sorted(p.quantity for p in result.parts), [1, 1, 2])

    def test_nested_assemblies_preserve_parent_and_level(self):
        nested = Component('/deep.sldasm', children=[Component('/part.sldprt')])
        self.assembly.root.children = [Component('/outer.sldasm', children=[nested])]
        result = bom.collect(self.app)
        self.assertEqual([(n.number, n.parent, n.level) for n in result.tree],
                         [('1', '', 1), ('1.1', '1', 2), ('1.1.1', '1.1', 3)])

    def test_suppressed_and_excluded_branches_are_omitted_hidden_are_counted(self):
        self.assembly.root.children += [Component('/supp.sldasm', state=0,
                children=[Component('/unreadable', state=1)]),
            Component('/excluded.sldasm', excluded=True, children=[Component('/only.sldprt')])]
        result = bom.collect(self.app)
        self.assertEqual((result.total_parts, result.suppressed, result.excluded), (5, 1, 1))
        included = bom.collect(self.app, include_excluded=True)
        self.assertEqual((included.total_parts, included.suppressed, included.excluded), (6, 1, 0))

    def test_modern_bom_exclusion_uses_active_configuration_and_rejects_unreadable_value(self):
        child = Component('/excluded.sldprt')
        child.GetExcludeFromBOM2 = lambda option, configs: (True,) if (option, configs) == (1, None) else None
        self.assembly.root.children = [child]
        self.assertEqual(bom.collect(self.app).excluded, 1)
        self.assertEqual(bom.collect(self.app, include_excluded=True).total_parts, 1)
        child.GetExcludeFromBOM2 = lambda *_: None
        with self.assertRaisesRegex(RuntimeError, 'BOM dışı durumu'):
            bom.collect(self.app)

    def test_required_api_errors_prevent_partial_result(self):
        cases = [Component('/a.sldprt', state=1), Component(''), Component('/a.sldprt', ''),
                 Component('/a.unknown'), Component('/sub.sldasm')]
        cases[-1].model = None
        for child in cases:
            with self.subTest(path=child.path):
                self.assembly.root.children = [child]
                with self.assertRaises(RuntimeError):
                    bom.collect(self.app)

    def test_cycles_speedpak_and_unresolved_assembly_fail(self):
        self.assembly.root.children = [Component(self.assembly.path, 'Assembly-Default')]
        with self.assertRaisesRegex(RuntimeError, 'döngü'):
            bom.collect(self.app)
        self.assembly.resolve_status = 1
        with self.assertRaisesRegex(RuntimeError, 'çözümlenemedi'):
            bom.collect(self.app)
        self.assembly.resolve_status = 0
        self.assembly.ConfigurationManager.ActiveConfiguration.IsSpeedPak = True
        with self.assertRaisesRegex(RuntimeError, 'SpeedPak'):
            bom.collect(self.app)

    def test_non_assembly_and_cancelled_scan_do_not_return_partial_bom(self):
        with self.assertRaisesRegex(RuntimeError, 'ana montaj'):
            bom.collect(FakeApp(None))
        stop = threading.Event()
        stop.set()
        with self.assertRaises(bom.Cancelled):
            bom.collect(self.app, stop)
        self.assembly.root.children = [Component(f'/p{i}.sldprt') for i in range(101)]
        stop.clear()
        with self.assertRaises(bom.Cancelled):
            bom.collect(self.app, stop_event=stop, progress=lambda _: stop.set())

    def test_excel_has_sorted_numeric_totals_and_grouped_tree_under_assembly(self):
        result = bom.collect(self.app)
        destination = self.folder/'bom.xlsx'
        bom.export_excel(result, destination)
        book = load_workbook(destination)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ['Parça listesi', 'Ürün ağacı'])
        parts, tree = book.worksheets
        self.assertEqual([parts.cell(i, 2).value for i in (5, 6)], ['Parça 2', 'Parça 10'])
        self.assertEqual([parts.cell(i, 3).value for i in (5, 6)], [3, 2])
        self.assertEqual(parts.cell(5, 3).data_type, 'n')
        self.assertEqual(parts.freeze_panes, 'A5')
        for index, node in enumerate(result.tree, 5):
            self.assertEqual(tree.cell(index, 1).value, node.number)
            self.assertEqual(tree.cell(index, 4).value, node.name)
            self.assertEqual(tree.cell(index, 4).alignment.indent, node.level-1)
            self.assertEqual(tree.row_dimensions[index].outlineLevel, node.level-1)
            if node.kind == 'Montaj':
                self.assertTrue(tree.cell(index, 4).font.bold)
                self.assertEqual(tree.cell(index+1, 2).value, node.number)
        self.assertFalse(tree.sheet_properties.outlinePr.summaryBelow)

    def test_excel_formula_like_names_stay_literal(self):
        result = bom.collect(self.app)
        result = replace(result, parts=(bom.Part('=SUM(1,2)', '=bad.sldprt', '+unsafe', 1),),
                         assembly='=unsafe.sldasm')
        destination = self.folder/'literal.xlsx'
        bom.export_excel(result, destination)
        book = load_workbook(destination, data_only=False)
        self.addCleanup(book.close)
        self.assertEqual(book.worksheets[0]['B5'].value, '=SUM(1,2)')
        self.assertEqual(book.worksheets[0]['B5'].data_type, 's')
        self.assertEqual(book.worksheets[0]['A2'].data_type, 's')

    def test_excel_failure_and_cancellation_preserve_existing_file_and_remove_staging(self):
        result = bom.collect(self.app)
        destination = self.folder/'existing.xlsx'
        destination.write_bytes(b'original')
        with patch('openpyxl.workbook.workbook.Workbook.save', side_effect=OSError('locked')):
            with self.assertRaises(OSError):
                bom.export_excel(result, destination)
        self.assertEqual(destination.read_bytes(), b'original')
        self.assertEqual(list(self.folder.glob('.gokworks*')), [])
        stop = threading.Event()
        real_replace = bom.os.replace
        with patch('bom.exporter.os.replace', side_effect=real_replace) as replace_file, \
                patch('openpyxl.workbook.workbook.Workbook.save', side_effect=lambda _: stop.set()):
            with self.assertRaises(bom.Cancelled):
                bom.export_excel(result, destination, stop)
            replace_file.assert_not_called()
        self.assertEqual(destination.read_bytes(), b'original')
        self.assertEqual(list(self.folder.glob('.gokworks*')), [])

    def test_worker_initializes_com_and_emits_only_complete_result(self):
        com, _ = install_com(self, bom, 'kr')
        events, stop = queue.Queue(), threading.Event()
        with patch('bom.exporter.connect_to_solidworks', return_value=self.app):
            bom.scan_worker(events, stop)
        event, outcome = events.get_nowait()
        self.assertEqual(event, 'bom_done')
        self.assertEqual(outcome['state'], 'completed')
        self.assertEqual(outcome['result'].total_parts, 5)
        com.CoInitialize.assert_called_once()
        com.CoUninitialize.assert_called_once()
        with patch('bom.exporter.collect', side_effect=RuntimeError('read failed')):
            with patch('bom.exporter.connect_to_solidworks', return_value=self.app):
                bom.scan_worker(events, stop)
        self.assertEqual(events.get_nowait()[1]['state'], 'error')

    def test_cancel_during_excel_styling_stops_before_next_row(self):
        from openpyxl.styles import Alignment
        stop = threading.Event()
        destination = self.folder/'existing.xlsx'
        destination.write_bytes(b'original')
        result = bom.collect(self.app)
        def cancel_on_alignment(*args, **kwargs):
            stop.set()
            return Alignment(*args, **kwargs)
        with patch('openpyxl.styles.Alignment', side_effect=cancel_on_alignment) as alignment, \
                patch('openpyxl.workbook.workbook.Workbook.save') as save:
            with self.assertRaises(bom.Cancelled):
                bom.export_excel(result, destination, stop)
            self.assertLessEqual(alignment.call_count, 6)
            save.assert_not_called()
        self.assertEqual(destination.read_bytes(), b'original')
        self.assertEqual(list(self.folder.glob('.gokworks*')), [])

    def test_cancel_before_scan_does_not_connect_to_solidworks(self):
        com, _ = install_com(self, bom, 'kr')
        events, stop = queue.Queue(), threading.Event()
        stop.set()
        with patch('bom.exporter.connect_to_solidworks') as connect:
            bom.scan_worker(events, stop)
            connect.assert_not_called()
        com.CoInitialize.assert_not_called()
        self.assertEqual(list(events.queue), [('bom_done', {'state': 'cancelled'})])


if __name__ == '__main__':
    unittest.main()
