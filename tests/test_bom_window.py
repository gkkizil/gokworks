"""Real Tk BOM controls, tree relationships, cached pages and Excel flow."""

from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

import customtkinter as ctk
from openpyxl import load_workbook
from bom import exporter as bom
from gui.text import BOM_TREE_TAB
from support.gui import GuiTestCase
from support.solidworks import FakeAssembly, FakeApp, install_com
from test_bom import Component


class BomWindowTests(GuiTestCase):
    def setUp(self):
        super().setUp()
        child = Component('/Parça 2.sldprt')
        assembly = FakeAssembly('/Ana.sldasm', [Component('/ÖRNEK-M001.sldasm', children=[child, child]),
                                               Component('/Parça 10.sldprt')])
        self.swapp = FakeApp(assembly)
        self.result = bom.collect(self.swapp)
        self.app.nav_buttons['bom'].invoke()
        self.app.update()

    def scan(self):
        install_com(self, bom, 'kr')
        with patch('bom.exporter.connect_to_solidworks', return_value=self.swapp):
            self.app.bom_widgets['scan'].invoke()
            self.wait_idle()

    def test_real_scan_nested_tree_filter_and_excel_export_preserve_full_bom(self):
        self.scan()
        widgets = self.app.bom_widgets
        rows = widgets['parts'].get_children()
        self.assertEqual(widgets['parts'].item(rows[0], 'values')[:2], ('Parça 2', '2'))
        assembly = next(n for n in self.result.tree if n.kind == 'Montaj')
        children = widgets['tree'].get_children(assembly.number)
        self.assertEqual(len(children), 2)
        self.assertTrue(widgets['tree'].item(assembly.number, 'text').startswith('ÖRNEK-M001'))
        self.app.bom_query.set('Parça 10')
        self.assertEqual(len(widgets['parts'].get_children()), 1)
        destination = self.root/'BOM.xlsx'
        with patch('bom.window.filedialog.asksaveasfilename', return_value=str(destination)):
            widgets['export'].invoke()
            self.wait_idle()
        book = load_workbook(destination)
        self.addCleanup(book.close)
        self.assertEqual(book['Parça listesi'].max_row, 6)  # Search does not filter the export.
        self.assertEqual(book['Parça listesi']['C5'].value, 2)
        self.assertEqual(book['Ürün ağacı'].max_row, 8)
        self.assertIn('Excel kaydedildi', self.app.status_label.cget('text'))

    def test_page_cache_themes_small_window_and_scaling_keep_controls_visible(self):
        self.scan()
        workspace, result = self.app.workspace, self.app.bom_result
        self.addCleanup(lambda: ctk.set_widget_scaling(1))
        self.app.bom_query.set('Parça 10')
        self.app._switch_mode('pdf')
        self.app._switch_mode('bom')
        self.assertIs(self.app.workspace, workspace)
        self.assertIs(self.app.bom_result, result)
        self.assertEqual(self.app.bom_query.get(), 'Parça 10')
        for theme in ('Açık', 'Koyu'):
            for size, scale in (('1000x690', 1), ('1240x860', 1.25)):
                self.app._change_theme(theme)
                ctk.set_widget_scaling(scale)
                # CTk locks min/max to the old size for 1 s during DPI changes.
                # Restore its normal bounds before checking the target size.
                self.app._set_scaled_min_max()
                self.app.geometry(size)
                self.settle()
                self.assertEqual(f'{self.app.winfo_width()}x{self.app.winfo_height()}', size)
                for key in ('scan', 'export', 'stop', 'parts', 'search'):
                    self.assert_inside(self.app.bom_widgets[key], self.app)
                for button in self.app.nav_buttons.values():
                    self.assert_inside(button, self.app.sidebar)
                self.assertGreater(self.app.bom_widgets['parts'].winfo_height()/scale, 180)
                self.assertFalse(self.app.bom_widgets['inspector'].winfo_viewable())
                self.app.bom_widgets['tabs'].set(BOM_TREE_TAB)
                self.settle()
                for key in ('tree', 'expand', 'collapse', 'details'):
                    self.assert_inside(self.app.bom_widgets[key], self.app)
                self.app.bom_widgets['tabs'].set('Parça listesi')

    def test_failed_rescan_discards_stale_results_and_disables_export(self):
        self.scan()
        with patch('bom.exporter.connect_to_solidworks', side_effect=RuntimeError('SolidWorks yok')):
            self.app.bom_widgets['scan'].invoke()
            self.wait_idle()
        self.assertIsNone(self.app.bom_result)
        self.assertFalse(self.app.bom_widgets['parts'].get_children())
        self.assertEqual(self.app.bom_widgets['export'].cget('state'), 'disabled')
        self.assertIn('SolidWorks yok', self.app.status_label.cget('text'))
        self.assertEqual(self.app.bom_widgets['assembly'].cget('text'), 'Tarama tamamlanamadı')
        self.assertEqual(self.app.nav_buttons['pdf'].cget('state'), 'normal')

    def test_detail_dialog_copies_its_row_and_hidden_or_cleared_selection_stays_consistent(self):
        self.scan()
        widgets = self.app.bom_widgets
        widgets['parts'].selection_set('part_0')
        self.settle()
        widgets['details'].invoke()
        dialog = self.app.detail_window
        widgets['parts'].selection_set('part_1')
        self.settle()
        self.assertEqual(widgets['detail_name'].cget('text'), self.result.parts[1].name)
        dialog.copy_path_button.invoke()
        self.assertEqual(self.app.clipboard_get(), self.result.parts[0].path)
        self.app._close_detail()
        widgets['tabs'].set(BOM_TREE_TAB)
        assembly = next(node for node in self.result.tree if node.kind == 'Montaj')
        widgets['tree'].selection_set(assembly.number)
        self.settle()
        widgets['parts'].selection_set('part_0')  # A queued event from the hidden tab.
        self.settle()
        self.assertEqual(widgets['detail_name'].cget('text'), assembly.name)
        widgets['tree'].selection_remove(assembly.number)
        self.settle()
        self.assertEqual(widgets['details'].cget('state'), 'disabled')
        self.assertEqual(widgets['copy_path'].cget('state'), 'disabled')
        self.app._show_bom_details('tree')
        self.assertIsNone(self.app.detail_window)

    def test_metrics_tree_controls_and_selected_component_details(self):
        self.scan()
        widgets = self.app.bom_widgets
        self.assertEqual([widgets['metric_'+key].cget('text') for key in
                          ('unique', 'quantity', 'assemblies')], ['2', '3', '1'])
        widgets['tabs'].set(BOM_TREE_TAB)
        assembly = next(node for node in self.result.tree if node.kind == 'Montaj')
        children = widgets['tree'].get_children(assembly.number)
        widgets['collapse'].invoke()
        self.pump_until(lambda: 'tree_toggle' not in widgets)
        self.assertFalse(widgets['tree'].item(assembly.number, 'open'))
        self.assertEqual(widgets['tree'].get_children(assembly.number), children)
        widgets['expand'].invoke()
        self.pump_until(lambda: 'tree_toggle' not in widgets)
        self.assertTrue(widgets['tree'].item(assembly.number, 'open'))
        widgets['tree'].selection_set(assembly.number)
        self.settle()
        self.assertEqual(widgets['detail_name'].cget('text'), 'ÖRNEK-M001')
        self.assertIn(assembly.instance, widgets['detail_info'].cget('text'))
        widgets['copy_path'].invoke()
        self.assertEqual(self.app.clipboard_get(), assembly.path)
        widgets['details'].invoke()
        self.settle()
        self.assertTrue(self.app.detail_window.winfo_viewable())
        self.app._close_detail()
        widgets['tabs'].set('Parça listesi')
        self.assertEqual(widgets['detail_name'].cget('text'), 'Bir satır seçin')

    def test_empty_search_can_be_cleared_without_changing_export_or_metrics(self):
        self.scan()
        widgets = self.app.bom_widgets
        widgets['parts'].selection_set('part_0')
        self.settle()
        self.app.bom_query.set('olmayan parça')
        self.settle()
        self.assertFalse(widgets['parts'].get_children())
        self.assertTrue(widgets['parts_empty'].winfo_viewable())
        self.assertEqual(widgets['parts_empty_heading'].cget('text'), 'Eşleşen parça yok')
        self.assertEqual(widgets['export'].cget('state'), 'normal')
        self.assertEqual(widgets['metric_quantity'].cget('text'), '3')
        self.assertEqual(widgets['details'].cget('state'), 'disabled')
        widgets['clear_search'].invoke()
        self.settle()
        self.assertEqual(len(widgets['parts'].get_children()), 2)
        self.assertFalse(widgets['parts_empty'].winfo_viewable())

    def test_cancelling_large_preview_clears_partial_rows_and_prevents_export(self):
        parts = tuple(bom.Part(f'Parça {i}', f'/p{i}.sldprt', 'Default', 1) for i in range(10000))
        self.app._bom_busy(True)
        self.app._start_bom_render(replace(self.result, parts=parts))
        self.app._render_bom_batch()
        self.assertIn('render', self.app.bom_widgets)
        self.app.bom_widgets['stop'].invoke()
        self.wait_idle()
        self.assertIsNone(self.app.bom_result)
        self.assertFalse(self.app.bom_widgets['parts'].get_children())
        self.assertEqual(self.app.bom_widgets['export'].cget('state'), 'disabled')
        self.assertEqual(self.app.bom_widgets['assembly'].cget('text'), 'Tarama durduruldu')

    def test_export_cancel_and_failure_preserve_preview_and_existing_file(self):
        self.scan()
        result = self.app.bom_result
        with patch('bom.window.filedialog.asksaveasfilename', return_value=''):
            self.app.bom_widgets['export'].invoke()
        self.assertFalse(self.app.busy)
        path = self.root/'existing.xlsx'
        path.write_bytes(b'original')
        with patch('bom.window.filedialog.asksaveasfilename', return_value=str(path)), \
                patch('bom.window.export_excel', side_effect=PermissionError('Excel dosyası açık')):
            self.app.bom_widgets['export'].invoke()
            self.wait_idle()
        self.assertEqual(path.read_bytes(), b'original')
        self.assertIs(self.app.bom_result, result)
        self.assertIn('Excel dosyası açık', self.app.status_label.cget('text'))
        self.assertEqual(self.app.bom_widgets['export'].cget('state'), 'normal')

    def test_close_during_scan_requests_stop_and_waits_for_worker(self):
        def wait_for_stop(events, stop, **kwargs):
            if not stop.wait(3):
                events.put(('bom_done', dict(state='error', error='Worker stop timeout')))
            else:
                events.put(('bom_done', dict(state='cancelled')))
        with patch('bom.window.scan_worker', side_effect=wait_for_stop):
            self.app.bom_widgets['scan'].invoke()
            with patch('pdf.window.messagebox.askyesno', return_value=True):
                self.app._on_close()
            self.pump_until(lambda: self.app._closed)
        self.assertTrue(self.app.stop_event.is_set())


if __name__ == '__main__':
    import unittest
    unittest.main()
