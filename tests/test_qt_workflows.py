"""Native Qt replacements for the archived Tk workflow and rendering regressions.

Only CAD boundaries are mocked here. Pure engines, files, models, Qt events,
widgets, geometry, settings and threads are real. CAD engine regressions remain
in the original test modules, including typed/raw dispatch, rollback and files.
"""

import csv
import json
import sys
import threading
import gc
from pathlib import Path
from time import perf_counter
from unittest.mock import patch

import pytest
from PySide6.QtCore import Qt, QThread, QTimer, QModelIndex, QCoreApplication, QEvent
from PySide6.QtTest import QAbstractItemModelTester
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox, QWidget

from gui_qt.app import MainWindow, MODULES, create_application
from gui_qt.models import Column, TableModel, TreeModel, FilterProxy
from gui_qt.settings import SettingsStore
from gui_qt.components import PathPicker
from bom.exporter import Part, Node, Result, export_excel
from naming.exporter import Snapshot, Document, path_key
from bolt import rules as bolt
from weight import rules as weight
from iso_fits import rules as iso
from geometric_tolerances.catalog import CATALOG


@pytest.fixture
def window(qtbot, tmp_path):
    app = create_application([])
    w = MainWindow(tmp_path / 'settings.json')
    w.resize(1360, 900)
    qtbot.addWidget(w)
    w.show()
    yield w
    if w.jobs.busy:
        if hasattr(w.jobs.owner, 'pending_apply'):
            w.jobs.owner.pending_apply = None
        w.jobs.cancel()
        qtbot.waitUntil(lambda: not w.jobs.busy, timeout=8000)
    w.settings.flush()
    w.close()
    w.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
    gc.collect()


@pytest.mark.parametrize('mode', [mode for mode, _, _ in MODULES])
@pytest.mark.parametrize('theme', ['dark', 'light'])
def test_every_page_theme_and_cached_navigation(window, qtbot, mode, theme):
    original = window.pages[mode]
    window.set_theme(theme)
    window.navigate(mode)
    assert window.stack.currentWidget() is original
    assert sum(nav.isChecked() for nav in window.navigation.values()) == 1
    assert window.navigation[mode].isChecked()
    assert original.isVisible()
    assert not window.grab().isNull()
    window.navigate('pdf')
    window.navigate(mode)
    assert window.pages[mode] is original
    assert not any(name == 'tkinter' or name == 'customtkinter' for name in sys.modules)


@pytest.mark.parametrize('size', [(1024, 700), (1366, 768), (1920, 1080), (2560, 1440)])
def test_resolution_all_pages_have_usable_viewport(window, qtbot, size):
    window.resize(*size)
    for mode in window.pages:
        window.navigate(mode)
        QApplication.processEvents()
        page = window.pages[mode]
        assert page.width() >= 550
        assert page.height() >= 450
        assert window.navigation[mode].isEnabled()
        assert window.stack.geometry().right() < window.width()


def test_settings_restart_preserves_all_modules_and_unknown_values(qtbot, tmp_path):
    path = tmp_path / 'settings.json'
    original = {'appearance': 'dark', 'pdf_source': '/old/pdf', 'dxf_dwg_format': 'DWG',
                'kr_manual_thickness': '1,5', 'kr_manual_k': '0,274', 'kr_manual_r': '1,5',
                'kr_manual_rows': '[{"thickness":"2","k":"0,33","r":"2"}]',
                'weight_material_group': 'Genel', 'weight_material': 'Pirinç',
                'bolt_unit': 'kgf', 'bolt_tension': '100',
                'naming_parent': str(tmp_path), 'naming_package': 'Yeni', 'future': {'keep': 1}}
    path.write_text(json.dumps(original))
    w = MainWindow(path)
    qtbot.addWidget(w)
    assert w.pages['pdf'].source.text() == '/old/pdf'
    assert w.pages['dxf_dwg'].form.value('format') == 'DWG'
    assert w.pages['k_factor'].recipe().rows[0].k_factor == .274
    assert len(w.pages['k_factor'].recipe().rows) == 2
    assert w.pages['weight'].result.material == 'Pirinç C23000'
    assert w.pages['weight'].result.density == 8750
    assert w.pages['bolt'].result.tension_n == pytest.approx(980.665)
    assert w.pages['naming'].destination.text() == str(tmp_path/'Yeni')
    w.navigate('weight')
    w.set_theme('light')
    w.close()
    data = json.loads(path.read_text())
    assert data['future'] == {'keep': 1}
    assert data['appearance'] == 'light'
    again = MainWindow(path)
    qtbot.addWidget(again)
    assert again.mode == 'weight'
    assert not again.pages['k_factor'].save_files.isChecked()
    again.close()


@pytest.mark.parametrize('raw', ['{broken', '[1,2]', '\udcff'])
def test_malformed_settings_are_retained_before_repair(qapp, tmp_path, raw):
    path = tmp_path/'settings.json'
    path.write_bytes(raw.encode('utf-8', errors='surrogatepass'))
    store = SettingsStore(path)
    store.set('appearance', 'dark')
    assert store.flush()
    assert json.loads(path.read_text())['appearance'] == 'dark'
    backups = list(tmp_path.glob('settings.invalid-*.json'))
    assert len(backups) == 1
    assert backups[0].read_bytes() == raw.encode('utf-8', errors='surrogatepass')


def test_settings_io_error_keeps_pending_changes(qapp, tmp_path):
    store = SettingsStore(tmp_path/'settings.json')
    store.set('pdf_source', 'keep')
    with patch('gui_qt.settings.save_settings', side_effect=OSError('disk unavailable')):
        assert not store.flush()
        assert store.dirty and 'disk unavailable' in store.last_error
    assert store.flush()


def test_invalid_preference_types_and_qt_geometry_fall_back_without_crash(qtbot, tmp_path):
    path = tmp_path/'settings.json'
    path.write_text(json.dumps(dict(appearance=[],mode={},weight_material=[],naming_parent=[1],future='keep')))
    settings = SettingsStore(path)
    settings.qt.setValue('window/geometry','invalid geometry')
    settings.qt.sync()
    w = MainWindow(path)
    qtbot.addWidget(w)
    assert w.mode == 'pdf' and w.appearance == 'dark'
    assert w.pages['weight'].result is not None
    assert w.settings.get('future') == 'keep'
    w.close()


@pytest.mark.parametrize('folder', [False, True])
def test_native_browse_accept_and_cancel(qapp, tmp_path, folder):
    widget = PathPicker('before', folder=folder)
    target = str(tmp_path)
    with patch.object(QFileDialog, 'getExistingDirectory', return_value=target), patch.object(QFileDialog, 'getOpenFileName', return_value=(target, '')):
        widget.choose.click()
    assert widget.text() == target
    with patch.object(QFileDialog, 'getExistingDirectory', return_value=''), patch.object(QFileDialog, 'getOpenFileName', return_value=('', '')):
        widget.choose.click()
    assert widget.text() == target


@pytest.mark.parametrize('mode,extension', [('pdf', '.SLDDRW'), ('dxf_dwg', '.SLDPRT')])
def test_folder_scan_worker_and_dwg_options(window, qtbot, tmp_path, mode, extension):
    page = window.pages[mode]
    (tmp_path/('part'+extension)).touch()
    (tmp_path/('~$temporary'+extension)).touch()
    page.source.setText(str(tmp_path))
    page.scan_button.click()
    qtbot.waitUntil(lambda: not window.jobs.busy)
    assert page.results.model.rowCount() == 1
    assert not page.stop_button.isEnabled()
    if mode == 'dxf_dwg':
        page.output.setText(str(tmp_path/'out'))
        page.form.set_value('format', 'DWG')
        page.form.set_value('group_by_thickness', True)
        calls = []
        def export(*args, **kwargs):
            calls.append((args, kwargs, QThread.currentThread()))
            args[4].put(('done', dict(state='completed', total=1, succeeded=1)))
        with patch('gui_qt.pages.export.export_parts', side_effect=export):
            page.start_button.click()
            qtbot.waitUntil(lambda: not window.jobs.busy)
        assert calls[0][1]['formats'] == ('dwg',)
        assert calls[0][1]['group_by_thickness']
        assert calls[0][2] is not window.thread()


def test_worker_is_responsive_navigable_and_cancellable(window, qtbot):
    owner = window.pages['pdf']
    started = threading.Event()
    worker_threads = []
    def operation(sink, stop):
        worker_threads.append(QThread.currentThread())
        started.set()
        while not stop.wait(.01):
            sink.put(('phase', 'working'))
        sink.put(('done', dict(state='cancelled', total=0)))
    assert window.jobs.start(owner, operation)
    qtbot.waitUntil(started.is_set)
    assert not owner.scan_button.isEnabled()
    for mode in window.pages:
        window.navigate(mode)
        assert window.stack.currentWidget() is window.pages[mode]
    fired = []
    QTimer.singleShot(0, lambda: fired.append(True))
    qtbot.waitUntil(lambda: bool(fired))
    assert not window.jobs.start(window.pages['bom'], operation)
    window.jobs.cancel()
    qtbot.waitUntil(lambda: not window.jobs.busy)
    assert worker_threads[0] is not QApplication.instance().thread()
    assert owner.scan_button.isEnabled()
    assert not owner.stop_button.isEnabled()


def test_worker_error_releases_controls_and_keeps_full_message(window, qtbot):
    page = window.pages['pdf']
    def fail(sink, stop):
        raise RuntimeError('API failure ' + 'details '*100)
    window.jobs.start(page, fail)
    qtbot.waitUntil(lambda: not window.jobs.busy)
    assert 'API failure' in page.status_text
    assert len(page.status.toolTip()) > 500
    assert page.status_details.isVisible() or not page.isVisible()
    assert page.scan_button.isEnabled()


def test_close_cancel_waits_without_terminating_worker(window, qtbot):
    page = window.pages['pdf']
    window.jobs.start(page, lambda sink, stop: stop.wait(3))
    with patch.object(QMessageBox, 'question', return_value=QMessageBox.No):
        window.close()
        assert window.jobs.busy and window.isVisible()
    with patch.object(QMessageBox, 'question', return_value=QMessageBox.Yes):
        window.close()
    qtbot.waitUntil(lambda: not window.jobs.busy)
    qtbot.waitUntil(lambda: not window.isVisible())


def sample_bom():
    return Result('C:/main.sldasm', 'Default',
        (Part('Parça 2','C:/p2.sldprt','Default',2,1.5,3.), Part('Parça 10','C:/p10.sldprt','Other',1,None,None,'API failed')),
        (Node('1','',1,'Ana','Ana','Montaj','C:/main.sldasm','Default'),
         Node('1.1','1',2,'Alt montaj','Alt','Montaj','C:/sub.sldasm','Default'),
         Node('1.1.1','1.1',3,'Parça 2','P2-1','Parça','C:/p2.sldprt','Default',1.5)),1,2,False)


def test_bom_hierarchy_missing_mass_filter_copy_and_complete_excel(window, qtbot, tmp_path):
    page = window.pages['bom']
    page.load_result(sample_bom())
    assert page.tree_model.rowCount() == 1
    parent = page.tree_model.index(0,0)
    assert page.tree_model.rowCount(parent) == 1
    assert page.parts_model.index(1,2).data() == 'Okunamadı'
    assert 'Okunamadı' in page.summary.text()
    page.search.setText('Parça 2')
    page.filter()
    qtbot.waitUntil(lambda: not page.expand_timer.isActive())
    assert page.parts_proxy.rowCount() == 1
    assert page.tree_proxy.rowCount() == 1
    page.parts.setCurrentIndex(page.parts_proxy.index(0,0))
    page.copy_path()
    assert QApplication.clipboard().text() == 'C:/p2.sldprt'
    target = tmp_path/'BOM.xlsx'
    with patch.object(QFileDialog,'getSaveFileName',return_value=(str(target),'')):
        page.export()
        qtbot.waitUntil(lambda: not window.jobs.busy)
    from openpyxl import load_workbook
    wb = load_workbook(target)
    assert len(wb.sheetnames) >= 2
    assert any('Parça 10' == cell.value for row in wb.worksheets[0] for cell in row)
    wb.close()
    page.clear_search()
    assert page.result is not None and page.parts_proxy.rowCount() == 2
    page.clear()
    assert page.result is None and page.parts_model.rowCount() == 0


def test_results_filter_sort_reports_full_errors_and_csv(window, tmp_path):
    panel = window.pages['pdf'].results
    panel.files(['p10','p2'])
    panel.result(dict(source='p2',status='error',detail='COM failure',output=''))
    panel.result(dict(source='p10',status='success',detail='done',output='out.pdf'))
    panel.proxy.sort(0)
    assert panel.proxy.index(0,0).data() == 'p2'
    panel.proxy.set_query('failure')
    assert panel.proxy.rowCount() == 1
    panel.proxy.set_status('success')
    assert panel.proxy.rowCount() == 0
    panel.copy_log_text()
    assert 'COM failure' in QApplication.clipboard().text()
    path = tmp_path/'results.csv'
    panel.save_csv(path)
    with path.open(encoding='utf-8-sig') as stream:
        rows = list(csv.reader(stream, delimiter=';'))
    assert len(rows) == 3 and 'COM failure' in rows[1]
    panel.clear_button.click()
    assert panel.proxy.rowCount() == 2 and len(panel.results) == 2
    panel.search.setText('missing')
    panel.filter.setCurrentIndex(1)
    panel.reset()
    panel.files(['new'])
    assert panel.proxy.rowCount() == 1
    assert not panel.search.text() and panel.filter.currentIndex() == 0


def naming_snapshot(tmp_path):
    src = tmp_path/'source'
    src.mkdir()
    paths = [str(src/name) for name in ('main.sldasm','sub.sldasm','p.sldprt','p.slddrw')]
    for path in paths:
        Path(path).touch()
    docs = tuple(Document(path, dependencies, (0,0)) for path, dependencies in zip(paths,
                 ((paths[1],paths[2]),(paths[2],),(),(paths[2],))))
    return Snapshot(paths[0],docs)


def test_naming_folder_source_exact_destination_hierarchy_and_collision(window, qtbot, tmp_path):
    page = window.pages['naming']
    snapshot = naming_snapshot(tmp_path)
    page.source_folder.setText(str(tmp_path/'source'))
    page.discover()
    qtbot.waitUntil(lambda: not window.jobs.busy)
    assert page.assemblies.count() == 2
    page.destination.setText(str(tmp_path/'target'))
    page.form.set_value('mode','Numaralı')
    page.form.set_value('prefix','BHV-R')
    page.form.set_value('suffix','630')
    page.form.set_value('digits','2')
    page.load_snapshot(snapshot)
    qtbot.waitUntil(lambda: page.valid_plan and not window.jobs.busy)
    assert page.valid_plan and len(page.entries) == 4
    assert page.model.rowCount() == 1
    assert all('BHV-R' in entry.destination and '630' in entry.destination for entry in page.entries)
    assert len({path_key(e.destination) for e in page.entries}) == 4
    page.overrides[path_key(snapshot.documents[1].path)] = 'same'
    page.overrides[path_key(snapshot.documents[0].path)] = 'same'
    page.update_preview()
    qtbot.waitUntil(lambda: not window.jobs.busy)
    assert not page.valid_plan and not page.save_button.isEnabled()
    page.reset_overrides()
    qtbot.waitUntil(lambda: not window.jobs.busy)
    assert page.valid_plan
    (tmp_path/'target').mkdir()
    (tmp_path/'target'/'keep.txt').write_text('keep')
    page.update_preview()
    qtbot.waitUntil(lambda: not window.jobs.busy)
    assert not page.valid_plan
    assert (tmp_path/'target'/'keep.txt').read_text() == 'keep'
    page.source_file.setText(snapshot.root)
    assert page.snapshot is None


def test_naming_plan_filesystem_checks_are_off_gui_thread(window, qtbot, tmp_path):
    page = window.pages['naming']
    snapshot = naming_snapshot(tmp_path)
    page.destination.setText(str(tmp_path/'new'))
    release = threading.Event()
    calls = []
    from naming.exporter import validate_destination
    def delayed(*args):
        calls.append(QThread.currentThread())
        release.wait(2)
        return validate_destination(*args)
    with patch('gui_qt.pages.naming.engine.validate_destination',side_effect=delayed):
        page.load_snapshot(snapshot)
        qtbot.waitUntil(lambda: bool(calls))
        window.navigate('bolt')
        assert window.stack.currentWidget() is window.pages['bolt']
        assert calls[0] is not QApplication.instance().thread()
        release.set()
        qtbot.waitUntil(lambda: page.valid_plan and not window.jobs.busy)


@pytest.mark.parametrize('outcome', ['completed','error','cancelled'])
def test_kr_exact_inputs_preview_labels_and_queued_scope(window, qtbot, outcome):
    page = window.pages['k_factor']
    page.rows_model.replace([dict(thickness='1,5',k='0,274',r='1,5'),dict(thickness='2',k='0,33',r='1')])
    page.persist_rows()
    assert page.recipe().rows[1].radius_mm == 1
    assert not page.recipe().all_bends
    release = threading.Event()
    calls = []
    def update(recipe, sink, stop, **kwargs):
        calls.append((recipe,kwargs))
        if kwargs['preview']:
            release.wait(2)
            sink.put(('result',dict(source='sheet [Default]',status='success',preview=True,detail='ready')))
        sink.put(('done',dict(state=outcome if kwargs['preview'] else 'completed',scope=('assembly','Default','signature'),succeeded=1)))
    with patch('gui_qt.pages.kr.engine.update_assembly',side_effect=update):
        page.scan_button.click()
        qtbot.waitUntil(lambda: bool(calls))
        assert page.tabs.currentIndex() == 2
        page.apply_button.click()
        assert page.pending_apply is not None
        release.set()
        qtbot.waitUntil(lambda: not window.jobs.busy)
        # Idle means the entire queued operation is complete, not the gap
        # before an unparented timer starts another job after this mock exits.
        if outcome == 'completed':
            assert len(calls) == 2
        QApplication.processEvents()
        qtbot.waitUntil(lambda: not window.jobs.busy)
    assert len(calls) == (2 if outcome == 'completed' else 1)
    assert calls[0][1]['save'] is False
    if outcome == 'completed':
        assert calls[1][1]['expected_scope'] == ('assembly','Default','signature')
    else:
        assert calls[0][1]['preview']


def test_kr_stop_discards_queued_apply_and_helper_never_overwrites(window, qtbot):
    page = window.pages['k_factor']
    page.rows_model.replace([dict(thickness='2',k='0,274',r='1,5')])
    original = page.recipe()
    page.helper.set_value('thickness','3')
    page.calculate_helper()
    assert page.recipe() == original
    calls = []
    def update(recipe,sink,stop,**kwargs):
        calls.append(kwargs)
        stop.wait(2)
        sink.put(('done',dict(state='cancelled')))
    with patch('gui_qt.pages.kr.engine.update_assembly',side_effect=update):
        page.scan()
        page.apply()
        page.stop_button.click()
        qtbot.waitUntil(lambda: not window.jobs.busy)
    assert len(calls) == 1 and page.pending_apply is None


def test_close_at_preview_completion_does_not_start_queued_cad_write(window, qtbot):
    page = window.pages['k_factor']
    page.rows_model.replace([dict(thickness='2',k='0,274',r='1,5')])
    release = threading.Event()
    calls = []
    def update(recipe, sink, stop, **kwargs):
        calls.append(kwargs)
        release.wait(2)
        sink.put(('done',dict(state='completed',scope=('assembly','Default','signature'))))
    window.jobs.finished.connect(window.close)
    with patch('gui_qt.pages.kr.engine.update_assembly',side_effect=update):
        page.scan()
        page.apply()
        release.set()
        qtbot.waitUntil(lambda: not window.jobs.busy)
        QApplication.processEvents()
    assert window.closing and not window.isVisible()
    assert len(calls) == 1 and page.pending_apply is None


@pytest.mark.parametrize('unit', ['kg','N','kN'])
@pytest.mark.parametrize('mode', bolt.MODES)
def test_bolt_unit_conversion_and_all_load_modes(window, unit, mode):
    page = window.pages['bolt']
    page.form.set_value('mode',mode)
    page.form.set_value('unit',unit)
    page.update_result()
    assert page.result is not None
    if mode != 'Kesme':
        assert page.result.tension_n == pytest.approx(10000)
    if mode != 'Çekme':
        assert page.result.shear_n == pytest.approx(4000)
    assert len(page.diagrams['dimensions'].scene().items()) > 15
    page.copy_report()
    assert 'Sp' in QApplication.clipboard().text() and 'Kapsam' in QApplication.clipboard().text()


@pytest.mark.parametrize('metric', bolt.THREADS)
def test_bolt_standard_and_fine_pitch_choices(window, metric):
    page = window.pages['bolt']
    page.form.set_value('metric',metric)
    assert 'Standart' in page.form.value('pitch')
    assert not page.form.fields['pitch'].isEnabled()
    page.form.set_value('thread_series','İnce diş')
    for index in range(page.form.fields['pitch'].count()):
        page.form.fields['pitch'].setCurrentIndex(index)
        page.update_result()
        assert page.result is not None
        assert page.result.pitch == page.form.fields['pitch'].currentData()


@pytest.mark.parametrize('ratio,verdict,progress', [(0, 'Yük yok', 0), (.5, 'Güvenli', 500), (1.5, 'Güvensiz', 1000)])
def test_bolt_safety_graph_reports_real_overload_and_resets_invalid_input(window, ratio, verdict, progress):
    page = window.pages['bolt']
    window.navigate('bolt')
    page.form.set_value('unit', 'N')
    capacity = page.result.tension_capacity_n
    page.form.set_value('tension', str(capacity * ratio))
    page.update_result()
    assert page.safety_card.isVisible()
    assert verdict in page.safety_verdict.text()
    assert page.safety_progress.value() == progress
    assert f'%{bolt.fmt(ratio*100, 1)}' in page.safety_usage.text()
    page.copy_report()
    assert 'aynı anda kullanılamaz' in QApplication.clipboard().text()
    if ratio > 1:
        assert '2 cıvata' in page.safety_usage.text()
        assert '−%' in page.safety_usage.text()
    window.set_theme('light')
    assert page.safety_progress.value() == progress
    page.form.set_value('tension', '')
    page.update_result()
    assert page.result is None and page.safety_progress.value() == 0
    assert 'Hesaplanamadı' in page.safety_verdict.text()
    assert not page.safety_usage.text()
    assert 'Çekme:' not in page.capacity_text.text()


def test_bolt_load_diagram_selection_is_restored_and_combined_limits_are_explained(qtbot, tmp_path):
    path = tmp_path/'settings.json'
    path.write_text(json.dumps(dict(bolt_view='Yük gösterimi', bolt_mode='Birleşik yük')))
    w = MainWindow(path)
    qtbot.addWidget(w)
    page = w.pages['bolt']
    assert page.diagram_tabs.currentWidget() is page.diagrams['load']
    assert page.diagrams['load'].scene().items()
    assert 'sabitken çekme sınırı' in page.capacity_text.text()
    assert 'sabitken kesme sınırı' in page.capacity_text.text()
    assert 'eklenebilecek yük değildir' in page.capacity_text.text()
    page.diagram_tabs.setCurrentIndex(0)
    w.close()
    assert json.loads(path.read_text())['bolt_view'] == 'Ölçüler'


@pytest.mark.parametrize('shaft,fragment', [('h6', 'Minimum boşluk sıfır; bu bir ara geçme değildir'),
                                         ('k6', 'ölçüsüne göre boşluk veya sıkılık'),
                                         ('s6', 'Sıkılık aralığı')])
def test_iso_fit_explanation_is_visible_and_follows_limits_and_units(window, shaft, fragment):
    page = window.pages['iso_fits']
    window.navigate('iso_fits')
    page.form.set_value('shaft', shaft)
    page.update_result()
    assert page.details.isVisible() and fragment in page.details.text()
    assert 'Kullanım yorumu' in page.usage_label.text()
    assert 'Kırmızı alan' in page.section_help.text()
    assert 'en sık oluşacak ölçü' in page.section_help.text()
    page.form.set_value('unit', 'µm')
    page.update_result()
    if shaft != 'h6':
        assert 'µm' in page.details.text()
    page.copy_report()
    assert fragment in QApplication.clipboard().text()
    page.form.set_value('diameter', '0')
    page.update_result()
    assert page.result is None and fragment not in page.details.text()
    assert not page.usage_label.text()


def test_manual_iso_keeps_full_reference_distance_notes_and_scrollable_tabs(window):
    page = window.pages['iso_fits']
    page.form.set_value('method', 'Manuel kontrol')
    for key, value in dict(hole_lower='-0,05', hole_upper='0', shaft_lower='-0,1', shaft_upper='0,1').items():
        page.form.set_value(key, value)
    page.update_result()
    assert (page.result.minimum_clearance_nm, page.result.maximum_clearance_nm) == (-150000, 100000)
    assert 'tolerans genişliğini aşıyor' in page.compatibility_label.text()
    assert 'En büyük tek limit farkı' in page.compatibility_label.text()
    assert 'Girilen − ISO' in page.comparison_label.text()
    assert '11 delik / 18 mil' in page.comparison_label.text()
    for index, title in enumerate(('Montaj uyumu', 'Ölçü sınırları', 'ISO karşılaştırması')):
        assert page.manual_tabs.tabText(index) == title
        assert page.manual_tabs.widget(index).widget().text()
    page.copy_report()
    assert 'tolerans genişliğini aşıyor' in QApplication.clipboard().text()


def test_gdt_retains_measurement_instructions_size_rule_and_boundary_assessment(window):
    page = window.pages['geometric_tolerances']
    window.navigate('geometric_tolerances')
    assert 'üst ve alt ölçü sınırlarının farkı' in page.size_rule_help.text()
    assert 'referans (datum) kullanılmaz' in page.datum_help.text()
    assert 'Ölçüm girişinin anlamı' in page.measurement_label.text()
    assert 'Boyut toleransı girilmedi' in page.assessment_label.text()
    page.form.set_value('size_tolerance', '0,1')
    page.form.set_value('measured', '0,05')
    page.update_result()
    assert 'tam sınırda' in page.assessment_label.text()
    assert 'MMC zarfı ölçülmedi' in page.assessment_label.text()
    assert 'ölçüm belirsizliği' in page.assessment_label.text()
    page.copy_report()
    assert 'MMC zarfı ölçülmedi' in QApplication.clipboard().text()
    page.form.set_value('tolerance', '')
    page.update_result()
    assert page.result is None
    assert 'tam sınırda' not in page.assessment_label.text()


@pytest.mark.parametrize('kind', weight.KINDS)
def test_weight_all_shapes_and_native_sections(window, kind):
    page = window.pages['weight']
    page.form.set_value('kind',kind)
    page.update_result()
    assert page.result is not None
    assert page.result.total_mass_kg > 0
    assert page.diagrams['section'].scene().items()
    page.form.set_value('method','Katalog alanından')
    page.form.set_value('area_cm2','20')
    page.update_result()
    assert page.result.area_mm2 == 2000


def test_weight_material_densities_tool_names_and_unit_conversion(window):
    page = window.pages['weight']
    page.form.set_value('kind','Sac levha')
    page.form.set_value('b','2000')
    page.form.set_value('t','10')
    page.form.set_value('length','1000')
    page.update_result()
    assert page.result.total_mass_kg == 157
    for unit in ('cm','m','mm'):
        page.form.set_value('unit',unit)
        page.update_result()
        assert page.result.total_mass_kg == pytest.approx(157)
    for group, materials in weight.MATERIAL_GROUPS.items():
        page.form.set_value('material_group',group)
        box = page.form.fields['material']
        assert box.count() == len(materials)
        for i, name in enumerate(materials):
            box.setCurrentIndex(i)
            page.update_result()
            assert page.result.density == weight.MATERIALS[name]
            assert 'kg/m³' in box.currentText()
    assert 'density' not in page.form.fields


@pytest.mark.parametrize('manual', [True,False])
def test_iso_views_samples_manual_limits_matching_and_report(window, manual):
    page = window.pages['iso_fits']
    page.form.set_value('method','Manuel kontrol' if manual else 'ISO seçimi')
    if manual:
        for key,value in [('hole_lower','-0,05'),('hole_upper','0'),('shaft_lower','-0,1'),('shaft_upper','0,1')]:
            page.form.set_value(key,value)
    for view in ('Montaj kesiti','Tolerans bölgeleri'):
        for sample in ('En az boşluk','Orta değer','En çok boşluk'):
            page.form.set_value('view',view)
            page.form.set_value('sample',sample)
            page.update_result()
            assert page.result is not None
            assert page.diagrams['fit'].scene().items()
            assert page.view_tabs.tabText(page.view_tabs.currentIndex()) == view
            assert page.sample_tabs.tabText(page.sample_tabs.currentIndex()) == sample
    page.view_tabs.setCurrentIndex(0)
    page.sample_tabs.setCurrentIndex(0)
    assert page.form.value('view') == 'Montaj kesiti'
    assert page.form.value('sample') == 'En az boşluk'
    if manual:
        assert (page.result.minimum_clearance_nm,page.result.maximum_clearance_nm) == (-150000,100000)
        assert page.manual_tabs.count() == 3
        assert 'referans' in page.compatibility_label.text()
        assert 'eşleşme' in page.comparison_label.text()
    else:
        assert page.preset_model.rowCount() == 3
        page.search.setText('H7')
        assert page.preset_proxy.rowCount() > 0
    page.copy_report()
    assert 'ISO 286' in QApplication.clipboard().text()


@pytest.mark.parametrize('item', CATALOG, ids=lambda item:item.key)
def test_gdt_catalog_frames_measurement_acceptance_and_legacy(window, item):
    page = window.pages['geometric_tolerances']
    page.kind = item.key
    page.form.set_value('datum','A')
    page.form.set_value('measured','0,01')
    page.update_result()
    if not item.active:
        assert page.result is None
        assert not page.diagrams['frame'].scene().items()
        return
    assert page.result is not None
    assert page.result.accepted
    assert page.diagrams['frame'].scene().items()
    assert page.diagrams['zone'].scene().items()
    assert page.form.fields['datum'].isEnabled() == (item.datum != 'none')
    page.form.set_value('measured','0,06')
    page.update_result()
    assert not page.result.accepted
    page.copy_report()
    assert item.name in QApplication.clipboard().text()
    assert 'gdt_measured' not in page.settings.values


def test_models_are_valid_and_large_data_creates_no_row_widgets(qapp):
    columns = (Column('Name','name'),Column('Mass','mass'))
    table = TableModel(columns)
    tree = TreeModel(columns)
    testers = [QAbstractItemModelTester(model, QAbstractItemModelTester.FailureReportingMode.Warning) for model in (table,tree)]
    before = len(QApplication.allWidgets())
    rows = [dict(id=str(i),parent=str((i-1)//5) if i else '',name=f'Parça {i}',mass=i/10) for i in range(50000)]
    table.replace(rows)
    tree.replace(rows)
    assert table.rowCount() == 50000 and tree.rowCount() == 1
    assert len(QApplication.allWidgets()) == before
    # The attached Qt diagnostic tester traverses the complete model on reset.
    # Keep that correctness audit, but measure production models separately:
    # the application does not attach QAbstractItemModelTester to user data.
    runtime_table = TableModel(columns)
    runtime_tree = TreeModel(columns)
    start = perf_counter()
    runtime_table.replace(rows)
    runtime_tree.replace(rows)
    proxy = FilterProxy(tree=True)
    proxy.setSourceModel(runtime_tree)
    proxy.set_query('Parça 49999')
    assert proxy.rowCount() == 1
    assert perf_counter()-start < 15


def test_repeated_worker_lifecycle_has_no_running_threads(window, qtbot):
    for _ in range(20):
        window.jobs.start(window.pages['pdf'], lambda sink, stop: sink.put(('phase','done')))
        qtbot.waitUntil(lambda: not window.jobs.busy)
        QApplication.processEvents()
    assert not any(thread.isRunning() for thread in window.findChildren(QThread))


def test_vector_labels_keep_engineering_center_coordinates(qapp):
    from gui_qt.diagrams import VectorScene
    scene = VectorScene(500,300)
    item = scene.create_text(250,150,text='MİL Ø50',anchor='center',font=('Segoe UI',-14),fill='#ABCDEF')
    center = item.sceneBoundingRect().center()
    assert center.x() == pytest.approx(250)
    assert center.y() == pytest.approx(150)
