"""Real geometry/files: import, NFP packing, stock, accounting and portable exports."""

from dataclasses import replace
import json
import math
import threading

import ezdxf
import pytest
from shapely.geometry import Polygon, box, MultiPoint

from nesting.accounting import mass, calculate, money
from nesting.dxf_import import read_dxf, MissingUnits
from nesting.dxf_export import export_dxf
from nesting.engine import nest
from nesting.geometry import placed, footprint, minkowski
from nesting.materials import new_project
from nesting.models import Part, Sheet, Material, Placement, number
from nesting.storage import save_project, load_project
from nesting.validation import validate_result


def drawing(tmp_path, build, *, unit=4, name='Part.dxf'):
    doc = ezdxf.new('R2018'); doc.units = unit; build(doc.modelspace(), doc)
    path = tmp_path/name; doc.saveas(path); return path


def rectangle(m, width=100, height=50):
    return m.add_lwpolyline([(0,0),(width,0),(width,height),(0,height)],close=True)


def project(parts=None, width=300, height=300):
    p = new_project(); p.sheets[0].width = width; p.sheets[0].height = height
    p.sheets[0].price_kg = 47.5
    p.options.seconds = 6; p.options.quality = 'Hızlı'; p.options.margin = 5
    p.parts = parts or [Part.create('A',box(0,0,60,40))]
    return p


def test_rectangle_dxf_area_and_mass(tmp_path):
    path = drawing(tmp_path,lambda m,d:rectangle(m))
    before = path.read_bytes(); p = read_dxf(path)[0]
    assert p.area == pytest.approx(5000)
    assert mass(p.area,2,7850) == pytest.approx(.0785)
    assert p.polygon.bounds == (0,0,100,50)
    assert path.read_bytes() == before


def test_tiny_join_uses_recorded_safety_allowance_without_rounding_closed_shapes(tmp_path):
    def build(m,d):
        m.add_line((0,0),(20,0)); m.add_line((20.0004,0),(20,10))
        m.add_line((20,10),(0,10)); m.add_line((0,10),(0,0))
    part = read_dxf(drawing(tmp_path,build))[0]
    assert part.endpoint_error == pytest.approx(.0004)
    assert part.geometry_error >= .0004-1e-9
    p = project([part]); r = nest(p)
    assert all(placed(part,x).bounds[0] >= p.options.margin+part.geometry_error for x in r.placements)


def test_circle_hole_net_area_and_saved_analytic_geometry(tmp_path):
    def build(m,d): rectangle(m); m.add_circle((30,25),10)
    p = read_dxf(drawing(tmp_path,build),tolerance=.001)[0]
    assert p.area == pytest.approx(5000-math.pi*100,abs=.05)
    assert len(p.polygon.interiors) == 1
    assert 'CIRCLE' in p.original_dxf


def test_arc_contour_area(tmp_path):
    def build(m,d): m.add_arc((0,0),10,0,180); m.add_line((-10,0),(10,0))
    p = read_dxf(drawing(tmp_path,build),tolerance=.001)[0]
    assert p.area == pytest.approx(math.pi*50,abs=.05)


@pytest.mark.parametrize('kind',['open','cross','duplicate','dangling','self'])
def test_bad_contours_are_not_silently_repaired(tmp_path,kind):
    def build(m,d):
        if kind == 'open': m.add_line((0,0),(10,0))
        elif kind == 'self': m.add_lwpolyline([(0,0),(10,10),(0,10),(10,0)],close=True)
        elif kind == 'duplicate': rectangle(m); rectangle(m)
        elif kind == 'dangling': rectangle(m); m.add_line((100,0),(120,0))
        else: rectangle(m); m.add_lwpolyline([(50,-20),(120,-20),(120,20),(50,20)],close=True)
    with pytest.raises(ValueError): read_dxf(drawing(tmp_path,build))


@pytest.mark.parametrize('unit,choice,scale',[(0,None,None),(0,'inch',25.4),(0,'cm',10),(4,'inch',1),(1,None,25.4)])
def test_insunits_or_explicit_missing_unit(tmp_path,unit,choice,scale):
    path = drawing(tmp_path,lambda m,d:rectangle(m,10,5),unit=unit)
    if scale is None:
        with pytest.raises(MissingUnits): read_dxf(path)
    else:
        assert read_dxf(path,unit=choice)[0].area == pytest.approx(50*scale**2)


def test_blocks_nonuniform_transforms_layers_and_annotations(tmp_path):
    def build(m,d):
        block = d.blocks.new('B'); rectangle(block,10,5)
        m.add_blockref('B',(20,30),dxfattribs={'rotation':90,'xscale':2,'yscale':3})
        d.layers.new('CONSTRUCTION'); m.add_line((0,0),(300,300),dxfattribs={'layer':'CONSTRUCTION'})
        m.add_text('Not a cutting line'); m.add_linear_dim(base=(0,20),p1=(0,0),p2=(10,0))
    path = drawing(tmp_path,build)
    p = read_dxf(path,excluded_layers=['CONSTRUCTION'])[0]
    assert p.area == pytest.approx(300)
    assert p.polygon.bounds == pytest.approx((0,0,15,20))


@pytest.mark.parametrize('kind',['polyline','bulge','ellipse','spline','lines'])
def test_supported_entity_types(tmp_path,kind):
    def build(m,d):
        if kind == 'polyline': m.add_polyline2d([(0,0),(40,0),(40,20),(0,20)],close=True)
        elif kind == 'bulge': m.add_lwpolyline([(0,0,1),(20,0,1)],format='xyb',close=True)
        elif kind == 'ellipse': m.add_ellipse((0,0),major_axis=(20,0,0),ratio=.5)
        elif kind == 'spline':
            m.add_spline([(0,0),(20,10),(40,0)]); m.add_line((40,0),(0,0))
        else:
            for a,b in [((0,0),(40,0)),((40,0),(40,20)),((40,20),(0,20)),((0,20),(0,0))]: m.add_line(a,b)
    parts = read_dxf(drawing(tmp_path,build))
    assert len(parts) == 1 and parts[0].area > 10


def test_separate_and_grouped_components_and_hole_islands(tmp_path):
    def build(m,d):
        rectangle(m,100,100); m.add_circle((50,50),30); m.add_circle((50,50),10)
        m.add_lwpolyline([(120,0),(140,0),(140,20),(120,20)],close=True)
    path = drawing(tmp_path,build)
    parts = read_dxf(path)
    assert len(parts) == 3
    grouped = read_dxf(path,grouped=True)
    assert len(grouped) == 1
    assert grouped[0].area == pytest.approx(sum(p.area for p in parts))


def test_ten_different_shapes_pack_without_overlap_and_edge_violation():
    parts = [Part.create(f'P{i}',Polygon([(0,0),(20+i,0),(22+i,12+i),(7,25+i),(0,15)])) for i in range(10)]
    p = project(parts); r = nest(p)
    assert len(r.placements) == 10 and not r.unplaced
    assert validate_result(p,r)
    # Separate direct checks, independent of the engine candidate calculations.
    for i,a in enumerate(r.placements):
        g = placed(p.part(a.part_id),a)
        assert g.bounds[0] >= p.options.margin
        for b in r.placements[i+1:]:
            if (a.sheet_id,a.sheet_number) == (b.sheet_id,b.sheet_number):
                h = placed(p.part(b.part_id),b)
                assert g.intersection(h).area < 1e-6
                assert g.distance(h) >= p.options.gap-1e-6


def test_fifty_copies_are_distinct_instances():
    p = project([Part.create('A',box(0,0,15,10),quantity=50)],width=250,height=250)
    r = nest(p)
    assert not r.unplaced and len(r.placements) == 50
    assert len({(p.part_id,p.instance) for p in r.placements}) == 50
    assert validate_result(p,r)


def test_second_sheet_and_stock_exhaustion():
    p = project([Part.create('A',box(0,0,75,75),quantity=2)],width=100,height=100)
    r = nest(p)
    assert {x.sheet_number for x in r.placements} == {1,2}
    p.sheets[0].stock = 1; r = nest(p)
    assert len(r.placements) == len(r.unplaced) == 1
    assert validate_result(p,r,complete=False)
    with pytest.raises(ValueError,match='yerleşmedi'): validate_result(p,r)


def test_search_considers_large_sheet_to_reduce_sheet_count():
    p = project([Part.create('A',box(0,0,75,75),quantity=4)],width=100,height=100)
    p.sheets.append(replace(p.sheets[0],id='large',name='Large',width=200,height=200))
    r = nest(p)
    assert not r.unplaced
    assert {x.sheet_id for x in r.placements} == {'large'}
    assert len({(x.sheet_id,x.sheet_number) for x in r.placements}) == 1


def test_true_shape_packs_inside_concavity_not_rectangles():
    # Bounding rectangles overlap 60×60; both actual contours fit one sheet.
    p = project([Part.create('L',Polygon([(0,0),(100,0),(100,40),(40,40),(40,100),(0,100)])),
                 Part.create('Square',box(0,0,60,60))],width=100,height=100)
    p.options.margin = p.options.gap = p.options.kerf = 0; p.options.rotation = False
    r = nest(p)
    assert not r.unplaced and {x.sheet_number for x in r.placements} == {1}
    a,b = [placed(p.part(x.part_id),x) for x in r.placements]
    assert box(*a.bounds).intersection(box(*b.bounds)).area == pytest.approx(3600)
    assert a.intersection(b).area == 0


def test_minkowski_matches_bruteforce_convex_sum():
    a = Polygon([(0,0),(10,0),(12,7),(4,12)])
    b = Polygon([(-3,0),(1,-5),(8,0),(3,4)])
    expected = MultiPoint([(x+u,y+v) for x,y in list(a.exterior.coords)[:-1] for u,v in list(b.exterior.coords)[:-1]]).convex_hull
    assert minkowski(a,b).symmetric_difference(expected).area < 1e-7


def test_weight_example_and_conservation():
    assert mass(1250*2500,2,7850) == pytest.approx(49.0625)
    p = project(); r = nest(p); t = calculate(p,r)
    assert t['gross_weight'] >= t['net_weight']
    assert t['waste_weight']+t['net_weight'] == pytest.approx(t['gross_weight'])
    assert t['waste_area']+t['net_area'] == pytest.approx(t['gross_area'])
    assert t['efficiency'] == pytest.approx(100*t['net_area']/t['gross_area'])


def test_different_materials_and_thicknesses_never_mix():
    p = project([Part.create('Steel2',box(0,0,30,40)),Part.create('Steel3',box(0,0,30,40),thickness=3),
                 Part.create('Al2',box(0,0,30,40),material='EN AW-5754')])
    p.sheets.extend([replace(p.sheets[0],id='steel3',thickness=3),replace(p.sheets[0],id='al2',material='EN AW-5754')])
    r = nest(p)
    assert len({x.sheet_id for x in r.placements}) == 3
    assert validate_result(p,r)


def test_costs_allocated_exactly_and_optional_operations_separate():
    p = project([Part.create('A',box(0,0,10,10),quantity=7),Part.create('B',box(0,0,20,10),quantity=3)])
    p.sheets[0].pricing = 'sheet'; p.sheets[0].price_sheet = 100.01
    p.operations.cut_price_m = 5; p.operations.piercing_price = 2; p.operations.setup = 20
    p.operations.fixed = 10; p.operations.profit_percent = 15; p.operations.vat_percent = 20
    r = nest(p); t = calculate(p,r)
    assert money(sum(x['allocation'] for x in t['parts'])) == money(t['material_cost'])
    assert t['material_cost'] == 100.01
    assert t['cut_cost'] > 0 and t['piercing_cost'] == 20
    assert t['total'] > t['material_cost']+t['cut_cost']+t['piercing_cost']


def test_no_part_in_hole_and_grain_rotation_restrictions():
    ring = Polygon([(0,0),(90,0),(90,90),(0,90)], [[(10,10),(80,10),(80,80),(10,80)]])
    p = project([Part.create('ring',ring),Part.create('inner',box(0,0,50,50))],width=110,height=110)
    r = nest(p); assert len({x.sheet_number for x in r.placements}) == 2
    p = project([Part.create('A',box(0,0,120,30))],width=60,height=160)
    r = nest(p); assert not r.unplaced
    assert r.placements[0].angle in (90,270)
    p.options.grain_locked = True
    assert nest(p).unplaced


@pytest.mark.parametrize('tamper',['overlap','edge','material','rotation','count','stock','nan'])
def test_independent_validation_rejects_tampered_placements(tamper):
    p = project([Part.create('A',box(0,0,40,30),quantity=2)]); r = nest(p)
    if tamper == 'overlap': r.placements[1].x = r.placements[0].x; r.placements[1].y = r.placements[0].y
    if tamper == 'edge': r.placements[0].x = -20
    if tamper == 'material': p.sheets[0].material = 'AISI 304'; r.digest = p.fingerprint()
    if tamper == 'rotation': r.placements[0].angle = 45
    if tamper == 'count': r.placements.append(r.placements[0])
    if tamper == 'stock': p.sheets[0].stock = 0; r.digest = p.fingerprint()
    if tamper == 'nan': r.placements[0].x = float('nan')
    with pytest.raises(ValueError): validate_result(p,r)


def test_portable_project_reports_and_exact_curve_export(tmp_path,qapp):
    def build(m,d):
        m.add_arc((0,0),30,0,180); m.add_line((-30,0),(30,0)); m.add_circle((0,12),3)
    source = drawing(tmp_path,build)
    p = project(read_dxf(source,tolerance=.005)); r = nest(p)
    path = tmp_path/'analysis.goknest.json'; save_project(p,path,r)
    source.unlink(); loaded, saved = load_project(path)
    assert loaded.fingerprint() == p.fingerprint(); assert saved == r
    out = export_dxf(loaded,saved,tmp_path/'cut')
    doc = ezdxf.readfile(out[0]); types = [e.dxftype() for e in doc.modelspace()]
    assert 'CIRCLE' in types and 'ARC' in types
    assert doc.units == 4
    assert set(['SHEET_BOUNDARY','PART_OUTER','PART_INNER','PART_LABELS','PART_REFERENCE']) <= {l.dxf.name for l in doc.layers}
    assert len(doc.modelspace().query('CIRCLE[layer=="PART_INNER"]')) == 1
    from nesting.report import export_excel,export_pdf
    export_excel(loaded,saved,tmp_path/'report.xlsx'); export_pdf(loaded,saved,tmp_path/'report.pdf')
    from openpyxl import load_workbook
    book = load_workbook(tmp_path/'report.xlsx'); assert 'Yerleşim 001' in book.sheetnames
    assert book['Yerleşim 001']._images; assert book['Parçalar'].max_row == 2; book.close()
    assert (tmp_path/'report.pdf').read_bytes().startswith(b'%PDF-')
    assert b'%%EOF' in (tmp_path/'report.pdf').read_bytes()[-1024:]


def test_invalid_or_stale_outputs_blocked_atomically(tmp_path):
    p = project(); r = nest(p); p.options.gap += 1
    with pytest.raises(ValueError,match='değişti'): export_dxf(p,r,tmp_path/'out')
    assert not (tmp_path/'out').exists()
    with pytest.raises(ValueError): save_project(p,tmp_path/'p.json',r)
    assert not (tmp_path/'p.json').exists()


def test_embedded_dxf_must_match_validated_geometry(tmp_path):
    path = drawing(tmp_path,lambda m,d:rectangle(m))
    p = project(read_dxf(path)); p.parts[0].geometry = __import__('shapely').geometry.mapping(box(0,0,60,30))
    r = nest(p)
    with pytest.raises(ValueError,match='eşleşmiyor'): export_dxf(p,r,tmp_path/'bad-export')
    assert not (tmp_path/'bad-export').exists()


def test_cancel_returns_only_verified_partial_and_progress():
    p = project([Part.create('A',box(0,0,10,10),quantity=50)])
    stop = threading.Event(); progress = []
    def callback(n,text):
        progress.append(n)
        if len(progress) == 5: stop.set()
    r = nest(p,stop=stop,progress=callback)
    assert r.stopped and r.unplaced
    assert len(r.placements)+len(r.unplaced) == 50
    assert validate_result(p,r,complete=False)
    assert progress[-1] == 100


@pytest.mark.parametrize('value',['nan','inf','-1','abc'])
def test_bad_numeric_input(value):
    with pytest.raises(ValueError): number(value)


def test_corrupt_project_is_rejected(tmp_path):
    p = project(); r = nest(p); f = tmp_path/'p.json'; save_project(p,f,r)
    data = json.loads(f.read_text()); data['project']['parts'][0]['quantity'] = 5
    f.write_text(json.dumps(data))
    with pytest.raises(ValueError): load_project(f)


def test_material_and_sheet_templates_can_be_saved_before_import(tmp_path):
    p = new_project(); p.sheets[0].width = 1234; p.materials.append(Material('custom','Özel kalite',7800))
    f = tmp_path/'template.goknest.json'; save_project(p,f)
    loaded,result = load_project(f)
    assert not loaded.parts and result is None
    assert loaded.sheets[0].width == 1234 and loaded.material('custom').density == 7800
