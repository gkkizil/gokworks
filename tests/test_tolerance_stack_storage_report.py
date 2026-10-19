from dataclasses import replace
from decimal import Decimal as D
import json
import threading

import pytest
from openpyxl import load_workbook

from test_tolerance_stack_math import example
from tolerance_stack import storage,report
from tolerance_stack.calculator import calculate,Cancelled


def test_project_roundtrip_preserves_sources_distribution_and_incomplete_work_but_invalidates_result(tmp_path):
    p=example();r=calculate(p)
    source=tmp_path/'Parça çizimi.SLDDRW';source.write_bytes(b'CAD snapshot')
    p.sources=[storage.fingerprint(source)]
    p.terms[0]=replace(p.terms[0],mean=D('49.99'),sigma=D('.02'))
    p.confirm();r=calculate(p)
    path=tmp_path/'a.gokstack.json';storage.save_project(p,path,result=r)
    loaded,archive=storage.load_project(path)
    assert loaded.fingerprint()==p.fingerprint()
    assert loaded.sources==p.sources
    assert loaded.terms[0].mean==D('49.99') and loaded.terms[0].sigma==D('.02')
    assert loaded.confirmed_digest=='' and archive['project_digest']==r.project_digest
    with pytest.raises(ValueError,match='doğrulayın'):calculate(loaded)
    loaded.confirm();assert calculate(loaded).minimum==r.minimum
    loaded.dimensions[0]=replace(loaded.dimensions[0],upper=None)
    storage.save_project(loaded,tmp_path/'incomplete.json')
    incomplete,_=storage.load_project(tmp_path/'incomplete.json')
    assert incomplete.dimensions[0].upper is None


def test_revision_check_hashes_content_and_deleted_sources(tmp_path):
    p=example();source=tmp_path/'part.SLDDRW';source.write_bytes(b'first')
    p.sources=[storage.fingerprint(source)];storage.check_sources(p)
    source.write_bytes(b'other')
    with pytest.raises(ValueError,match='değişti'):storage.check_sources(p)
    source.unlink()
    with pytest.raises(ValueError,match='bulunamadı'):storage.check_sources(p)


def test_atomic_writes_keep_existing_file_on_cancel_failure_or_race(tmp_path):
    path=tmp_path/'keep.json';path.write_text('original')
    with pytest.raises(FileExistsError):storage.atomic_write(path,lambda p:p.write_text('new'))
    def fail(p):p.write_text('partial');raise RuntimeError('disk full')
    with pytest.raises(RuntimeError):storage.atomic_write(path,fail,overwrite=True)
    assert path.read_text()=='original'
    stop=threading.Event()
    def cancel(p):p.write_text('partial');stop.set()
    with pytest.raises(Cancelled):storage.atomic_write(path,cancel,overwrite=True,stop=stop)
    assert path.read_text()=='original' and not list(tmp_path.glob('.gokworks-stack-*'))
    target=tmp_path/'new.json'
    def race(p):p.write_text('ours');target.write_text('theirs')
    with pytest.raises(FileExistsError):storage.atomic_write(target,race)
    assert target.read_text()=='theirs'


@pytest.mark.parametrize('mutation',[lambda d:d.update(version=99),lambda d:d.update(unit='rad'),
    lambda d:d['project'].update(dimensions=[dict(id='bad',name='x',nominal='NaN',lower='0',upper='0')]),
    lambda d:d['project'].update(terms=[dict(unknown='x')])])
def test_invalid_or_future_project_data_is_rejected(tmp_path,mutation):
    path=tmp_path/'project.json';storage.save_project(example(),path)
    data=json.loads(path.read_text());mutation(data);path.write_text(json.dumps(data))
    with pytest.raises(ValueError):storage.load_project(path)


def test_excel_contains_all_terms_assumptions_and_never_evaluates_cad_labels_as_formulas(tmp_path):
    p=example();p.dimensions[0]=replace(p.dimensions[0],name='=HYPERLINK("bad")');p.confirm();r=calculate(p)
    path=tmp_path/'rapor.xlsx';report.export_excel(p,r,path)
    wb=load_workbook(path,data_only=False)
    assert wb['Ölçü zinciri'].max_row==4
    assert wb['Ölçü zinciri']['B2'].data_type=='s'
    assert wb['Ölçü zinciri']['B2'].value.startswith('=HYPERLINK')
    assert wb['Ölçü zinciri']['G2'].value=='50,000000'
    values=[row for row in wb['Analiz sonucu'].values]
    assert ('Worst-Case minimum (mm)','0,020000') in values
    assert any('bağımsız' in str(v) for row in values for v in row)
    wb.close()
    p.acceptance_max=D('.5')
    with pytest.raises(ValueError,match='güncel değil'):report.export_excel(p,r,tmp_path/'stale.xlsx')
    assert not (tmp_path/'stale.xlsx').exists()


def test_real_unicode_pdf_is_complete_and_cancel_does_not_leave_a_report(qapp,tmp_path):
    p=example();p.name='Ölçü zinciri · İmalat';p.confirm();r=calculate(p)
    path=tmp_path/'rapor.pdf';report.export_pdf(p,r,path)
    raw=path.read_bytes()
    assert raw.startswith(b'%PDF-') and b'%%EOF' in raw[-1024:] and len(raw)>3000
    stop=threading.Event();stop.set()
    with pytest.raises(Cancelled):report.export_pdf(p,r,tmp_path/'cancelled.pdf',stop=stop)
    assert not (tmp_path/'cancelled.pdf').exists()
