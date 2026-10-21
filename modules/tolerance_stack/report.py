"""Unicode PDF and XLSX reports with source provenance and explicit assumptions."""

from html import escape
from pathlib import Path

from modules.tolerance_stack.calculator import check_cancel
from modules.tolerance_stack.models import fmt, serializable
from modules.tolerance_stack.storage import atomic_write, check_sources


def verdict(value):
    return 'Kriter girilmedi' if value is None else 'Uygun' if value else 'Uygun değil'


def assert_current(project,result,*,stop=None):
    if result is None or result.project_digest!=project.fingerprint() or project.confirmed_digest!=project.fingerprint():
        raise ValueError('Sonuç güncel değil; zinciri doğrulayıp yeniden hesaplayın.')
    check_sources(project,stop=stop)


def summary_rows(project,result):
    rows=[('Analiz',project.name),('Kritik ölçü',project.critical),('Zincir başlangıcı / bitişi',project.start+' → '+project.end),
        ('Eksen / ortak datum',project.axis+' / '+project.datum),('Nominal sonuç (mm)',fmt(result.nominal)),
        ('Worst-Case minimum (mm)',fmt(result.minimum)),('Worst-Case maksimum (mm)',fmt(result.maximum)),
        ('Worst-Case değerlendirmesi',verdict(result.worst_case_pass)),('Boşluk / sıkılık',result.kind),
        ('Dağılım merkezi (mm)',fmt(result.mean)),('RSS σ (mm)',fmt(result.sigma)),
        ('RSS minimum (mm)',fmt(result.rss_minimum)),('RSS maksimum (mm)',fmt(result.rss_maximum)),
        ('RSS değerlendirmesi',verdict(result.rss_pass)),
        ('Kabul alt sınırı (mm)','Yok' if project.acceptance_min is None else fmt(project.acceptance_min)),
        ('Kabul üst sınırı (mm)','Yok' if project.acceptance_max is None else fmt(project.acceptance_max)),
        ('Yöntem',project.method),('Varsayılan dağılım',project.distribution),('Normal kesme',str(project.truncate_normal)),
        ('Giriş / sonuç σ katsayısı',f'{project.input_sigma} / {project.output_sigma}'),
        ('Hesap zamanı (UTC)',result.timestamp),('GOKWORKS analiz kimliği',project.id),('Veri imzası',result.project_digest)]
    if result.monte_carlo:
        mc=result.monte_carlo
        for title,key in [('MC ortalama (mm)','mean'),('MC σ (mm)','sigma'),('MC örnek minimum (mm)','sample_minimum'),
            ('MC örnek maksimum (mm)','sample_maximum'),('Uygunluk oranı','acceptance_rate'),('Tahmini uygunsuzluk (PPM)','ppm'),
            ('Varsayımsal Cp','cp'),('Varsayımsal Cpk','cpk'),('Örnek sayısı','samples'),('Seed','seed')]:
            rows.append((title,mc[key]))
        rows += [('PPM %95 güven aralığı',str(mc['ppm_confidence_95'])),('Kabiliyet yorumu',mc['capability_basis'])]
        rows.extend((f'MC yüzdelik %{p} (mm)',fmt(v)) for p,v in mc['percentiles'].items())
    rows += [('Varsayım',text) for text in result.assumptions]
    return rows


def dimension_rows(project):
    dimensions={d.id:d for d in project.dimensions}
    for index,term in enumerate(project.terms,1):
        d=dimensions[term.dimension_id]
        yield (f'L{index}',d.name,term.direction,fmt(term.factor),term.start,term.end,fmt(d.nominal),fmt(d.lower),fmt(d.upper),
               fmt(d.minimum),fmt(d.maximum),d.tolerance_type,d.kind,term.relation,
               project.distribution if term.distribution=='inherit' else term.distribution,
               '' if term.mean is None else fmt(term.mean),'' if term.sigma is None else fmt(term.sigma),
               '' if term.mode is None else fmt(term.mode),d.source,d.sheet,d.view,d.model,d.configuration,d.api_name or d.id,d.provenance,d.note,
               'Kullanıcı doğruladı' if term.alignment_confirmed else 'İlişki açıklamasında' if term.relation else 'Manuel ölçü')


HEADERS=('Etiket','Ölçü','Yön','Katsayı','Başlangıç','Bitiş','Nominal mm','Alt sapma mm','Üst sapma mm',
         'Minimum mm','Maksimum mm','Tolerans tipi','Tür','Fonksiyonel ilişki','Dağılım','Merkez mm','σ mm','Tepe mm',
         'Kaynak çizim','Sayfa','Görünüm','Model','Konfigürasyon','Ölçü kimliği','Veri kaynağı','Not','Ortak doğrultu / referans kontrolü')


def export_excel(project,result,path,*,overwrite=False,stop=None):
    assert_current(project,result,stop=stop)
    def write(temporary):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        book=Workbook();summary=book.active;summary.title='Analiz sonucu'
        summary.append(('Alan','Değer'))
        for key,value in summary_rows(project,result):
            check_cancel(stop);summary.append((key,value))
        dimensions=book.create_sheet('Ölçü zinciri');dimensions.append(HEADERS)
        for row in dimension_rows(project):
            check_cancel(stop);dimensions.append(row)
        contributions=book.create_sheet('Tolerans katkıları');contributions.append(('Ölçü','Worst-Case %','Varyans %'))
        for c in result.contributions: contributions.append((c.name,c.worst_case_percent,c.variance_percent))
        sources=book.create_sheet('Kaynak kontrolü');sources.append(('Dosya','SHA256','Boyut','Değişiklik ns','Okunduğunda kayıt dışı değişiklik','Çizim notları'))
        for source in project.sources: sources.append((source.path,source.sha256,source.size,str(source.modified_ns),source.dirty,'\n'.join(source.notes)))
        if result.monte_carlo:
            histogram=book.create_sheet('MC histogramı');histogram.append(('Aralık alt mm','Aralık üst mm','Örnek sayısı'))
            h=result.monte_carlo['histogram'];width=(h['maximum']-h['minimum'])/len(h['counts'])
            for i,count in enumerate(h['counts']): histogram.append((h['minimum']+i*width,h['minimum']+(i+1)*width,count))
        for sheet in book:
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24466C')
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width=min(70,max(16,max(len(str(c.value or '')) for c in column)+2))
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    # CAD labels and filenames must remain text, never formulas.
                    if isinstance(cell.value,str): cell.data_type='s'
                    cell.alignment=Alignment(vertical='top',wrap_text=True)
        book.save(temporary);book.close()
        check=load_workbook(temporary,read_only=True,data_only=False)
        try:
            if check['Ölçü zinciri'].max_row!=len(project.terms)+1: raise RuntimeError('Excel satır doğrulaması başarısız.')
        finally: check.close()
    return atomic_write(path,write,overwrite=overwrite,stop=stop)


def export_pdf(project,result,path,*,overwrite=False,stop=None):
    assert_current(project,result,stop=stop)
    def write(temporary):
        from PySide6.QtCore import QRectF, QMarginsF
        from PySide6.QtGui import QPdfWriter,QPageSize,QPainter,QTextDocument,QPageLayout
        rows=summary_rows(project,result)
        html=['<h1>GOKWORKS · Tolerans Yığılması</h1><table border="1" cellpadding="5" width="100%">']
        html += ['<tr><td>'+escape(str(k))+'</td><td>'+escape(str(v))+'</td></tr>' for k,v in rows]
        html += ['</table><h2>Ölçü zinciri · mm</h2>']
        for row in dimension_rows(project):
            check_cancel(stop)
            html.append('<p><b>'+escape(f'{row[0]} · {row[1]} · yön {row[2]:+d} × {row[3]}')+'</b><br>'+escape(
                f'Nominal {row[6]} / alt {row[7]} / üst {row[8]} mm\n{row[18]} · {row[19]} · {row[20]} · {row[22]}\n{row[13]} · {row[14]} · {row[25]}\nDoğrultu kontrolü: {row[26]}').replace('\n','<br>')+'</p>')
        writer=QPdfWriter(str(temporary));writer.setResolution(96);writer.setPageSize(QPageSize(QPageSize.A4))
        writer.setPageMargins(QMarginsF(14,14,14,14),QPageLayout.Millimeter)
        writer.setTitle(project.name);writer.setCreator('GOKWORKS Tolerance Stack-Up 1.0')
        document=QTextDocument();document.setDefaultStyleSheet('body {font-family: DejaVu Sans; font-size:10pt;} h1 {color:#24466C;} td {font-size:9pt;}')
        document.setHtml(''.join(html));document.setTextWidth(writer.width())
        painter=QPainter(writer)
        try:
            height=writer.height()
            for offset in range(0,int(document.size().height())+1,height):
                check_cancel(stop)
                if offset and not writer.newPage(): raise RuntimeError('PDF sayfası oluşturulamadı.')
                painter.save();painter.setClipRect(QRectF(0,0,writer.width(),height));painter.translate(0,-offset)
                document.drawContents(painter,QRectF(0,offset,writer.width(),height));painter.restore()
        finally:
            painter.end()
        del writer
        content=temporary.read_bytes()
        if not content.startswith(b'%PDF-') or b'%%EOF' not in content[-1024:]: raise RuntimeError('PDF doğrulanamadı.')
    return atomic_write(path,write,overwrite=overwrite,stop=stop)
