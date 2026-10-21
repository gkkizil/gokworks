"""Atomic Excel/PDF exports with live accounting and every sheet's actual layout."""

from io import BytesIO

from modules.tolerance_stack.storage import atomic_write
from modules.nesting.accounting import calculate
from modules.nesting.dxf_import import check_cancel
from modules.nesting.render import sheet_png, draw_sheet


def summary_rows(project, result, totals):
    return [('Proje', project.name), ('Rapor zamanı (UTC)', result.timestamp),
            ('Motor', result.engine+' '+result.engine_version), ('Sonuç', 'Bağımsız geometri kontrolü geçti; en iyi bulunan sonuç, optimum garantisi yok.'),
            ('Parça adedi', totals['part_count']), ('Kullanılan levha', totals['sheet_count']),
            ('Brüt ağırlık kg', totals['gross_weight']), ('Net ağırlık kg', totals['net_weight']),
            ('Fire ağırlığı kg', totals['waste_weight']), ('Fire %', 100-totals['efficiency']),
            ('Alan verimi %', totals['efficiency']), ('Malzeme maliyeti TL', totals['material_cost']),
            ('Net malzeme karşılığı TL', totals['net_material_value']), ('Fire malzeme karşılığı TL', totals['waste_material_value']),
            ('Kesim uzunluğu m (nominal kontur)', totals['cut_length_m']), ('Kontur / piercing adedi', totals['piercing_count']),
            ('Kesim TL', totals['cut_cost']), ('Piercing TL', totals['piercing_cost']), ('Hazırlık TL', totals['setup']),
            ('Sabit gider TL', totals['fixed']), ('Kâr TL', totals['profit']), ('KDV TL', totals['vat']), ('Genel toplam TL', totals['total']),
            ('Kenar payı mm', project.options.margin), ('Bitmiş parça aralığı mm', project.options.gap), ('Kerf mm', project.options.kerf),
            ('Kerf açıklaması', 'Nominal kontur; CAM kesim yolu telafisi uygulanmadı. Mesafe max(aralık, kerf) + eğri hata payları.'),
            ('Maliyet dağıtımı', 'Her levhanın satın alma bedeli, yerleşen parçaların net alanı oranında dağıtılır; toplam kuruşu kuruşuna korunur.'),
            ('Kalıntılar', 'Yerleşen parçalar dışındaki alan fire/kalan alanıdır. Deliklere yerleşen parçalar net alana dahildir; hurda kredisi düşülmedi.'),
            ('Süre tahmini', 'Hesaplanmaz; kesim ve piercing bedelleri girilen birim fiyatlardan gelir.'),
            ('Hadde yönü', '0°/180° ile sınırlı' if project.options.grain_locked else 'Parça bazlı / serbest'),
            ('Aynalama', 'Kapalı'), ('Delik içi yerleşim', 'Otomatik'),
            ('Kesim sırası', 'Önce iç içe küçük konturlar; DXF yanındaki KESIM-SIRASI.csv CAM içinde uygulanmalıdır.'),
            ('Hesap birimi', 'mm; kg; TL')]


PART_COLUMNS = [('Parça kodu','code'),('Açıklama','description'),('Örnek','instance'),('Malzeme','material'),
                ('Kalınlık mm','thickness'),('Net alan mm²','area'),('Tek parça kg','weight'),('Levha','sheet'),
                ('Dağıtılmış maliyet TL','allocation'),('Dış çevre mm','outer_perimeter'),('Delik çevresi mm','inner_perimeter'),
                ('Kaynak','source'),('SHA256','sha256')]
SHEET_COLUMNS = [('Plaka','plate'),('Sac adı','name'),('En mm','width'),('Boy mm','height'),('Kalınlık mm','thickness'),
                 ('Malzeme','material'),('Yoğunluk kg/m³','density'),('Brüt kg','gross_weight'),('Net kg','net_weight'),
                 ('Fire kg','waste_weight'),('Verim %','efficiency'),('Malzeme TL','material_cost'),('Fiyat modeli','pricing'),('Birim fiyat','price')]


def export_excel(project, result, path, *, overwrite=False, stop=None):
    totals = calculate(project, result, stop=stop)
    def write(temporary):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.drawing.image import Image
        book = Workbook(); summary = book.active; summary.title = 'Özet'
        summary.append(['GOKWORKS · Nesting ve sac maliyeti', 'Değer'])
        for row in summary_rows(project, result, totals): summary.append(row)
        overview = book.create_sheet('Parça özeti')
        overview.append(['Parça kodu','Açıklama','Malzeme','Kalınlık mm','Adet','Tek parça kg','Toplam kg','Net alan mm²','X mm','Y mm','Levhalar','Dağıtılmış maliyet TL'])
        for p in totals['part_summary']:
            overview.append([p[k] for k in ('code','description','material','thickness','quantity','weight','total_weight','area','width','height','sheets','allocation')])
        for title, columns, rows in [('Parçalar', PART_COLUMNS, totals['parts']), ('Saclar', SHEET_COLUMNS, totals['sheets'])]:
            sheet = book.create_sheet(title); sheet.append([c[0] for c in columns])
            for row in rows:
                check_cancel(stop); sheet.append([row[c[1]] for c in columns])
        materials = book.create_sheet('Malzemeler'); materials.append(['Malzeme', 'Yoğunluk kg/m³ (yaklaşık)', 'Kaynak / kullanıcı tanımı'])
        for m in project.materials: materials.append([m.name, m.density, m.source])
        sources = book.create_sheet('Geometri bilgisi'); sources.append(['Parça', 'Giriş birimi', 'Eğri toleransı mm', 'Uç toleransı mm', 'Gerçek uç hareketi mm', 'Not'])
        for p in project.parts: sources.append([p.code, p.source_unit, p.curve_tolerance, p.join_tolerance, p.endpoint_error, '\n'.join(p.notes)])
        for sheet in book:
            sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='24466C')
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(65, max(18, max(len(str(c.value or '')) for c in column)+2))
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str): cell.data_type = 's'
                    elif isinstance(cell.value, float): cell.number_format = '0.0000'
        for index, s in enumerate(totals['sheets'], 1):
            check_cancel(stop)
            sheet = book.create_sheet(f'Yerleşim {index:03d}'); sheet['A1'] = f"{s['plate']}. plaka · {s['name']} · {s['material']} · {s['thickness']} mm"
            sheet['A1'].data_type = 's'
            image = Image(BytesIO(sheet_png(project, result, (s['id'], s['number'])))); image.width = 900; image.height = 630
            sheet.add_image(image, 'A3')
        book.save(temporary); book.close()
        check = load_workbook(temporary, read_only=True)
        try:
            if check['Parçalar'].max_row != totals['part_count']+1: raise ValueError('Excel satır sayısı doğrulanamadı.')
        finally: check.close()
    return atomic_write(path, write, overwrite=overwrite, stop=stop)


def export_pdf(project, result, path, *, overwrite=False, stop=None):
    totals = calculate(project, result, stop=stop)
    def write(temporary):
        from PySide6.QtCore import QRectF, QMarginsF, QSizeF, Qt
        from PySide6.QtGui import (QAbstractTextDocumentLayout, QPdfWriter, QPainter, QPageSize,
                                  QPageLayout, QTextDocument, QFont, QColor, QPalette)
        from modules.nesting.pdf_report import report_html
        writer = QPdfWriter(str(temporary)); writer.setResolution(96); writer.setPageSize(QPageSize(QPageSize.A4))
        writer.setPageMargins(QMarginsF(12,12,12,12), QPageLayout.Millimeter)
        writer.setTitle(project.name); writer.setCreator('GOKWORKS Nesting')
        # Give Qt a real page size before layout so tables wrap and paginate as
        # tables, rather than cutting a single long document at arbitrary rows.
        footer_height = 28
        page_height = writer.height()-footer_height
        document = QTextDocument(); document.setDocumentMargin(0)
        body_font = QFont('Arial'); body_font.setPointSizeF(10.5)
        document.setDefaultFont(body_font)
        document.documentLayout().setPaintDevice(writer)
        document.setPageSize(QSizeF(writer.width(), page_height))
        document.setHtml(report_html(project, result, totals))
        context = QAbstractTextDocumentLayout.PaintContext()
        for role in (QPalette.Text, QPalette.WindowText, QPalette.ButtonText):
            context.palette.setColor(role, QColor('#17212b'))
        pages = document.pageCount()+len(totals['sheets'])
        painter = QPainter(writer)
        try:
            def white_page():
                # Paint the actual paper, including margins. QtPdf renders
                # unpainted PDF regions as transparent; PDF consumers should
                # always receive a white production report.
                layout = writer.pageLayout()
                full = layout.fullRectPixels(writer.resolution())
                area = layout.paintRectPixels(writer.resolution())
                painter.fillRect(QRectF(-area.x(), -area.y(), full.width(), full.height()), QColor('#ffffff'))

            def footer(number):
                painter.setPen(QColor('#334155')); painter.setFont(QFont('Arial', 9))
                painter.drawText(QRectF(0, writer.height()-22, writer.width(), 22),
                                 Qt.AlignVCenter | Qt.AlignRight, f'GOKWORKS · {number} / {pages}')

            for index in range(document.pageCount()):
                check_cancel(stop)
                if index and not writer.newPage(): raise RuntimeError('PDF sayfası oluşturulamadı.')
                white_page()
                offset = index*page_height
                painter.save(); painter.setClipRect(QRectF(0,0,writer.width(),page_height)); painter.translate(0,-offset)
                context.clip = QRectF(0,offset,writer.width(),page_height)
                document.documentLayout().draw(painter, context); painter.restore()
                footer(index+1)
            for index, s in enumerate(totals['sheets'], document.pageCount()+1):
                check_cancel(stop)
                if not writer.newPage(): raise RuntimeError('PDF yerleşim sayfası oluşturulamadı.')
                white_page()
                painter.setPen(QColor('#17212b')); painter.setFont(QFont('Arial', 12, QFont.Bold))
                painter.drawText(QRectF(0,0,writer.width(),60), Qt.TextWordWrap,
                                 f"{s['plate']}. plaka · {s['name']} · {s['material']} · {s['width']:g} × {s['height']:g} × {s['thickness']:g} mm")
                painter.setFont(body_font)
                painter.drawText(QRectF(0,65,writer.width(),85), Qt.TextWordWrap,
                                 f"Brüt {s['gross_weight']:.5f} kg · Net {s['net_weight']:.5f} kg · Fire {s['waste_weight']:.5f} kg\n"
                                 f"Verim %{s['efficiency']:.2f} · Malzeme {s['material_cost']:.2f} TL · {s['pricing']} {s['price']:g}\n"
                                 'Nominal kesim konturları; kerf telafisi CAM/tezgâhta uygulanır.')
                draw_sheet(painter, QRectF(0,150,writer.width(),writer.height()-190), project, result, (s['id'],s['number']))
                footer(index)
        finally: painter.end()
        del writer
        content = temporary.read_bytes()
        if not content.startswith(b'%PDF-') or b'%%EOF' not in content[-1024:]: raise ValueError('PDF doğrulanamadı.')
    return atomic_write(path, write, overwrite=overwrite, stop=stop)
