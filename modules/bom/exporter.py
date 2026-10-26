"""Read-only, instance-aware assembly traversal and atomic XLSX export."""

from dataclasses import dataclass, replace
from datetime import datetime, timedelta, timezone
import os
import math
from pathlib import Path, PureWindowsPath
import re
from tempfile import NamedTemporaryFile

from modules.dxf_dwg import exporter as sw
from solidworks.connection import connect_to_solidworks

pythoncom, win32com = sw.pythoncom, sw.win32com


class Cancelled(Exception):
    pass


def _check_cancel(stop_event):
    if stop_event is not None and stop_event.is_set():
        raise Cancelled()


def name_key(value):
    """Turkish alphabet, numeric segments: Parça 2 precedes Parça 10."""
    alphabet = 'abcçdefgğhıijklmnoöprsştuüvyz'
    text = str(value).replace('I', 'ı').replace('İ', 'i').casefold()
    return tuple((0, int(chunk)) if chunk.isdecimal() else
                 (1, tuple(alphabet.index(c) if c in alphabet else ord(c)+100 for c in chunk))
                 for chunk in re.split(r'(\d+)', text) if chunk)


@dataclass(frozen=True)
class Part:
    name: str
    path: str
    configuration: str
    quantity: int
    mass_kg: float | None = None
    total_mass_kg: float | None = None
    mass_note: str = ''


@dataclass(frozen=True)
class Node:
    number: str
    parent: str
    level: int
    name: str
    instance: str
    kind: str
    path: str
    configuration: str
    mass_kg: float | None = None
    mass_note: str = ''


@dataclass(frozen=True)
class Result:
    assembly: str
    configuration: str
    parts: tuple
    tree: tuple
    suppressed: int
    excluded: int
    include_excluded: bool

    @property
    def total_parts(self):
        return sum(part.quantity for part in self.parts)

    @property
    def total_mass_kg(self):
        # An incomplete sum must not look like the weight of the whole BOM.
        if any(part.total_mass_kg is None for part in self.parts):
            return None
        return math.fsum(part.total_mass_kg for part in self.parts)

    @property
    def missing_mass_parts(self):
        return sum(node.kind == 'Parça' and node.mass_kg is None for node in self.tree)


class _MassReader:
    """Read live occurrence mass without selecting or switching documents.

    GetModelDoc2 can point at a different active part configuration. Use the
    assembly's IMassProperty2 with exactly one component instead. Reuse the
    calculator, but never cache by file/config: assembly cuts can make two
    occurrences of the same part weigh differently.
    """
    def __init__(self, assembly):
        self.properties = None
        self.error = ''
        try:
            extension = sw._value(assembly, 'Extension')
            properties = sw._value(extension, 'CreateMassProperty2')
            if properties is None:
                raise RuntimeError('SolidWorks ağırlık hesaplayıcısını oluşturamadı.')
            properties.UseSystemUnits = True  # meters, radians, kilograms
            properties.IncludeHiddenBodiesOrComponents = True
            self.properties = properties
        except Exception as error:
            self.error = f'SolidWorks ağırlık API’sine erişilemiyor: {error}'

    def read(self, component, stop_event=None):
        _check_cancel(stop_event)
        if self.properties is None:
            return None, self.error
        try:
            properties = self.properties
            items = (component,)
            if hasattr(component, '_oleobj_'):
                # Preserve IDispatch entries in the SAFEARRAY on pywin32;
                # plain Python doubles use the same logical one-item list.
                items = win32com.client.VARIANT(pythoncom.VT_ARRAY | pythoncom.VT_DISPATCH, items)
            properties.SelectedItems = items
            selected = sw._value(properties, 'SelectedItems')
            if not isinstance(selected, (tuple, list)) or len(selected) != 1:
                raise RuntimeError('Ağırlık hesabı tek bileşenle sınırlandırılamadı.')
            for member in ('Name2', 'GetPathName', 'ReferencedConfiguration'):
                if sw._value(selected[0], member) != sw._value(component, member):
                    raise RuntimeError('Ağırlık hesabında farklı bir bileşen seçildi.')
            if not sw._value(properties, 'UseSystemUnits'):
                raise RuntimeError('Ağırlık birimi kilogram olarak ayarlanamadı.')
            if not sw._value(properties, 'Recalculate'):
                raise RuntimeError('SolidWorks ağırlığı hesaplayamadı.')
            value = sw._value(properties, 'Mass')
            if isinstance(value, bool) or value is None:
                raise ValueError('Geçerli bir ağırlık değeri alınamadı.')
            mass = float(value)
            if not math.isfinite(mass) or mass < 0:
                raise ValueError('Geçerli bir ağırlık değeri alınamadı.')
        except Exception as error:
            return None, f'Ağırlık okunamadı: {error}'
        _check_cancel(stop_event)
        return mass, ''


def _weighted_part(original, occurrences):
    masses = [mass for mass, _ in occurrences]
    if any(mass is None for mass in masses):
        note = next(note for mass, note in occurrences if mass is None)
        return Part(*original, len(masses), mass_note=note)
    unit = masses[0] if all(math.isclose(mass, masses[0], rel_tol=1e-9, abs_tol=1e-12)
                            for mass in masses) else None
    note = '' if unit is not None else 'Bileşen örneklerinin ağırlıkları farklı.'
    return Part(*original, len(masses), unit, math.fsum(masses), note)


def _roll_up_mass(tree):
    """Assembly rows sum included leaf occurrences; never add them twice."""
    totals = {}
    weighted = []
    for node in reversed(tree):
        if node.kind == 'Montaj':
            children = totals.get(node.number, [])
            mass = None if any(value is None for value in children) else math.fsum(children)
            node = replace(node, mass_kg=mass,
                           mass_note='Alt parçalardan bazılarının ağırlığı okunamadı.' if mass is None else '')
        weighted.append(node)
        totals.setdefault(node.parent, []).append(node.mass_kg)
    return tuple(reversed(weighted))


def _assembly(app):
    model = sw._value(app, 'ActiveDoc')
    if model is None or sw._value(model, 'GetType') != sw.ASSEMBLY:
        raise RuntimeError("BOM için SolidWorks'te ana montajı etkinleştirin.")
    api = model
    if not hasattr(api, 'ResolveAllLightWeightComponents'):
        for interface in ('IAssemblyDoc', 'AssemblyDoc'):
            try:
                api = win32com.client.CastTo(model, interface)
                if hasattr(api, 'ResolveAllLightWeightComponents'):
                    break
            except Exception:
                continue
        else:
            raise RuntimeError('SolidWorks montaj API’sine erişilemiyor.')
    return model, api


def _excluded(component):
    if hasattr(component, 'GetExcludeFromBOM2'):
        values = component.GetExcludeFromBOM2(1, None)  # swThisConfiguration
        if not isinstance(values, (tuple, list)) or len(values) != 1:
            raise RuntimeError('Bileşenin aktif konfigürasyondaki BOM dışı durumu okunamadı.')
        return bool(values[0])
    # Documented legacy property for older dispatch interfaces.
    return bool(sw._value(component, 'ExcludeFromBOM'))


def collect(app, stop_event=None, *, include_excluded=False, progress=None):
    """Count each physical occurrence; never deduplicate subassembly traversal.

    Hidden components are included; suppressed branches are omitted. Excluded
    BOM components and their descendants are omitted unless explicitly enabled.
    No feature edits, configuration switches, file saves or document closes.
    """
    _check_cancel(stop_event)
    assembly, api = _assembly(app)
    resolved = api.ResolveAllLightWeightComponents(True)
    if resolved not in (0, 2):
        raise RuntimeError(f'Montaj bileşenleri çözümlenemedi. API hata kodu: {resolved}')
    config = assembly.ConfigurationManager.ActiveConfiguration
    if hasattr(config, 'IsSpeedPak') and sw._value(config, 'IsSpeedPak'):
        raise RuntimeError('SpeedPak yerine tam montaj konfigürasyonunu etkinleştirin.')
    root = config.GetRootComponent3(True)
    if root is None:
        raise RuntimeError('Montaj bileşen ağacı okunamadı.')
    configuration = str(sw._value(config, 'Name'))
    assembly_name = str(sw._document_name(assembly))
    masses, originals, tree = {}, {}, []
    reader = _MassReader(assembly)
    suppressed = excluded = visited = 0
    root_key = sw._part_key(assembly_name, configuration)
    # Sibling records are read once and sorted before entering the traversal.
    def children(component):
        nonlocal suppressed, excluded
        records = []
        for child in sw._value(component, 'GetChildren') or ():
            _check_cancel(stop_event)
            state = sw._value(child, 'GetSuppression')
            if state == 0:
                suppressed += 1
                continue
            if state not in (2, 3):
                raise RuntimeError('Çözümlenmemiş bileşen var; eksik BOM oluşturulmadı.')
            if not include_excluded and _excluded(child):
                excluded += 1
                continue
            path = str(sw._value(child, 'GetPathName') or '')
            ref = str(sw._value(child, 'ReferencedConfiguration') or '')
            if not path or not ref:
                raise RuntimeError('Bir bileşenin dosya yolu veya konfigürasyonu okunamadı.')
            filename = PureWindowsPath(path)
            kind = {'.sldprt': 'Parça', '.sldasm': 'Montaj'}.get(filename.suffix.lower())
            if kind is None:
                raise RuntimeError(f'Tanınmayan bileşen dosyası: {path}')
            instance = str(sw._value(child, 'Name2'))
            records.append((child, path, ref, filename.stem, instance, kind))
        return sorted(records, key=lambda r: (name_key(r[3]), name_key(r[2]), r[1], name_key(r[4])))

    stack = [(record, str(index), '', 1, frozenset((root_key,)))
             for index, record in reversed(list(enumerate(children(root), 1)))]
    while stack:
        _check_cancel(stop_event)
        (component, path, ref, name, instance, kind), number, parent, level, ancestors = stack.pop()
        key = sw._part_key(path, ref)
        if kind == 'Parça':
            mass, note = reader.read(component, stop_event)
            tree.append(Node(number, parent, level, name, instance, kind, path, ref, mass, note))
            masses.setdefault(key, []).append((mass, note))
            originals.setdefault(key, (name, path, ref))
        else:
            tree.append(Node(number, parent, level, name, instance, kind, path, ref))
            if key in ancestors:
                raise RuntimeError(f'Montaj ağacında döngü bulundu: {path}')
            if sw._value(component, 'GetModelDoc2') is None:
                raise RuntimeError(f'Alt montaj yüklenmemiş: {path}')
            if hasattr(component, 'IsSpeedPak') and sw._value(component, 'IsSpeedPak'):
                raise RuntimeError(f'SpeedPak alt montajın tam ağacını açın: {path}')
            stack.extend((record, f'{number}.{index}', number, level+1, ancestors | {key})
                         for index, record in reversed(list(enumerate(children(component), 1))))
        visited += 1
        if progress and visited % 100 == 0:
            progress(f'Montaj taranıyor · {visited} bileşenin parça ve ağırlık bilgileri okundu')
    _check_cancel(stop_event)
    parts = tuple(sorted((_weighted_part(originals[key], values) for key, values in masses.items()),
                         key=lambda p: (name_key(p.name), name_key(p.configuration), p.path)))
    return Result(assembly_name, configuration, parts, _roll_up_mass(tree), suppressed, excluded, include_excluded)


def export_excel(result, destination, stop_event=None, *, previews=None):
    """Write literal strings and numeric quantities; replace only after success."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    _check_cancel(stop_event)
    destination = Path(destination)
    if destination.suffix.lower() != '.xlsx':
        raise ValueError('Excel dosyası .xlsx uzantılı olmalı.')
    if max(len(result.parts), len(result.tree)) > 1048572:
        raise ValueError('Excel satır sınırı aşıldı; dosya oluşturulmadı.')
    book = Workbook()
    book.remove(book.active)
    created = datetime.now(timezone.utc).astimezone(timezone(timedelta(hours=3)))
    def sheet(title, headers, rows, widths):
        ws = book.create_sheet(title)
        ws.append([title])
        ws.append([result.assembly, result.configuration, created.isoformat(timespec='seconds')])
        ws.append([f'Bastırılmış dal: {result.suppressed}', f'BOM dışı dal: {result.excluded}',
                   'BOM dışında bırakılanlar dahil' if result.include_excluded else 'BOM dışında bırakılanlar hariç'])
        ws.append(headers)
        for values in rows:
            _check_cancel(stop_event)
            ws.append(values)
        # Names and paths are user-controlled strings, never Excel formulas.
        for row in ws:
            _check_cancel(stop_event)
            for cell in row:
                _check_cancel(stop_event)
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top')
        for cell in ws[4]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='177C65')
        ws['A1'].font = Font(size=16, bold=True, color='177C65')
        ws.freeze_panes = 'A5'
        ws.auto_filter.ref = f'A4:{get_column_letter(len(headers))}{max(4, ws.max_row)}'
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        return ws
    temporary = None
    try:
        image_headers = ['Kaynak parça · izometrik'] if previews is not None else []
        image_widths = (28,) if previews is not None else ()
        part_sheet = sheet('Parça listesi', ['Sıra', 'Parça adı', 'Adet', 'Konfigürasyon', 'Dosya adı', 'Dosya yolu',
                          'Birim ağırlık (kg)', 'Toplam ağırlık (kg)', 'Ağırlık durumu']+image_headers,
              ([i, p.name, p.quantity, p.configuration, PureWindowsPath(p.path).name, p.path,
                p.mass_kg if p.mass_kg is not None else 'Değişken' if p.total_mass_kg is not None else 'Okunamadı',
                p.total_mass_kg if p.total_mass_kg is not None else 'Okunamadı',
                p.mass_note or ('Okunamadı' if p.total_mass_kg is None else 'Okundu')]
               for i, p in enumerate(result.parts, 1)), (8, 34, 10, 28, 40, 80, 22, 22, 65)+image_widths)
        tree_sheet = sheet('Ürün ağacı', ['Pozisyon', 'Üst pozisyon', 'Seviye', 'Parça / montaj adı',
                            'Bileşen örneği', 'Tür', 'Adet', 'Konfigürasyon', 'Dosya yolu',
                            'Ağırlık (kg)', 'Ağırlık durumu']+image_headers,
              ([n.number, n.parent, n.level, n.name, n.instance, n.kind, 1, n.configuration, n.path,
                n.mass_kg if n.mass_kg is not None else 'Okunamadı',
                n.mass_note or ('Okunamadı' if n.mass_kg is None else
                               'Listelenen alt parçaların toplamı' if n.kind == 'Montaj' else 'Okundu')]
               for n in result.tree), (16, 16, 10, 34, 40, 12, 10, 28, 80, 22, 65)+image_widths)
        for ws, columns in ((part_sheet, (7, 8)), (tree_sheet, (10,))):
            for row in ws.iter_rows(min_row=5):
                _check_cancel(stop_event)
                for column in columns:
                    row[column-1].number_format = '0.000000'
            # Metadata stays outside the filterable rows. Use the flat list to
            # avoid summing both subassemblies and their children.
            ws.cell(1, 4, 'BOM toplam ağırlığı (kg)')
            total_cell = ws.cell(1, 5, result.total_mass_kg if result.total_mass_kg is not None else 'Eksik ağırlık bilgisi')
            total_cell.number_format = '0.000000'
            ws.cell(2, 4, 'Alt montaj ağırlıkları, listelenen alt parçaların toplamıdır.')
            ws.cell(3, 4, f'Ağırlığı okunamayan parça adedi: {result.missing_mass_parts}')
        tree_sheet.sheet_properties.outlinePr.summaryBelow = False
        for row, node in enumerate(result.tree, 5):
            _check_cancel(stop_event)
            tree_sheet.row_dimensions[row].outlineLevel = min(node.level-1, 7)
            tree_sheet.cell(row, 4).alignment = Alignment(indent=min(node.level-1, 250))
            if node.kind == 'Montaj':
                for cell in tree_sheet[row]:
                    cell.font = Font(bold=True, color='177C65')
                    cell.fill = PatternFill('solid', fgColor='E7F4EE')
        if previews is not None:
            from io import BytesIO
            from openpyxl.drawing.image import Image
            from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
            from openpyxl.drawing.xdr import XDRPositiveSize2D
            from openpyxl.utils.units import pixels_to_EMU
            for ws, records, column in ((part_sheet, result.parts, 10), (tree_sheet, result.tree, 12)):
                ws.cell(2, column, 'Dosya ve konfigürasyona göre izometrik görünüş')
                ws.cell(3, column, 'Kaynak parça geometrisi; montaj bağlamındaki kesimler dahil değildir.')
                ws.cell(3, column).alignment = Alignment(wrap_text=True, vertical='top')
                for row, item in enumerate(records, 5):
                    _check_cancel(stop_event)
                    cell = ws.cell(row, column)
                    if getattr(item, 'kind', 'Parça') != 'Parça':
                        cell.value = 'Alt montaj'
                        continue
                    preview = previews.get(sw._part_key(item.path, item.configuration))
                    if preview is None or not preview.png:
                        cell.value = preview.note if preview is not None else 'Görsel alınamadı'
                        cell.data_type = 's'
                        cell.alignment = Alignment(wrap_text=True, vertical='center')
                    else:
                        image = Image(BytesIO(preview.png))
                        image.width, image.height = 160, 121
                        image.anchor = OneCellAnchor(
                            _from=AnchorMarker(col=column-1, row=row-1,
                                               colOff=pixels_to_EMU(8), rowOff=pixels_to_EMU(8)),
                            ext=XDRPositiveSize2D(pixels_to_EMU(160), pixels_to_EMU(121)))
                        ws.add_image(image)
                    ws.row_dimensions[row].height = 108
        _check_cancel(stop_event)
        with NamedTemporaryFile(dir=destination.parent, prefix='.gokworks-bom-', suffix='.xlsx', delete=False) as stream:
            temporary = Path(stream.name)
        book.save(temporary)
        _check_cancel(stop_event)
        os.replace(temporary, destination)
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def export_worker(result, destination, events, stop_event):
    """Capture CAD images and write the complete BOM in the existing worker."""
    from .images import collect_previews
    try:
        _check_cancel(stop_event)
        previews = collect_previews(result, stop_event,
                                    progress=lambda text: events.put(('bom_phase', text)))
        events.put(('bom_phase', 'İzometrik görseller Excel dosyasına yerleştiriliyor…'))
        export_excel(result, destination, stop_event, previews=previews)
        images = sum(bool(previews.get(sw._part_key(part.path, part.configuration))
                          and previews[sw._part_key(part.path, part.configuration)].png) for part in result.parts)
        outcome = dict(state='exported', path=str(destination), images=images,
                       missing_images=len(result.parts)-images)
    except Cancelled:
        outcome = dict(state='cancelled')
    except Exception as error:
        outcome = dict(state='error', error=str(error))
    events.put(('bom_done', outcome))


def scan_worker(events, stop_event, *, include_excluded=False):
    initialized = False
    try:
        _check_cancel(stop_event)
        if pythoncom is None or win32com is None:
            raise RuntimeError('Montaj taraması Windows, pywin32 ve açık SolidWorks oturumu gerektirir.')
        pythoncom.CoInitialize()
        initialized = True
        result = collect(connect_to_solidworks(), stop_event, include_excluded=include_excluded,
                         progress=lambda text: events.put(('bom_phase', text)))
        outcome = dict(state='completed', result=result)
    except Cancelled:
        outcome = dict(state='cancelled')
    except Exception as error:
        outcome = dict(state='error', error=str(error))
    finally:
        if initialized:
            try:
                pythoncom.CoUninitialize()
            except Exception as error:
                outcome = dict(state='error', error=f'SolidWorks bağlantısı temizlenemedi: {error}')
        events.put(('bom_done', outcome))
