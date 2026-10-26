"""Real BMP/PNG/XLSX with mocked CAD calls; no licensed SOLIDWORKS acceptance."""

from dataclasses import replace
from io import BytesIO
from pathlib import Path
from queue import Queue
from types import SimpleNamespace as NS
from threading import Event
import math
import zipfile

from PIL import Image, ImageDraw
from openpyxl import load_workbook
import pytest

from modules.bom import exporter as bom, images
from modules.cad_converter import api
from modules.cad_converter.tests.fakes import App, C


IDENTITY = (1., 0., 0., 0., 1., 0., 0., 0., 1., 0., 0., 0., 1., 0., 0., 0.)
ISO = (math.sqrt(.5), -math.sqrt(1/6), math.sqrt(1/3),
       0., math.sqrt(2/3), math.sqrt(1/3),
       -math.sqrt(.5), -math.sqrt(1/6), math.sqrt(1/3), 0., 0., 0., 1., 0., 0., 0.)


class Vector:
    def __init__(self, values): self.ArrayData = tuple(values)
    def MultiplyTransform(self, transform):
        a = transform.ArrayData
        return Vector(tuple(a[12]*sum(self.ArrayData[i]*a[3*i+j] for i in range(3)) for j in range(3)))


def transpose_rotation(values):
    return tuple(values[3*j+i] for i in range(3) for j in range(3))+tuple(values[9:])


class View:
    """Orientation/camera state differs from the documented model-to-pixel map."""
    def __init__(self):
        self.Orientation3 = NS(ArrayData=IDENTITY)
        self.Translation3 = Vector((.1, -.2, 0))
        self.Scale2, self.perspective, self.Camera = 1.7, True, None

    @property
    def Transform(self):
        values = list(transpose_rotation(self.Orientation3.ArrayData))
        # Pixel Y points down; the view scale is part of the full mapping.
        for i in range(3): values[3*i+1] *= -1
        values[12] = self.Scale2
        return NS(ArrayData=tuple(values))

    def HasPerspective(self): return self.perspective
    def RemovePerspective(self): self.perspective = False
    def AddPerspective(self): self.perspective = True


@pytest.fixture
def native(tmp_path, monkeypatch):
    app = App()
    model = app.add(tmp_path/'Parça.SLDPRT', ('Original', 'Thin', 'Thick'), opened=True)
    previous = app.add(tmp_path/'Assembly.SLDASM', opened=True)
    app.ActiveDoc = previous
    view = View()
    def doubles(values):
        assert values.varianttype == 0x2005  # native MathUtility requires SAFEARRAY(double)
        return tuple(values.value)
    utility = NS(CreateTransform=lambda values: NS(ArrayData=doubles(values)),
                 CreateVector=lambda values: Vector(doubles(values)))
    app.GetMathUtility = lambda: utility
    model.ActiveView = view
    model.Visible = False
    model.GraphicsRedraw2 = lambda: None
    captures = []
    def show_view(name, enum):
        assert name == '' and enum == 7 and app.ActiveDoc is model and model.Visible
        view.Orientation3 = NS(ArrayData=transpose_rotation(ISO))
    model.ShowNamedView2 = show_view
    def zoom():
        view.Scale2 = 3.
        view.Translation3 = Vector((0, 0, 0))
    model.ViewZoomtofit2 = zoom
    def bitmap(path, width, height):
        assert (width, height) == (512, 384) and not view.perspective
        captures.append((model.path, model.active, Path(path)))
        with Image.new('RGB', (width, height), 'white') as image:
            ImageDraw.Draw(image).rectangle((80, 70, 430, 310),
                                           fill='blue' if model.active == 'Thin' else 'green')
            image.save(path, format='BMP')
        return True
    model.SaveBMP = bitmap
    real_switch = model.ShowConfiguration2
    def switch(name):
        answer = real_switch(name)
        # A configuration switch can itself change the view: snapshot before it.
        view.Orientation3 = NS(ArrayData=IDENTITY)
        view.Scale2 = 2.5
        return answer
    model.ShowConfiguration2 = switch
    monkeypatch.setattr(api, 'interface', lambda obj, iid: obj)
    monkeypatch.setattr(api, 'out_int', lambda: NS(value=0))
    class Session(images.Session):
        def __enter__(self):
            self.app, self.constants = app, dict(C, swIsometricView=7)
            self.owned, self.load_warnings, self.previous = {}, {}, previous.path
            self.pythoncom = NS(CoUninitialize=lambda: None)
            return self
    return app, model, previous, view, captures, Session


@pytest.mark.parametrize('scale', [1e-12, 1., 1e12])
def test_isometric_validation_uses_model_to_pixel_transform_at_any_zoom(native, scale):
    view = native[3]
    view.Orientation3 = NS(ArrayData=transpose_rotation(ISO))
    view.Scale2 = scale
    assert images._isometric(view, native[0].GetMathUtility())


def test_isometric_validation_does_not_accept_an_isometric_orientation_if_pixel_view_is_front(native):
    view = NS(Orientation3=NS(ArrayData=ISO), Transform=NS(ArrayData=IDENTITY))
    assert not images._isometric(view, native[0].GetMathUtility())


def test_native_view_is_redrawn_before_projection_is_validated(native):
    model, view = native[1], native[3]
    pending = []
    model.ShowNamedView2 = lambda *args: pending.append(transpose_rotation(ISO))
    def redraw():
        if pending: view.Orientation3 = NS(ArrayData=pending.pop())
    model.GraphicsRedraw2 = redraw
    report = images.collect_previews(result_for(native), session_factory=native[-1])
    assert all(p.png for p in report.values())
    assert_restored(native)


def result_for(native):
    model = native[1]
    return bom.Result(native[2].path, 'Default',
                      (bom.Part('Part', model.path, 'Thin', 2, 1., 2.),
                       bom.Part('Part', model.path, 'Thick', 1, 3., 3.)),
                      (bom.Node('1', '', 1, 'Part', 'P1', 'Parça', model.path, 'Thin', 1.),
                       bom.Node('2', '', 1, 'Part', 'P2', 'Parça', model.path, 'Thin', 1.),
                       bom.Node('3', '', 1, 'Part', 'P3', 'Parça', model.path, 'Thick', 3.)),
                      0, 0, False)


def assert_restored(native):
    app, model, previous, view, captures, _ = native
    assert app.ActiveDoc is previous and model.active == 'Original'
    assert model.Visible is False and view.perspective is True
    assert tuple(view.Orientation3.ArrayData) == IDENTITY and view.Scale2 == 1.7
    assert tuple(view.Translation3.ArrayData) == (.1, -.2, 0)
    assert model.SelectionManager.depth == 0 and not app.closed and not app.saves
    assert all(not path.exists() for _, _, path in captures)


def test_images_match_exact_configuration_and_restore_source_view_and_selection(native):
    report = images.collect_previews(result_for(native), session_factory=native[-1])
    thin = report[bom.sw._part_key(native[1].path, 'Thin')]
    thick = report[bom.sw._part_key(native[1].path, 'Thick')]
    with Image.open(BytesIO(thin.png)) as a, Image.open(BytesIO(thick.png)) as b:
        assert a.format == b.format == 'PNG' and a.size == b.size == (264, 200)
        assert a.getpixel((132, 100)) == (0, 0, 255)
        assert b.getpixel((132, 100)) == (0, 128, 0)
    assert [config for _, config, _ in native[4]] == ['Thin', 'Thick']
    assert_restored(native)


def test_repeated_part_configuration_is_captured_once(native):
    result = result_for(native)
    result = replace(result, parts=(result.parts[0], result.parts[0]))
    report = images.collect_previews(result, session_factory=native[-1])
    assert len(report) == len(native[4]) == 1
    assert_restored(native)


def test_named_camera_is_reassociated_after_capture(native):
    camera = NS(name='Original named camera')
    native[3].Camera = camera
    show = native[1].ShowNamedView2
    def change_camera(*args):
        native[3].Camera = None
        show(*args)
    native[1].ShowNamedView2 = change_camera
    report = images.collect_previews(result_for(native), session_factory=native[-1])
    assert all(p.png for p in report.values()) and native[3].Camera is camera
    assert_restored(native)


def test_only_newly_opened_document_is_closed_after_its_picture(native):
    app, model, previous, _, _, factory = native
    model.opened = False
    report = images.collect_previews(replace(result_for(native), parts=(result_for(native).parts[0],)),
                                     session_factory=factory)
    assert all(p.png for p in report.values()) and app.closed == [Path(model.path).name]
    assert not model.opened and previous.opened and app.ActiveDoc is previous and not app.saves


def test_missing_connection_records_every_missing_image(native):
    class Unavailable:
        def __enter__(self): raise RuntimeError('SOLIDWORKS is not open')
        def __exit__(self, *args): pass
    result = result_for(native)
    report = images.collect_previews(result, session_factory=Unavailable)
    assert len(report) == len(result.parts)
    assert all(not p.png and 'SOLIDWORKS is not open' in p.note for p in report.values())


def test_raw_dispatch_property_write_uses_dispatch_propertyput():
    writes = []
    class Raw:
        def GetIDsOfNames(self, locale, name): return name
        def Invoke(self, member, locale, flags, result, *args): writes.append((member, flags, args))
    images._put(Raw(), 'Visible', True)
    assert writes == [('Visible', 4, (True,))]


@pytest.mark.parametrize('fault', ['missing_configuration', 'failed_bitmap', 'not_isometric', 'needs_rebuild'])
def test_unreadable_or_wrong_view_is_explicit_without_losing_bom_rows(native, fault):
    result = result_for(native)
    if fault == 'missing_configuration':
        result = replace(result, parts=(replace(result.parts[0], configuration='Missing'),))
    if fault == 'failed_bitmap': native[1].SaveBMP = lambda *args: False
    if fault == 'not_isometric': native[1].ShowNamedView2 = lambda *args: None
    if fault == 'needs_rebuild': native[1].Extension.NeedsRebuild2 = 1
    report = images.collect_previews(result, session_factory=native[-1])
    assert len(report) == len(result.parts) and all(not p.png and 'alınamadı' in p.note for p in report.values())
    assert_restored(native)


def test_cancel_during_bitmap_restores_everything_and_discards_images(native):
    stop = Event()
    save = native[1].SaveBMP
    def cancel(*args):
        answer = save(*args)
        stop.set()
        return answer
    native[1].SaveBMP = cancel
    with pytest.raises(bom.Cancelled):
        images.collect_previews(result_for(native), stop, session_factory=native[-1])
    assert_restored(native)


def test_view_restore_failure_aborts_instead_of_claiming_success(native, monkeypatch):
    original = images._put
    def fail(obj, name, value):
        if name == 'Scale2' and value == 1.7: raise RuntimeError('Restore failed')
        return original(obj, name, value)
    monkeypatch.setattr(images, '_put', fail)
    with pytest.raises(api.RestoreError, match='görünümü'):
        images.collect_previews(result_for(native), session_factory=native[-1])
    assert native[0].ActiveDoc is native[2] and not native[0].saves


def test_xlsx_contains_embedded_pngs_in_both_sheets_with_matching_rows(native, tmp_path):
    result = result_for(native)
    report = images.collect_previews(result, session_factory=native[-1])
    path = tmp_path/'BOM.xlsx'
    bom.export_excel(result, path, previews=report)
    book = load_workbook(path)
    try:
        flat, tree = book.worksheets
        assert flat['J4'].value == tree['L4'].value == 'Kaynak parça · izometrik'
        assert len(flat._images) == 2 and len(tree._images) == 3
        assert [(p.anchor._from.col, p.anchor._from.row) for p in flat._images] == [(9, 4), (9, 5)]
        assert [(p.anchor._from.col, p.anchor._from.row) for p in tree._images] == [(11, 4), (11, 5), (11, 6)]
        assert flat['C5'].value == 2 and flat['H5'].value == 2. and flat['E1'].value == 5.
        assert flat.row_dimensions[5].height == 108 and flat.column_dimensions['J'].width == 28
        assert all(not cell.hyperlink for row in flat for cell in row)
    finally:
        book.close()
    with zipfile.ZipFile(path) as archive:
        assert len([n for n in archive.namelist() if n.startswith('xl/media/')]) == 5


def test_missing_image_note_is_literal_and_row_quantity_is_preserved(native, tmp_path):
    result = result_for(native)
    key = bom.sw._part_key(result.parts[0].path, result.parts[0].configuration)
    path = tmp_path/'BOM.xlsx'
    bom.export_excel(result, path, previews={key: images.Preview(None, '=Unreadable')})
    book = load_workbook(path)
    try:
        assert book.worksheets[0]['J5'].value == '=Unreadable'
        assert book.worksheets[0]['J5'].data_type == 's' and book.worksheets[0]['C5'].value == 2
        assert not book.worksheets[0]._images and book.worksheets[0]['J6'].value == 'Görsel alınamadı'
    finally:
        book.close()


@pytest.mark.parametrize('cancel', [False, True])
def test_image_export_failure_or_cancellation_preserves_existing_workbook(native, tmp_path, monkeypatch, cancel):
    result = result_for(native)
    report = images.collect_previews(result, session_factory=native[-1])
    path = tmp_path/'BOM.xlsx'; path.write_bytes(b'Existing workbook')
    stop = Event()
    if cancel:
        import openpyxl.drawing.image
        original = openpyxl.drawing.image.Image
        def image(*args, **kwargs):
            value = original(*args, **kwargs)
            stop.set()
            return value
        monkeypatch.setattr(openpyxl.drawing.image, 'Image', image)
    else:
        def save(*args, **kwargs): raise OSError('Workbook is locked')
        monkeypatch.setattr('openpyxl.workbook.workbook.Workbook.save', save)
    with pytest.raises(bom.Cancelled if cancel else OSError):
        bom.export_excel(result, path, stop, previews=report)
    assert path.read_bytes() == b'Existing workbook' and not list(tmp_path.glob('.gokworks-bom-*'))


def test_worker_embeds_images_and_reports_partial_image_failures(native, tmp_path, monkeypatch):
    result = result_for(native)
    collect = images.collect_previews
    def capture(result, stop_event, **kwargs):
        report = collect(result, stop_event, session_factory=native[-1], **kwargs)
        report[bom.sw._part_key(result.parts[1].path, 'Thick')] = images.Preview(None, 'No image')
        return report
    monkeypatch.setattr(images, 'collect_previews', capture)
    events = Queue()
    path = tmp_path/'BOM.xlsx'
    bom.export_worker(result, path, events, Event())
    assert list(events.queue)[-1] == ('bom_done', dict(state='exported', path=str(path), images=1, missing_images=1))
    assert path.is_file()


@pytest.mark.parametrize('failure', [bom.Cancelled(), api.RestoreError('Restore failed')])
def test_worker_cancel_or_restore_failure_preserves_existing_workbook(native, tmp_path, monkeypatch, failure):
    def fail(*args, **kwargs): raise failure
    monkeypatch.setattr(images, 'collect_previews', fail)
    path = tmp_path/'BOM.xlsx'; path.write_bytes(b'Existing workbook')
    events = Queue()
    bom.export_worker(result_for(native), path, events, Event())
    assert list(events.queue)[-1][1]['state'] == ('cancelled' if isinstance(failure, bom.Cancelled) else 'error')
    assert path.read_bytes() == b'Existing workbook' and not list(tmp_path.glob('.gokworks-bom-*'))
