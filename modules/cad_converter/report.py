"""CSV/XLSX reports use the same rows and publish only complete files."""

import csv
import os
from pathlib import Path
from uuid import uuid4

HEADERS = (
    ('Kaynak dosya', 'source'), ('Konfigürasyon', 'configuration'), ('Format', 'format'),
    ('STEP protokolü', 'protocol'), ('Durum', 'status'), ('Açıklama', 'detail'),
    ('Çıktı dosyası', 'output'), ('Boyut (bayt)', 'size'), ('Süre (sn)', 'seconds'),
    ('SW hata kodu', 'errors'), ('SW uyarı kodu', 'warnings'),
    ('SW açma uyarı kodu', 'open_warnings'),
)


def csv_value(value):
    # File/configuration names are untrusted text, never spreadsheet formulas.
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + value
    return value


def write_report(rows, path):
    path = Path(path)
    if path.suffix.casefold() not in ('.csv', '.xlsx'):
        raise ValueError('Rapor uzantısı CSV veya XLSX olmalı.')
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name('.gokworks-' + uuid4().hex + path.suffix)
    counts = {k: sum(row['status'] == k for row in rows) for k in ('success', 'partial', 'skipped', 'error')}
    try:
        if path.suffix.casefold() == '.csv':
            with temporary.open('w', encoding='utf-8-sig', newline='') as stream:
                writer = csv.writer(stream, delimiter=';')
                writer.writerow([title for title, _ in HEADERS])
                for row in rows: writer.writerow([csv_value(row.get(key, '')) for _, key in HEADERS])
                writer.writerow([]); writer.writerow(['Özet', *[f'{key}={value}' for key, value in counts.items()]])
        else:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            workbook = Workbook(); sheet = workbook.active; sheet.title = 'CAD aktarımı'
            sheet.append([title for title, _ in HEADERS])
            for row in rows:
                sheet.append([row.get(key, '') for _, key in HEADERS])
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='244968')
            # CAD names are data; never execute formula-like file/configuration names.
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str): cell.data_type = 's'
            sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
            for column in ('A', 'F', 'G'): sheet.column_dimensions[column].width = 60
            for column in ('B', 'C', 'D', 'E', 'H', 'I', 'J', 'K'): sheet.column_dimensions[column].width = 22
            summary = workbook.create_sheet('Özet')
            for key, value in counts.items(): summary.append([key, value])
            workbook.save(temporary); workbook.close()
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
